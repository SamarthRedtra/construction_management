# Copyright (c) 2026, Construction Management
# License: MIT

from unittest.mock import patch

import frappe
from frappe.tests.utils import FrappeTestCase
from frappe.utils import flt, money_in_words, today

from construction_management.api.so_boq_progress_excel import export_sales_order_boq_progress_excel
from construction_management.so_boq_progress_print_context import build


class TestSOBOQProgressPrint(FrappeTestCase):
	def _doc(self, **kwargs):
		item = frappe._dict(
			item_code="WATERPROOFING",
			item_name="Waterproofing",
			description="Supply and apply 1 coat of Primer and 1 layer of 4mm thick membrane",
			uom="M2",
			qty=1,
			rate=165775.35,
			amount=165775.35,
			boq_item="BOQ-ITEM-1",
		)
		defaults = {
			"name": "SO-TEST-BOQ-PROGRESS-PRINT",
			"customer": "CUST-1",
			"customer_name": "M1 Construction LLC SOC",
			"company": "Test Company",
			"project": "PROJ-1",
			"currency": "AED",
			"transaction_date": today(),
			"total_taxes_and_charges": 8288.768,
			"discount_amount": 0,
			"grand_total": 174064.118,
			"in_words": money_in_words(174064.118, "AED"),
			"items": [item],
			"taxes": [frappe._dict(rate=5)],
			"terms": "",
		}
		defaults.update(kwargs)
		return frappe._dict(defaults)

	def _ledger(self):
		return [
			frappe._dict(
				boq_item="BOQ-ITEM-1",
				prev_amount=0,
				current_amount=165775.35,
				accumulated_amount=165775.35,
				retention_amount=16577.535,
				advance_deduction=0,
				prev_qty=0,
				percentage=0,
				tax_invoice_amount=0,
			)
		]

	def test_net_c_matches_print_formula_and_words_use_net_not_grand_total(self):
		doc = self._doc()
		with patch(
			"construction_management.so_boq_progress_print_context.frappe.get_all",
			return_value=self._ledger(),
		), patch(
			"construction_management.so_boq_progress_print_context.frappe.db.has_column",
			return_value=False,
		):
			ctx = build(doc)

		t = ctx["totals"]
		expected_net = t["gross_c"] - t["ret_c"] - t["adv_c"] + t["tax_c"] - t["disc_c"]
		self.assertAlmostEqual(t["gross_c"], 165775.35, places=3)
		self.assertAlmostEqual(t["ret_c"], 16577.535, places=3)
		self.assertAlmostEqual(t["net_c"], 157486.583, places=3)
		self.assertAlmostEqual(t["net_c"], expected_net, places=6)
		self.assertEqual(ctx["amount_in_words"], money_in_words(flt(t["net_c"]), "AED"))
		self.assertNotEqual(ctx["amount_in_words"], doc.in_words)

	def test_excel_export_contains_net_amount(self):
		doc = self._doc()
		with patch(
			"construction_management.so_boq_progress_print_context.frappe.get_all",
			return_value=self._ledger(),
		), patch(
			"construction_management.so_boq_progress_print_context.frappe.db.has_column",
			return_value=False,
		):
			ctx = build(doc)

		real_get_doc = frappe.get_doc

		def fake_get_doc(*args, **kwargs):
			first = args[0] if args else None
			if first == "Sales Order":
				return doc
			return real_get_doc(*args, **kwargs)

		with patch(
			"construction_management.api.so_boq_progress_excel.frappe.has_permission",
			return_value=True,
		), patch(
			"construction_management.api.so_boq_progress_excel.frappe.get_doc",
			side_effect=fake_get_doc,
		), patch(
			"construction_management.api.so_boq_progress_excel.build",
			return_value=ctx,
		):
			url = export_sales_order_boq_progress_excel(doc.name)

		self.assertTrue(url)
		self.assertIn(".xlsx", url)
		file_doc = real_get_doc("File", {"file_url": url})
		from openpyxl import load_workbook

		wb = load_workbook(file_doc.get_full_path())
		ws = wb.active
		values = []
		for row in ws.iter_rows(values_only=True):
			values.extend([str(v) for v in row if v is not None])
		joined = " ".join(values)
		self.assertIn("NET AMOUNT DUE", joined)
		self.assertIn("157486", joined.replace(",", ""))
		self.assertIn("Amount in words", joined)
