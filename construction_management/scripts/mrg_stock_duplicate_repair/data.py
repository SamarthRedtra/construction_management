"""Read immutable migration evidence and company-scoped ledger balances."""

from __future__ import annotations

import hashlib
import re
from collections import defaultdict
from pathlib import Path

import frappe
from openpyxl import load_workbook


COMPANY = "M R G INSULATION WORKS L.L.C"
POSTING_DATE = "2026-09-30"
POSTING_TIME = "00:00:01"
CUTOFF = f"{POSTING_DATE} {POSTING_TIME}"
PREFIX = "MRG-DUP-20260930-"
# Difference account for retiring duplicate legacy stock (BS liability, not P&L).
EXPENSE_ACCOUNT = "Stock Received But Not Billed - MRG"
DATA_DIRECTORY = Path(__file__).resolve().parents[2] / "data"
SOURCES = {
	"mrg_item_linking_2026_10_22.xlsx": "2e37f5fdfec254d1e22fdf17ebb8a900db0356d7535d2eeafd6e157cf7f7ae59",
	"mrg_stock_report_2026_10_07.xlsx": "204dd85d3a49c67b96a46972339d84a5b13063df522c4e71da4cbd4f4aa55225",
}


class Evidence:
	"""Load evidence once; planning performs no writes."""

	def __init__(self):
		self.items = {
			row.name: row for row in frappe.get_all(
				"Item", fields=["name", "item_name", "stock_uom", "is_stock_item",
					"has_serial_no", "has_batch_no", "disabled", "is_customer_provided_item"],
			)
		}
		self.snapshot = frappe.db.sql("""
			SELECT s.name AS voucher, c.item_code, c.warehouse, c.qty
			FROM `tabStock Reconciliation` s
			JOIN `tabStock Reconciliation Item` c ON c.parent=s.name
			WHERE s.company=%s AND s.docstatus=1 AND s.posting_date=%s
			  AND c.item_code LIKE 'CP-%%' AND s.name NOT LIKE %s
		""", (COMPANY, POSTING_DATE, PREFIX + "%"), as_dict=True)
		self.new_codes = {row.item_code for row in self.snapshot}
		self.snapshot_pairs = {(row.item_code, row.warehouse) for row in self.snapshot}
		self.balances = self._balances()
		self.future = self._future_transactions()
		self.pending = frappe.get_all(
			"Repost Item Valuation", filters={"company": COMPANY, "status": ["in", ["Queued", "In Progress", "Failed"]]},
			fields=["name", "status", "item_code", "warehouse", "voucher_type", "voucher_no"],
			order_by="name",
		)
		self.uoms = defaultdict(dict)
		for row in frappe.get_all("UOM Conversion Detail", fields=["parent", "uom", "conversion_factor"]):
			self.uoms[row.parent][normalize(row.uom)] = float(row.conversion_factor)
		self.reserved = {
			(row.item_code, row.warehouse): float(row.reserved_qty or 0)
			for row in frappe.db.sql("""
				SELECT b.item_code,b.warehouse,b.reserved_qty FROM `tabBin` b
				JOIN `tabWarehouse` w ON w.name=b.warehouse WHERE w.company=%s
			""", (COMPANY,), as_dict=True)
		}
		self.mapping, self.mapping_rows = self._mapping()
		self.source_rows = self._source_rows()

	def _balances(self):
		rows = frappe.db.sql("""
			WITH ranked AS (
				SELECT name,item_code,warehouse,qty_after_transaction,valuation_rate,
				       stock_value,voucher_type,voucher_no,posting_datetime,
				       ROW_NUMBER() OVER (PARTITION BY item_code,warehouse
				           ORDER BY posting_datetime DESC,creation DESC,name DESC) rn
				FROM `tabStock Ledger Entry`
				WHERE company=%s AND is_cancelled=0 AND posting_datetime<%s
			)
			SELECT * FROM ranked WHERE rn=1
		""", (COMPANY, CUTOFF), as_dict=True)
		return {(row.item_code, row.warehouse): row for row in rows}

	def _future_transactions(self):
		rows = frappe.db.sql("""
			SELECT item_code,warehouse,COUNT(*) AS count,MIN(posting_datetime) AS first_date
			FROM `tabStock Ledger Entry`
			WHERE company=%s AND is_cancelled=0 AND posting_datetime>=%s
			GROUP BY item_code,warehouse
		""", (COMPANY, CUTOFF), as_dict=True)
		return {(row.item_code, row.warehouse): row for row in rows}

	def _mapping(self):
		mapping, source_rows = defaultdict(set), defaultdict(list)
		book = verified_workbook("mrg_item_linking_2026_10_22.xlsx")
		try:
			headers = next(book.active.values)
			if headers[:5] != ("LPO References", "Old Item Code", "Old Item Name", "New Item Code", "New Item Name"):
				raise ValueError("Unexpected item-linking workbook layout")
			for number, row in enumerate(book.active.iter_rows(min_row=2, values_only=True), 2):
				old, new = str(row[1] or "").strip(), str(row[3] or "").strip()
				if old and new and old != new:
					mapping[new].add(old)
					source_rows[new, old].append(number)
		finally:
			book.close()
		return mapping, source_rows

	def _source_rows(self):
		new_by_name = defaultdict(list)
		for code in self.new_codes:
			new_by_name[normalize(self.items[code].item_name)].append(code)
		grouped = {}
		book = verified_workbook("mrg_stock_report_2026_10_07.xlsx")
		try:
			for number, row in enumerate(book.active.iter_rows(min_row=2, values_only=True), 2):
				if not row[0] or row[5] != COMPANY:
					continue
				matches = new_by_name.get(normalize(row[0]), [])
				if len(matches) != 1:
					continue
				key = matches[0], str(row[2]).strip()
				target = grouped.setdefault(key, {"new_code": key[0], "warehouse": key[1],
					"source_qty": 0.0, "source_uoms": set(), "source_rows": []})
				target["source_qty"] += float(row[4] or 0)
				target["source_uoms"].add(str(row[3] or "").strip())
				target["source_rows"].append(number)
		finally:
			book.close()
		return [dict(row, source_uoms=sorted(row["source_uoms"])) for _, row in sorted(grouped.items())]


def normalize(value: str) -> str:
	return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def verified_workbook(filename):
	path = DATA_DIRECTORY / filename
	if hashlib.sha256(path.read_bytes()).hexdigest() != SOURCES[filename]:
		raise ValueError(f"Source workbook checksum mismatch: {filename}")
	return load_workbook(path, read_only=True, data_only=True)
