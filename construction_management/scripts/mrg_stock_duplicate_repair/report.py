"""Produce accounting and inventory evidence for CA review."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import frappe
from openpyxl import Workbook

from .data import COMPANY


def capture_balances():
	from erpnext.accounts.report.profit_and_loss_statement.profit_and_loss_statement import execute

	filters = frappe._dict(company=COMPANY, filter_based_on="Date Range",
		period_start_date="2026-09-01", period_end_date="2026-10-31", periodicity="Monthly",
		accumulated_values=0, presentation_currency="AED", include_default_book_entries=1)
	result = execute(filters)
	accounts = {row["account"]: {"September": row.get("sep_2026", 0), "October": row.get("oct_2026", 0)}
		for row in result[1] if row.get("account") in (
			"Stock Adjustment - MRG", "Cost of Goods Sold - MRG", "Stock Expenses - MRG")}
	return {"pnl": accounts, "inventory": {date: inventory_at(date) for date in ["2026-09-30", "2026-10-31"]}}


def inventory_at(date):
	stock = frappe.db.sql("""
		WITH ranked AS (
			SELECT stock_value,ROW_NUMBER() OVER (PARTITION BY item_code,warehouse
				ORDER BY posting_datetime DESC,creation DESC,name DESC) rn
			FROM `tabStock Ledger Entry`
			WHERE company=%s AND is_cancelled=0 AND posting_date<=%s
		)
		SELECT COALESCE(SUM(stock_value),0) FROM ranked WHERE rn=1
	""", (COMPANY, date))[0][0]
	gl = frappe.db.sql("""
		SELECT COALESCE(SUM(g.debit-g.credit),0) FROM `tabGL Entry` g
		JOIN `tabAccount` a ON a.name=g.account
		WHERE g.company=%s AND g.is_cancelled=0 AND g.posting_date<=%s
		  AND a.account_type='Stock' AND a.is_group=0
	""", (COMPANY, date))[0][0]
	return {"stock_ledger": float(stock), "inventory_gl": float(gl), "gl_minus_stock": float(gl - stock)}


def write_report(directory, plan, result):
	path = Path(directory).expanduser().resolve()
	path.mkdir(parents=True, exist_ok=True)
	(path / "plan.json").write_text(json.dumps(plan, default=str, sort_keys=True, indent=2))
	(path / "result.json").write_text(json.dumps(result, default=str, sort_keys=True, indent=2))
	save_csv(path / "eligible_corrections.csv", plan["eligible"])
	save_csv(path / "exceptions.csv", plan["exceptions"])
	book = Workbook()
	sheet = book.active
	sheet.title = "Read me"
	sheet.append(["MRG local correction review", "Value"])
	for key, value in {"Site": result["site"], "Status": result["status"], "Plan SHA256": plan["plan_hash"],
		"Company": plan["company"], "Historical correction timestamp": plan["cutoff"],
		"Correction documents": len(result.get("documents", [])), **plan["summary"]}.items():
		label = "Captured valuation jobs at start" if key == "pending_revaluations" else key
		sheet.append([label, cell_value(value)])
	sheet.append(["Scope", "Verified legacy duplicates only; ambiguous rows are held in Exceptions."])
	sheet.append(["No positive legacy balance", "Informational: no legacy balance was eligible for removal; includes merged/missing old codes."])
	sheet.append(["Production", "Not run. CA must review held items, final revaluation status and inventory reconciliation."])
	if "remaining_captured_jobs" in result:
		sheet.append(["Remaining captured revaluations", len(result["remaining_captured_jobs"])])
	if "inventory_reconciled" in result:
		sheet.append(["Inventory GL matches stock ledger", result["inventory_reconciled"]])
	sheet = book.create_sheet("P&L before and after")
	sheet.append(["Account", "September before", "September after", "October before", "October after"])
	for account, before in result.get("before", {}).get("pnl", {}).items():
		after = result.get("after", result["before"])["pnl"][account]
		sheet.append([account, before["September"], after["September"], before["October"], after["October"]])
	for title, rows in [("Eligible corrections", plan["eligible"]), ("Exceptions", plan["exceptions"]),
		("Voucher checks", result.get("voucher_checks", []))]:
		add_sheet(book, title, rows)
	if "before" in result and "after" in result:
		inventory = []
		for date, before in result["before"]["inventory"].items():
			after = result["after"]["inventory"][date]
			inventory.append({"Date": date, "Stock ledger before": before["stock_ledger"],
				"Inventory GL before": before["inventory_gl"], "Difference before": before["gl_minus_stock"],
				"Stock ledger after": after["stock_ledger"], "Inventory GL after": after["inventory_gl"],
				"Difference after": after["gl_minus_stock"]})
		add_sheet(book, "Inventory reconciliation", inventory)
	if result.get("jobs"):
		add_sheet(book, "Revaluation jobs", result["jobs"])
	if result.get("accounting_jobs"):
		add_sheet(book, "Accounting rebuild jobs", result["accounting_jobs"])
	if result.get("additional_chain_gaps"):
		add_sheet(book, "Additional stale chains", result["additional_chain_gaps"])
	if result.get("additional_accounting_jobs"):
		add_sheet(book, "Additional GL jobs", result["additional_accounting_jobs"])
	if "snapshot_timing_review" in result:
		add_sheet(book, "Snapshot timing review", result["snapshot_timing_review"])
		save_csv(path / "snapshot_timing_review.csv", result["snapshot_timing_review"])
	for sheet in book:
		sheet.freeze_panes = "A2"
		sheet.auto_filter.ref = sheet.dimensions
		for column in sheet.columns:
			sheet.column_dimensions[column[0].column_letter].width = min(65, max(15,
				max(len(str(cell.value or "")) for cell in column) + 2))
		for row in sheet.iter_rows(min_row=2):
			for cell in row:
				if isinstance(cell.value, float):
					cell.number_format = "#,##0.000"
	book.save(path / "CA_review.xlsx")
	return str(path)


def snapshot_timing_review():
	"""Flag September movements after the backdated count for CA date verification."""
	return frappe.db.sql("""
		WITH snapshots AS (
			SELECT s.name,c.item_code,c.warehouse,c.qty,s.posting_date,s.posting_time,
				ROW_NUMBER() OVER (PARTITION BY c.item_code,c.warehouse
					ORDER BY s.posting_time DESC,s.creation DESC,s.name DESC) rn
			FROM `tabStock Reconciliation` s JOIN `tabStock Reconciliation Item` c ON c.parent=s.name
			WHERE s.company=%s AND s.docstatus=1 AND s.posting_date='2026-09-30'
			  AND c.item_code LIKE 'CP-%%'
		)
		SELECT c.item_code,c.warehouse,c.name AS count_voucher,c.qty AS count_quantity,
			c.posting_time AS count_time,l.voucher_type,l.voucher_no,l.posting_time AS later_time,
			l.actual_qty,l.qty_after_transaction,l.stock_value
		FROM snapshots c JOIN `tabStock Ledger Entry` l
			ON l.item_code=c.item_code AND l.warehouse=c.warehouse
		WHERE c.rn=1 AND l.company=%s AND l.is_cancelled=0 AND l.posting_date=c.posting_date
			AND l.posting_time>c.posting_time AND l.voucher_type!='Stock Reconciliation'
		ORDER BY c.item_code,c.warehouse,l.posting_datetime,l.creation
	""", (COMPANY, COMPANY), as_dict=True)


def add_sheet(book, title, rows):
	sheet = book.create_sheet(title)
	if not rows:
		sheet.append(["No rows"])
		return
	fields = list(dict.fromkeys(key for row in rows for key in row))
	sheet.append(fields)
	for row in rows:
		sheet.append([cell_value(row.get(field)) for field in fields])


def save_csv(path, rows):
	if not rows:
		path.write_text("No rows\n")
		return
	fields = list(dict.fromkeys(key for row in rows for key in row))
	with path.open("w", newline="") as stream:
		writer = csv.DictWriter(stream, fieldnames=fields)
		writer.writeheader()
		writer.writerows([{key: cell_value(value) for key, value in row.items()} for row in rows])


def cell_value(value):
	if isinstance(value, (dict, list, tuple, set)):
		return json.dumps(value, default=str)
	if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
		return "'" + value
	return value
