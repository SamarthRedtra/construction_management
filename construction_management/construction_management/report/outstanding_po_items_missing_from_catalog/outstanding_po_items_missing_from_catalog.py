"""Open Purchase Order lines that are absent from an imported catalog workbook."""

from openpyxl import load_workbook

import frappe
from frappe import _
from frappe.query_builder.functions import IfNull

from construction_management.api.controlled_procurement import _uploaded_file_path


def execute(filters=None):
	filters = frappe._dict(filters or {})
	request = _validate_filters(filters)
	sheet_names = _sheet_item_names(request.full_source_file)
	return get_columns(), _get_open_po_rows(filters, sheet_names)


def _validate_filters(filters):
	if not filters.company:
		frappe.throw(_("Company is required."))
	if not filters.catalog_request:
		frappe.throw(_("Catalog Request is required."))
	request = frappe.get_doc("Controlled Catalog Request", filters.catalog_request)
	frappe.has_permission("Controlled Catalog Request", "read", doc=request, throw=True)
	if request.company != filters.company:
		frappe.throw(_("The selected Catalog Request belongs to another company."))
	if not request.full_source_file:
		frappe.throw(_("The selected Catalog Request has no full source workbook."))
	return request


def _sheet_item_names(file_url: str) -> set[str]:
	book = load_workbook(_uploaded_file_path(file_url), read_only=True, data_only=True)
	if "Sheet2" not in book:
		frappe.throw(_("The catalog workbook must contain Sheet2."))
	sheet = book["Sheet2"]
	header = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True), ())
	try:
		description_index = [str(value or "").strip() for value in header].index("Description")
	except ValueError:
		frappe.throw(_("Sheet2 must contain a Description column."))
	return {
		_normalize(row[description_index])
		for row in sheet.iter_rows(min_row=3, values_only=True)
		if len(row) > description_index and _normalize(row[description_index])
	}


def _get_open_po_rows(filters, sheet_names: set[str]) -> list[dict]:
	po = frappe.qb.DocType("Purchase Order")
	poi = frappe.qb.DocType("Purchase Order Item")
	received_qty = IfNull(poi.received_qty, 0)
	query = (
		frappe.qb.from_(po)
		.join(poi)
		.on(poi.parent == po.name)
		.select(
			po.name.as_("purchase_order"),
			po.supplier,
			po.supplier_name,
			po.project.as_("po_project"),
			po.schedule_date.as_("po_schedule_date"),
			poi.project.as_("item_project"),
			poi.schedule_date.as_("item_schedule_date"),
			poi.item_code,
			poi.item_name,
			poi.warehouse,
			poi.uom,
			poi.qty.as_("ordered_qty"),
			received_qty.as_("received_qty"),
			(poi.qty - received_qty).as_("outstanding_qty"),
		)
		.where(po.docstatus == 1)
		.where(po.company == filters.company)
		.where(IfNull(poi.closed, 0) == 0)
		.where(IfNull(poi.delivered_by_supplier, 0) == 0)
		.where(poi.qty > received_qty)
	)
	if filters.get("purchase_order"):
		query = query.where(po.name == filters.purchase_order)
	if filters.get("supplier"):
		query = query.where(po.supplier == filters.supplier)
	if filters.get("project"):
		query = query.where((po.project == filters.project) | (poi.project == filters.project))
	rows = query.orderby(po.schedule_date).orderby(po.name).orderby(poi.idx).run(as_dict=True)
	return [
		{
			**row,
			"project": row.item_project or row.po_project,
			"schedule_date": row.item_schedule_date or row.po_schedule_date,
		}
		for row in rows
		if _normalize(row.item_name) not in sheet_names
	]


def get_columns():
	return [
		{"fieldname": "purchase_order", "label": _("Purchase Order"), "fieldtype": "Link", "options": "Purchase Order", "width": 170},
		{"fieldname": "supplier", "label": _("Supplier"), "fieldtype": "Link", "options": "Supplier", "width": 160},
		{"fieldname": "supplier_name", "label": _("Supplier Name"), "fieldtype": "Data", "width": 180},
		{"fieldname": "project", "label": _("Project"), "fieldtype": "Link", "options": "Project", "width": 120},
		{"fieldname": "schedule_date", "label": _("Schedule Date"), "fieldtype": "Date", "width": 110},
		{"fieldname": "item_code", "label": _("Item Code"), "fieldtype": "Link", "options": "Item", "width": 170},
		{"fieldname": "item_name", "label": _("Item Name"), "fieldtype": "Data", "width": 240},
		{"fieldname": "warehouse", "label": _("Warehouse"), "fieldtype": "Link", "options": "Warehouse", "width": 160},
		{"fieldname": "uom", "label": _("UOM"), "fieldtype": "Link", "options": "UOM", "width": 80},
		{"fieldname": "ordered_qty", "label": _("Ordered Qty"), "fieldtype": "Float", "width": 110},
		{"fieldname": "received_qty", "label": _("Received Qty"), "fieldtype": "Float", "width": 110},
		{"fieldname": "outstanding_qty", "label": _("Outstanding Qty"), "fieldtype": "Float", "width": 120},
	]


def _normalize(value) -> str:
	return " ".join(str(value or "").split()).casefold()
