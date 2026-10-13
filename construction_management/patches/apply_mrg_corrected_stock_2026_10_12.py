"""Apply the approved 12 October 2026 MRG controlled-catalog stock correction."""

from __future__ import annotations

import csv
import hashlib
from collections import defaultdict
from contextlib import contextmanager
from io import StringIO
from pathlib import Path

import frappe
from frappe.utils import flt, getdate
from frappe.utils.file_manager import save_file
from openpyxl import load_workbook

from erpnext.stock.doctype.stock_reconciliation.stock_reconciliation import StockReconciliation


COMPANY = "M R G INSULATION WORKS L.L.C"
POSTING_DATE = "2026-10-12"
POSTING_TIME = "00:00:00"
SOURCE_PATH = Path(__file__).resolve().parents[1] / "data" / "mrg_corrected_stock_2026_10_12.xlsx"
SOURCE_SHA256 = "f9b6a38d258483d825c6a9098b01688d1859022b74caf3b911b064a6f75af398"
SHEET_NAME = "MRG - Qty Revision"
EXPECTED_HEADERS = (
	"New Item Name",
	"New Group",
	"Warehouse",
	"Stock UOM",
	"Corrected Qty",
	"Company",
)


def execute():
	"""Run only on sites that contain the MRG company."""
	if not frappe.db.exists("Company", COMPANY):
		return

	result = run()
	frappe.logger("controlled_procurement").info({"event": "mrg_corrected_stock_2026_10_12", **result})


def run() -> dict:
	"""Validate and reconcile all corrected stock balances in the packaged workbook."""
	frappe.set_user("Administrator")
	result = preflight()
	if result["errors"]:
		frappe.throw(_preflight_error_message(result["errors"]))

	mismatches = _mismatches(result["targets"])
	with _allow_negative_stock_for_reconciliations():
		documents = _create_reconciliations(mismatches)
		if documents:
			for document in documents:
				document.submit()

	audit_file = _save_audit(result, mismatches, documents)
	_save_source_workbook(documents)
	return {
		"posting_date": POSTING_DATE,
		"source_rows": result["source_rows"],
		"target_lines": result["target_lines"],
		"matched_balances": result["target_lines"] - len(mismatches),
		"reconciliation_balances": len(mismatches),
		"stock_reconciliations": [document.name for document in documents],
		"audit_file": audit_file.file_url,
		"no_op": not documents,
	}


def preflight() -> dict:
	"""Build exact item/warehouse targets without writing stock transactions."""
	rows = _source_rows()
	items = _controlled_stock_items()
	warehouses = _active_mrg_warehouses()
	targets, errors = _targets(rows, items, warehouses)
	_apply_valuation_rates(targets)
	errors.extend(
		{
			"source_rows": target["source_rows"],
			"reason": "No existing Bin valuation rate or approved Item Tree rate",
		}
		for target in targets
		if target["valuation_rate"] is None
	)
	return {
		"source_rows": len(rows),
		"target_lines": len(targets),
		"warehouses": len({target["warehouse"] for target in targets}),
		"targets": targets,
		"errors": errors,
	}


def _source_rows() -> list[dict]:
	"""Read the immutable correction workbook and validate its layout."""
	if not SOURCE_PATH.exists():
		frappe.throw(f"Missing stock correction workbook: {SOURCE_PATH.name}")
	if hashlib.sha256(SOURCE_PATH.read_bytes()).hexdigest() != SOURCE_SHA256:
		frappe.throw(f"Unexpected content in stock correction workbook: {SOURCE_PATH.name}")

	book = load_workbook(SOURCE_PATH, read_only=True, data_only=True)
	try:
		if SHEET_NAME not in book.sheetnames:
			frappe.throw(f"Workbook must contain sheet '{SHEET_NAME}'.")
		sheet = book[SHEET_NAME]
		headers = tuple(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))[:6])
		if headers != EXPECTED_HEADERS:
			frappe.throw("Stock correction workbook headers do not match the approved layout.")

		rows = []
		for source_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
			item_name, item_group, warehouse, report_uom, quantity, company = values[:6]
			if not any(values[:6]):
				continue
			if not item_name or not warehouse or quantity is None or company != COMPANY:
				frappe.throw(f"Invalid stock correction row {source_row}.")
			try:
				quantity = float(quantity)
			except (TypeError, ValueError):
				frappe.throw(f"Stock correction quantity must be numeric on row {source_row}.")
			rows.append({
				"source_row": source_row,
				"item_name": str(item_name).strip(),
				"item_group": str(item_group or "").strip(),
				"warehouse": str(warehouse).strip(),
				"report_uom": str(report_uom or "").strip(),
				"qty": flt(quantity),
			})
	finally:
		book.close()

	if len(rows) != 192:
		frappe.throw(f"Expected 192 populated stock correction rows, found {len(rows)}.")
	return rows


def _controlled_stock_items() -> dict[str, list[dict]]:
	items = defaultdict(list)
	for item in frappe.get_all(
		"Item",
		filters={"disabled": 0, "controlled_procurement_catalog": 1, "is_stock_item": 1},
		fields=["name", "item_name", "stock_uom"],
		limit_page_length=0,
	):
		items[_key(item.item_name)].append(item)
	return items


def _active_mrg_warehouses() -> set[str]:
	return set(
		frappe.get_all(
			"Warehouse",
			filters={"company": COMPANY, "is_group": 0, "disabled": 0},
			pluck="name",
			limit_page_length=0,
		)
	)


def _targets(rows: list[dict], items: dict, warehouses: set[str]) -> tuple[list[dict], list[dict]]:
	"""Resolve catalog items by name only and aggregate duplicate report rows."""
	targets, errors = {}, []
	for row in rows:
		matches = items.get(_key(row["item_name"]), [])
		if len(matches) != 1:
			errors.append({**row, "reason": "Controlled stock item was not uniquely resolved by item name"})
			continue
		if row["warehouse"] not in warehouses:
			errors.append({**row, "reason": "Active MRG warehouse was not found"})
			continue

		item = matches[0]
		key = (item.name, row["warehouse"])
		target = targets.setdefault(key, {
			"item_code": item.name,
			"item_name": item.item_name,
			"stock_uom": item.stock_uom,
			"warehouse": row["warehouse"],
			"qty": 0.0,
			"source_rows": [],
			"valuation_rate": None,
		})
		target["qty"] += row["qty"]
		target["source_rows"].append(row["source_row"])

	return sorted(targets.values(), key=lambda target: (target["warehouse"], target["item_code"])), errors


def _apply_valuation_rates(targets: list[dict]) -> None:
	"""Use current Bin rates first, then the approved Item Tree buying rate."""
	from construction_management.api import mrg_stock_snapshot

	tree_rates, _, _ = mrg_stock_snapshot._tree_metadata(mrg_stock_snapshot.ITEM_TREE_PATH)
	bins = {
		(row.item_code, row.warehouse): flt(row.valuation_rate)
		for row in frappe.get_all(
			"Bin",
			fields=["item_code", "warehouse", "valuation_rate"],
			limit_page_length=0,
		)
	}
	for target in targets:
		target["valuation_rate"] = bins.get((target["item_code"], target["warehouse"])) or tree_rates.get(
			_key(target["item_name"])
		)


def _mismatches(targets: list[dict]) -> list[dict]:
	bins = {
		(row.item_code, row.warehouse): flt(row.actual_qty)
		for row in frappe.get_all(
			"Bin",
			fields=["item_code", "warehouse", "actual_qty"],
			limit_page_length=0,
		)
	}
	return [
		{**target, "actual_qty": bins.get((target["item_code"], target["warehouse"]), 0.0)}
		for target in targets
		if abs(bins.get((target["item_code"], target["warehouse"]), 0.0) - target["qty"]) > 0.000001
	]


def _create_reconciliations(targets: list[dict]) -> list:
	by_warehouse = defaultdict(list)
	for target in targets:
		by_warehouse[target["warehouse"]].append(target)

	documents = []
	for warehouse, warehouse_targets in sorted(by_warehouse.items()):
		document = frappe.new_doc("Stock Reconciliation")
		document.company = COMPANY
		document.purpose = "Stock Reconciliation"
		document.set_posting_time = 1
		document.posting_date = getdate(POSTING_DATE)
		document.posting_time = POSTING_TIME
		for target in warehouse_targets:
			document.append("items", {
				"item_code": target["item_code"],
				"warehouse": warehouse,
				"qty": target["qty"],
				"valuation_rate": target["valuation_rate"],
			})
		document.flags.ignore_permissions = True
		document.insert()
		documents.append(document)
	return documents


@contextmanager
def _allow_negative_stock_for_reconciliations():
	"""Allow approved negative balances only while this patch creates its documents."""
	original_update_stock_ledger = StockReconciliation.update_stock_ledger
	original_validate_data = StockReconciliation.validate_data

	def update_stock_ledger(document, allow_negative_stock=False):
		return original_update_stock_ledger(document, allow_negative_stock=True)

	def validate_data(document):
		try:
			return original_validate_data(document)
		except frappe.ValidationError:
			document.validation_messages = [
				message
				for message in document.validation_messages
				if not str(message).endswith("Negative Quantity is not allowed")
			]
			if document.validation_messages:
				raise frappe.ValidationError(document.validation_messages)

	StockReconciliation.update_stock_ledger = update_stock_ledger
	StockReconciliation.validate_data = validate_data
	try:
		yield
	finally:
		StockReconciliation.update_stock_ledger = original_update_stock_ledger
		StockReconciliation.validate_data = original_validate_data


def _save_audit(result: dict, mismatches: list[dict], documents: list):
	by_warehouse = {document.items[0].warehouse: document.name for document in documents}
	mismatch_keys = {(target["item_code"], target["warehouse"]) for target in mismatches}
	current_qty_by_key = {
		(target["item_code"], target["warehouse"]): target["actual_qty"] for target in mismatches
	}
	content = StringIO()
	writer = csv.writer(content)
	writer.writerow(["MRG Corrected Stock", "Value"])
	writer.writerow(["Posting Date", POSTING_DATE])
	writer.writerow(["Workbook SHA256", SOURCE_SHA256])
	writer.writerow(["Source Rows", result["source_rows"]])
	writer.writerow(["Target Balances", result["target_lines"]])
	writer.writerow([])
	writer.writerow([
		"Status", "Stock Reconciliation", "Warehouse", "Item Code", "Item Name", "Catalog UOM",
		"Target Qty", "Current Qty", "Valuation Rate", "Source Rows",
	])
	for target in result["targets"]:
		key = (target["item_code"], target["warehouse"])
		writer.writerow([
			"Reconciled" if key in mismatch_keys else "Matched",
			by_warehouse.get(target["warehouse"], ""),
			target["warehouse"],
			target["item_code"],
			target["item_name"],
			target["stock_uom"],
			target["qty"],
			current_qty_by_key.get(key, target["qty"]),
			target["valuation_rate"],
			", ".join(map(str, target["source_rows"])),
		])

	filename = f"MRG_Corrected_Stock_{POSTING_DATE}_{SOURCE_SHA256[:8]}.csv"
	existing = frappe.get_all(
		"File",
		filters={"file_name": filename, "attached_to_doctype": "Company", "attached_to_name": COMPANY},
		pluck="name",
		limit_page_length=1,
	)
	if existing:
		return frappe.get_doc("File", existing[0])
	return save_file(filename, content.getvalue().encode(), "Company", COMPANY, is_private=1)


def _save_source_workbook(documents: list) -> None:
	doctype, name = (documents[0].doctype, documents[0].name) if documents else ("Company", COMPANY)
	if frappe.db.exists(
		"File",
		{"file_name": SOURCE_PATH.name, "attached_to_doctype": doctype, "attached_to_name": name},
	):
		return
	save_file(SOURCE_PATH.name, SOURCE_PATH.read_bytes(), doctype, name, is_private=1)


def _preflight_error_message(errors: list[dict]) -> str:
	rows = ", ".join(str(error.get("source_row", error.get("source_rows", "?"))) for error in errors[:10])
	return f"MRG corrected stock preflight failed for {len(errors)} row(s): {rows}."


def _key(value) -> str:
	return " ".join(str(value or "").split()).casefold()
