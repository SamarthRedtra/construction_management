# Copyright (c) 2026, Construction Management
# License: MIT

"""Export Sales Order BOQ Progress print layout to Excel."""

from __future__ import annotations

import io

import frappe
from frappe import _
from frappe.utils import flt, now_datetime

from construction_management.so_boq_progress_print_context import build

PRINT_FORMAT = "Sales Order BOQ Progress"


@frappe.whitelist()
def export_sales_order_boq_progress_excel(sales_order: str) -> str:
	if not sales_order:
		frappe.throw(_("Sales Order is required"))
	if not frappe.has_permission("Sales Order", "read", sales_order):
		frappe.throw(_("Not permitted"), frappe.PermissionError)

	try:
		from openpyxl import Workbook
		from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
		from openpyxl.utils import get_column_letter
	except ImportError:
		frappe.throw(_("openpyxl is required for Excel export. Please install it."))

	doc = frappe.get_doc("Sales Order", sales_order)
	ctx = build(doc)
	wb = Workbook()
	ws = wb.active
	ws.title = "BOQ Progress"

	thin = Border(
		left=Side(style="thin"),
		right=Side(style="thin"),
		top=Side(style="thin"),
		bottom=Side(style="thin"),
	)
	header_fill = PatternFill("solid", fgColor="F2F2F2")
	progress_fill = PatternFill("solid", fgColor="FAFAFA")
	net_fill = PatternFill("solid", fgColor="E6E6E6")
	bold = Font(bold=True)
	wrap = Alignment(wrap_text=True, vertical="top")
	right = Alignment(horizontal="right", vertical="center")
	center = Alignment(horizontal="center", vertical="center")

	ws.merge_cells("A1:K1")
	ws["A1"] = "PROFORMA INVOICE"
	ws["A1"].font = Font(bold=True, size=16)
	ws["A1"].alignment = Alignment(horizontal="center")

	header = ctx["header"]
	ws["A3"] = header["customer_name"]
	ws["A3"].font = bold
	ws["A4"] = f"Tel: {header['customer_tel']}" if header["customer_tel"] else ""
	ws["A5"] = f"Customer TRN: {header['customer_trn']}"
	ws["A6"] = f"Subject: {header['subject']}"
	ws["G3"] = f"Date: {header['date']}"
	ws["G4"] = f"Invoice No.: {header['invoice_no']}"
	ws["G5"] = f"Project: {header['project_name']}"
	if header["company_trn"]:
		ws["G6"] = f"{header['company_trn_label']}: {header['company_trn']}"

	headers_top = [
		"#",
		"Description",
		"Payment Terms",
		"Unit",
		"Progress (Qty %)",
		"",
		"",
		"Unit Price",
		f"Amount ({ctx['currency']})",
		"",
		"",
	]
	headers_sub = [
		"",
		"",
		"",
		"",
		"Previous",
		"Current",
		"Total",
		"",
		"Previous",
		"Current",
		"Total",
	]
	for col, value in enumerate(headers_top, 1):
		cell = ws.cell(row=8, column=col, value=value)
		cell.font = bold
		cell.fill = header_fill
		cell.border = thin
		cell.alignment = center
	for col, value in enumerate(headers_sub, 1):
		cell = ws.cell(row=9, column=col, value=value)
		cell.font = bold
		cell.fill = header_fill
		cell.border = thin
		cell.alignment = center
	ws.merge_cells("E8:G8")
	ws.merge_cells("I8:K8")
	ws.merge_cells("A8:A9")
	ws.merge_cells("B8:B9")
	ws.merge_cells("C8:C9")
	ws.merge_cells("D8:D9")
	ws.merge_cells("H8:H9")

	row = 10
	for item in ctx["rows"]:
		desc_row = [
			item["idx"],
			item["description"],
			"",
			item["uom"],
			"",
			"",
			"",
			"",
			"",
			"",
			"",
		]
		for col, value in enumerate(desc_row, 1):
			cell = ws.cell(row=row, column=col, value=value)
			cell.border = thin
			cell.alignment = wrap if col == 2 else center
		ws.row_dimensions[row].height = 42

		row += 1
		progress_row = [
			"",
			"Progress",
			f"{flt(item['percentage'])}%",
			"",
			item["prev_qty"],
			item["current_qty"],
			item["total_qty"],
			item["rate"],
			item["prev_amount"],
			item["current_amount"],
			item["total_amount"],
		]
		for col, value in enumerate(progress_row, 1):
			cell = ws.cell(row=row, column=col, value=value)
			cell.border = thin
			cell.fill = progress_fill
			cell.alignment = right if col >= 5 else center
			if col >= 8:
				cell.number_format = "#,##0.00"
			elif col >= 5:
				cell.number_format = "#,##0.00"
		row += 1

	t = ctx["totals"]
	row = _totals_row(ws, row, "TOTAL WORK DONE", t["gross_p"], t["gross_c"], t["gross_t"], thin, bold)
	row = _totals_row(
		ws,
		row,
		f"Retention ({t['ret_percent']}%)",
		-t["ret_p"],
		-t["ret_c"],
		-t["ret_t"],
		thin,
	)
	row = _totals_row(ws, row, "Advance Deduction", -t["adv_p"], -t["adv_c"], -t["adv_t"], thin)
	row = _totals_row(ws, row, f"VAT ({t['vat_rate']}%)", t["tax_p"], t["tax_c"], t["tax_t"], thin)
	row = _totals_row(ws, row, "Additional Deduction / Discount", 0, -t["disc_c"], -t["disc_c"], thin)
	row = _totals_row(ws, row, "Previous Payment (Invoiced)", -t["prev_paid"], 0, -t["prev_paid"], thin)

	ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
	label = ws.cell(row=row, column=1, value="NET AMOUNT DUE")
	label.font = bold
	label.alignment = Alignment(horizontal="right", vertical="center")
	label.fill = net_fill
	ws.merge_cells(start_row=row, start_column=10, end_row=row, end_column=11)
	net_cell = ws.cell(row=row, column=10, value=t["net_c"])
	net_cell.font = Font(bold=True, size=12)
	net_cell.number_format = "#,##0.00"
	net_cell.fill = net_fill
	net_cell.alignment = right
	for col in range(1, 12):
		ws.cell(row=row, column=col).border = thin
		ws.cell(row=row, column=col).fill = net_fill
	row += 2

	ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=11)
	ws.cell(row=row, column=1, value=f"Amount in words: {ctx['amount_in_words']}")
	row += 1
	ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=11)
	ws.cell(row=row, column=1, value=f"Note: All cheques to be made in the name of {ctx['company']}")

	widths = [6, 42, 14, 10, 12, 12, 12, 14, 14, 14, 14]
	for i, width in enumerate(widths, 1):
		ws.column_dimensions[get_column_letter(i)].width = width

	output = io.BytesIO()
	wb.save(output)
	output.seek(0)

	file_name = f"{PRINT_FORMAT}_{sales_order}_{now_datetime().strftime('%Y%m%d_%H%M%S')}.xlsx"
	file_args = {
		"doctype": "File",
		"file_name": file_name,
		"content": output.getvalue(),
		"is_private": 1,
	}
	if frappe.db.exists("Sales Order", sales_order):
		file_args["attached_to_doctype"] = "Sales Order"
		file_args["attached_to_name"] = sales_order
	file_doc = frappe.get_doc(file_args)
	file_doc.save(ignore_permissions=True)
	return file_doc.file_url


def _totals_row(ws, row, label, prev, current, total, thin, font=None):
	from openpyxl.styles import Alignment

	ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
	cell = ws.cell(row=row, column=1, value=label)
	cell.alignment = Alignment(horizontal="right")
	if font:
		cell.font = font
	values = (prev, current, total)
	for offset, value in enumerate(values, 9):
		num = ws.cell(row=row, column=offset, value=value)
		num.number_format = "#,##0.00"
		num.alignment = Alignment(horizontal="right")
		if font:
			num.font = font
	for col in range(1, 12):
		ws.cell(row=row, column=col).border = thin
	return row + 1
