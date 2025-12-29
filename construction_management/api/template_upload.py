# Copyright (c) 2024, Construction Management
# License: MIT

"""
BOQ Template Upload Processor

Processes Excel templates to create BOQ Bills and Items.

Properties validated:
- Property 21: Template Upload Validation and Creation

Requirements: 11.1, 11.2, 11.3, 11.4, 11.5, 11.6
"""

import frappe
from frappe import _
from frappe.utils import flt, cint
from typing import List, Dict, Tuple, Optional
from dataclasses import dataclass, field
import io


@dataclass
class ValidationError:
	"""Validation error details"""
	row: int
	column: str
	message: str


@dataclass
class ValidationResult:
	"""Template validation result"""
	is_valid: bool
	errors: List[ValidationError] = field(default_factory=list)
	warnings: List[str] = field(default_factory=list)
	row_count: int = 0
	valid_row_count: int = 0


@dataclass
class CreationResult:
	"""BOQ creation result"""
	success: bool
	bills_created: int = 0
	items_created: int = 0
	errors: List[str] = field(default_factory=list)
	created_bills: List[str] = field(default_factory=list)
	created_items: List[str] = field(default_factory=list)


# Required columns in template
REQUIRED_COLUMNS = ["bill_no", "description", "unit", "quantity", "rate"]
OPTIONAL_COLUMNS = ["item_code", "total_estimated_cost", "estimated_material_cost", "estimated_labour_cost", 
					"estimated_subcontract_cost", "estimated_asset_cost", "estimated_other_cost"]


def parse_template(file_content: bytes, filename: str = None) -> Tuple[List[Dict], List[str]]:
	"""
	Parse Excel template into structured data.
	
	Requirements: 11.2
	
	Args:
		file_content: File content as bytes
		filename: Optional filename for format detection
		
	Returns:
		Tuple of (data rows, column headers)
	"""
	try:
		import openpyxl
		from io import BytesIO
		
		wb = openpyxl.load_workbook(BytesIO(file_content), data_only=True)
		ws = wb.active
		
		# Get headers from first row
		headers = []
		for cell in ws[1]:
			if cell.value:
				# Normalize header names
				header = str(cell.value).lower().strip().replace(" ", "_")
				headers.append(header)
			else:
				headers.append(None)
		
		# Parse data rows
		data = []
		for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
			row_data = {}
			has_data = False
			
			for col_idx, cell in enumerate(row):
				if col_idx < len(headers) and headers[col_idx]:
					value = cell.value
					if value is not None:
						has_data = True
					row_data[headers[col_idx]] = value
			
			if has_data:
				row_data["_row_number"] = row_idx
				data.append(row_data)
		
		return data, headers
		
	except ImportError:
		frappe.throw(_("openpyxl library is required for Excel file processing"))
	except Exception as e:
		frappe.throw(_("Error parsing template: {0}").format(str(e)))


def validate_template_data(data: List[Dict], headers: List[str]) -> ValidationResult:
	"""
	Validate all rows in template.
	
	Property 21: Template Upload Validation and Creation
	Requirements: 11.3
	
	Args:
		data: Parsed data rows
		headers: Column headers
		
	Returns:
		ValidationResult with errors and warnings
	"""
	result = ValidationResult(is_valid=True, row_count=len(data))
	
	# Check required columns exist
	missing_columns = []
	for col in REQUIRED_COLUMNS:
		if col not in headers:
			missing_columns.append(col)
	
	if missing_columns:
		result.is_valid = False
		result.errors.append(ValidationError(
			row=0,
			column="headers",
			message=_("Missing required columns: {0}").format(", ".join(missing_columns))
		))
		return result
	
	# Validate each row
	seen_item_codes = set()
	
	for row in data:
		row_num = row.get("_row_number", 0)
		
		# Validate bill_no
		if not row.get("bill_no"):
			result.errors.append(ValidationError(
				row=row_num,
				column="bill_no",
				message=_("Bill No is required")
			))
		
		# Validate description
		if not row.get("description"):
			result.errors.append(ValidationError(
				row=row_num,
				column="description",
				message=_("Description is required")
			))
		
		# Validate unit
		if not row.get("unit"):
			result.errors.append(ValidationError(
				row=row_num,
				column="unit",
				message=_("Unit is required")
			))
		
		# Validate quantity
		qty = row.get("quantity")
		if qty is None:
			result.errors.append(ValidationError(
				row=row_num,
				column="quantity",
				message=_("Quantity is required")
			))
		elif not isinstance(qty, (int, float)) or flt(qty) < 0:
			result.errors.append(ValidationError(
				row=row_num,
				column="quantity",
				message=_("Quantity must be a positive number")
			))
		
		# Validate rate
		rate = row.get("rate")
		if rate is None:
			result.errors.append(ValidationError(
				row=row_num,
				column="rate",
				message=_("Rate is required")
			))
		elif not isinstance(rate, (int, float)) or flt(rate) < 0:
			result.errors.append(ValidationError(
				row=row_num,
				column="rate",
				message=_("Rate must be a positive number")
			))
		
		# Check for duplicate item codes
		item_code = row.get("item_code")
		if item_code:
			if item_code in seen_item_codes:
				result.errors.append(ValidationError(
					row=row_num,
					column="item_code",
					message=_("Duplicate item code: {0}").format(item_code)
				))
			else:
				seen_item_codes.add(item_code)
		
		# Validate estimated costs (if provided)
		for cost_field in ["total_estimated_cost", "estimated_material_cost", "estimated_labour_cost", 
						   "estimated_subcontract_cost", "estimated_asset_cost", "estimated_other_cost"]:
			cost_value = row.get(cost_field)
			if cost_value is not None and not isinstance(cost_value, (int, float)):
				result.warnings.append(
					_("Row {0}: Invalid {1} value, will be set to 0").format(row_num, cost_field)
				)
	
	# Set validity
	result.is_valid = len(result.errors) == 0
	result.valid_row_count = len(data) if result.is_valid else 0
	
	return result


def create_boq_records(project: str, data: List[Dict]) -> CreationResult:
	"""
	Create BOQ Bills and Items from template data.
	
	Property 21: The number of BOQ Items created SHALL equal the number of valid 
	data rows in the template, and each item's fields SHALL match the template values.
	
	Requirements: 11.4
	
	Args:
		project: Project name
		data: Validated template data
		
	Returns:
		CreationResult with created records
	"""
	result = CreationResult(success=True)
	
	# Get or create Project BOQ
	project_boq = frappe.db.get_value("Project BOQ", {"project": project})
	if not project_boq:
		boq = frappe.new_doc("Project BOQ")
		boq.project = project
		boq.boq_name = f"BOQ - {project}"
		boq.insert(ignore_permissions=True)
		project_boq = boq.name
	
	# Group data by bill_no
	bills_data = {}
	for row in data:
		bill_no = str(row.get("bill_no", "")).strip()
		if bill_no not in bills_data:
			bills_data[bill_no] = []
		bills_data[bill_no].append(row)
	
	# Create bills and items
	for bill_no, items in bills_data.items():
		try:
			# Check if bill exists
			existing_bill = frappe.db.get_value(
				"BOQ Bill", 
				{"project": project, "bill_no": bill_no}
			)
			
			if existing_bill:
				bill_name = existing_bill
			else:
				# Create new bill
				bill = frappe.new_doc("BOQ Bill")
				bill.project = project
				bill.project_boq = project_boq
				bill.bill_no = bill_no
				bill.description = f"Bill {bill_no}"
				bill.insert(ignore_permissions=True)
				bill_name = bill.name
				result.bills_created += 1
				result.created_bills.append(bill_name)
			
			# Create items for this bill
			for item_data in items:
				try:
					item = frappe.new_doc("BOQ Item")
					item.parent_bill = bill_name
					item.project = project
					item.project_boq = project_boq
					
					# Set required fields
					item.description = str(item_data.get("description", "")).strip()
					item.unit = str(item_data.get("unit", "Nos")).strip()
					item.total_qty = flt(item_data.get("quantity", 0))
					item.rate = flt(item_data.get("rate", 0))
					
					# Set optional fields
					if item_data.get("item_code"):
						item.item_code = str(item_data.get("item_code")).strip()
					
					# Set estimated costs - support both total cost and breakdown
					# If total_estimated_cost is provided and no breakdown, use it directly
					# If breakdown is provided, it will be auto-calculated by the doctype
					total_cost = flt(item_data.get("total_estimated_cost", 0))
					material_cost = flt(item_data.get("estimated_material_cost", 0))
					labour_cost = flt(item_data.get("estimated_labour_cost", 0))
					subcontract_cost = flt(item_data.get("estimated_subcontract_cost", 0))
					asset_cost = flt(item_data.get("estimated_asset_cost", 0))
					other_cost = flt(item_data.get("estimated_other_cost", 0))
					
					breakdown_total = material_cost + labour_cost + subcontract_cost + asset_cost + other_cost
					
					if breakdown_total > 0:
						# Use breakdown values - total will be auto-calculated
						item.estimated_material_cost = material_cost
						item.estimated_labour_cost = labour_cost
						item.estimated_subcontract_cost = subcontract_cost
						item.estimated_asset_cost = asset_cost
						item.estimated_other_cost = other_cost
					elif total_cost > 0:
						# Only total cost provided - put it in "other" cost for now
						# This allows users to import with just total cost
						item.estimated_other_cost = total_cost
						item.estimated_material_cost = 0
						item.estimated_labour_cost = 0
						item.estimated_subcontract_cost = 0
						item.estimated_asset_cost = 0
					
					item.insert(ignore_permissions=True)
					result.items_created += 1
					result.created_items.append(item.name)
					
				except Exception as e:
					result.errors.append(
						_("Error creating item from row {0}: {1}").format(
							item_data.get("_row_number", "?"), str(e)
						)
					)
					result.success = False
					
		except Exception as e:
			result.errors.append(
				_("Error creating bill {0}: {1}").format(bill_no, str(e))
			)
			result.success = False
	
	return result


@frappe.whitelist()
def upload_boq_template(project: str, file_url: str = None, file_content: str = None) -> dict:
	"""
	API endpoint to upload and process BOQ template.
	
	Requirements: 11.1, 11.5, 11.6
	
	Args:
		project: Project name
		file_url: URL of uploaded file
		file_content: Base64 encoded file content
		
	Returns:
		dict with upload results
	"""
	import base64
	
	# Get file content
	if file_url:
		# Get file from URL
		file_doc = frappe.get_doc("File", {"file_url": file_url})
		content = file_doc.get_content()
	elif file_content:
		# Decode base64 content
		content = base64.b64decode(file_content)
	else:
		frappe.throw(_("No file provided"))
	
	# Parse template
	data, headers = parse_template(content)
	
	if not data:
		return {
			"success": False,
			"message": _("No data found in template"),
			"errors": []
		}
	
	# Validate data
	validation = validate_template_data(data, headers)
	
	if not validation.is_valid:
		return {
			"success": False,
			"message": _("Validation failed"),
			"errors": [
				{"row": e.row, "column": e.column, "message": e.message}
				for e in validation.errors
			],
			"warnings": validation.warnings,
			"row_count": validation.row_count
		}
	
	# Create records
	creation = create_boq_records(project, data)
	
	return {
		"success": creation.success,
		"message": _("Created {0} bills and {1} items").format(
			creation.bills_created, creation.items_created
		),
		"bills_created": creation.bills_created,
		"items_created": creation.items_created,
		"created_bills": creation.created_bills,
		"created_items": creation.created_items,
		"errors": creation.errors,
		"warnings": validation.warnings
	}


@frappe.whitelist()
def get_template_format() -> dict:
	"""
	Get the expected template format for download.
	
	Returns:
		dict with column information
	"""
	return {
		"required_columns": REQUIRED_COLUMNS,
		"optional_columns": OPTIONAL_COLUMNS,
		"sample_data": [
			{
				"bill_no": "Bill 1",
				"item_code": "ITEM-001",
				"description": "Excavation work",
				"unit": "CuM",
				"quantity": 100,
				"rate": 500,
				"total_estimated_cost": 50000,  # Option 1: Just total cost
				"estimated_material_cost": 0,
				"estimated_labour_cost": 0,
				"estimated_subcontract_cost": 0,
				"estimated_asset_cost": 0,
				"estimated_other_cost": 0
			},
			{
				"bill_no": "Bill 1",
				"item_code": "ITEM-002",
				"description": "Concrete work",
				"unit": "CuM",
				"quantity": 50,
				"rate": 8000,
				"total_estimated_cost": 0,  # Option 2: Breakdown (total auto-calculated)
				"estimated_material_cost": 200000,
				"estimated_labour_cost": 100000,
				"estimated_subcontract_cost": 50000,
				"estimated_asset_cost": 20000,
				"estimated_other_cost": 30000
			}
		],
		"notes": [
			"You can provide either total_estimated_cost OR the breakdown fields",
			"If breakdown fields have values, they will be used and total will be auto-calculated",
			"If only total_estimated_cost is provided, it will be stored as 'Other Cost'"
		]
	}


@frappe.whitelist()
def download_template() -> str:
	"""
	Generate and return a sample template file.
	
	Returns:
		File URL
	"""
	try:
		import openpyxl
		from openpyxl.styles import Font, PatternFill
		from io import BytesIO
		
		wb = openpyxl.Workbook()
		ws = wb.active
		ws.title = "BOQ Template"
		
		# Headers
		headers = REQUIRED_COLUMNS + OPTIONAL_COLUMNS
		header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
		header_font = Font(color="FFFFFF", bold=True)
		
		for col, header in enumerate(headers, 1):
			cell = ws.cell(row=1, column=col, value=header.replace("_", " ").title())
			cell.fill = header_fill
			cell.font = header_font
		
		# Sample data
		sample = get_template_format()["sample_data"]
		for row_idx, row_data in enumerate(sample, 2):
			for col_idx, header in enumerate(headers, 1):
				ws.cell(row=row_idx, column=col_idx, value=row_data.get(header, ""))
		
		# Adjust column widths
		for col in ws.columns:
			max_length = 0
			column = col[0].column_letter
			for cell in col:
				try:
					if len(str(cell.value)) > max_length:
						max_length = len(str(cell.value))
				except:
					pass
			ws.column_dimensions[column].width = max_length + 2
		
		# Save to bytes
		output = BytesIO()
		wb.save(output)
		output.seek(0)
		
		# Create file
		file_doc = frappe.get_doc({
			"doctype": "File",
			"file_name": "boq_template.xlsx",
			"content": output.getvalue(),
			"is_private": 0
		})
		file_doc.insert(ignore_permissions=True)
		
		return file_doc.file_url
		
	except ImportError:
		frappe.throw(_("openpyxl library is required for template generation"))
