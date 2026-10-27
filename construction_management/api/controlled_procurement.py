# Copyright (c) 2026, Construction Management
# License: MIT

"""Controlled catalog import and standard ERPNext procurement document creation."""

from __future__ import annotations

from collections import Counter, defaultdict
import csv
import hashlib
import json
import re
from io import BytesIO, StringIO
from pathlib import Path
from uuid import UUID

import frappe
from frappe import _
from frappe.utils.file_manager import save_file
from frappe.utils import cint, flt, getdate, now_datetime, today
from openpyxl import Workbook


ADMIN_ROLES = {"System Manager", "CEO"}
PURCHASE_ROLES = ADMIN_ROLES | {"Purchase User", "Purchase Manager"}
STOCK_ROLES = ADMIN_ROLES | {"Stock User", "Stock Manager"}
ACCOUNT_ROLES = {"Accounts User", "Accounts Manager"}
CATALOG_CEO_USERS = {"faraaz@mrggroup.ae", "salman@mrggroup.ae"}
CATALOG_ROLES = PURCHASE_ROLES | STOCK_ROLES | ACCOUNT_ROLES
OPENING_STOCK_ROLES = ACCOUNT_ROLES | ADMIN_ROLES
STOCKABLE = "Stockable"
CONSUMABLE = "Consumable"
STOCK_TYPES = {STOCKABLE, CONSUMABLE}
ASSET = "Asset"
SERVICE = "Service"
SUPPLIER_TYPES = {"Company", "Individual", "Partnership"}
ALLOCATION_FIELDS = ("project", "bill_no", "boq_item")
EMPTY_ALLOCATION_VALUES = {"", "null", "undefined"}


def _roles() -> set[str]:
	return set(frappe.get_roles())


def _require(roles: set[str]) -> None:
	if not _roles().intersection(roles):
		frappe.throw(_("Not permitted"), frappe.PermissionError)


def _as_dict(value: str | dict | None) -> dict:
	if isinstance(value, str):
		return frappe.parse_json(value)
	return value or {}


def _as_name_list(value, limit: int = 50) -> list[str]:
	if isinstance(value, str):
		value = frappe.parse_json(value) if value.strip().startswith("[") else [value]
	if not isinstance(value, list) or not value:
		frappe.throw(_("Select at least one document."))
	names = []
	for item in value:
		name = item.get("name") if isinstance(item, dict) else item
		if not name or not isinstance(name, str):
			frappe.throw(_("Invalid selection."))
		names.append(name)
	if len(names) > limit:
		frappe.throw(_("Select at most {0} documents at a time.").format(limit))
	return names


def _bulk_error_message(exc: Exception) -> str:
	message = getattr(exc, "message", None) or (exc.args[0] if exc.args else str(exc))
	if isinstance(message, (list, tuple)):
		message = message[0] if message else str(exc)
	return frappe.utils.strip_html(str(message))[:400]


def _run_bulk(items: list, worker) -> dict:
	"""Run worker per item; commit successes so one failure does not undo others."""
	results = []
	approved = 0
	failed = 0
	for item in items:
		name = item if isinstance(item, str) else (item.get("name") or "")
		try:
			out = worker(item)
			frappe.db.commit()
			status = None
			if isinstance(out, dict):
				status = out.get("status") or out.get("approval_status")
			results.append({"name": name, "ok": True, "status": status})
			approved += 1
		except Exception as exc:
			frappe.db.rollback()
			results.append({"name": name, "ok": False, "error": _bulk_error_message(exc)})
			failed += 1
	return {"results": results, "approved": approved, "failed": failed}


def _normalized_allocation(values: dict) -> dict:
	"""Remove browser placeholder values before applying BOQ allocation rules."""
	allocation = {}
	for field in ALLOCATION_FIELDS:
		value = values.get(field)
		value = value.strip() if isinstance(value, str) else value
		allocation[field] = "" if value is None or str(value).casefold() in EMPTY_ALLOCATION_VALUES else value
	return allocation


def _catalog_item(item_code: str) -> dict:
	item = frappe.db.get_value(
		"Item", item_code,
		["name", "stock_uom", "is_stock_item", "is_fixed_asset", "controlled_procurement_catalog", "controlled_item_type"],
		as_dict=True,
	)
	if not item or not cint(item.controlled_procurement_catalog) or not item.stock_uom:
		frappe.throw(_("Item {0} is not in the controlled catalog.").format(item_code))
	return item


def _payment_terms_template(value: str | None) -> str | None:
	value = (value or "").strip()
	if value and not frappe.db.exists("Payment Terms Template", value):
		frappe.throw(_("Payment Terms Template {0} does not exist.").format(value))
	return value or None


def _order_discount(data: dict) -> dict:
	percentage = flt(data.get("additional_discount_percentage"))
	amount = flt(data.get("discount_amount"))
	if percentage < 0 or percentage > 100 or amount < 0:
		frappe.throw(_("Discount must be a non-negative amount or a percentage between 0 and 100."))
	if percentage and amount:
		frappe.throw(_("Enter either an order discount percentage or an order discount amount, not both."))
	return {"apply_discount_on": data.get("apply_discount_on") or "Net Total", "additional_discount_percentage": percentage, "discount_amount": amount}


def _set_row_type(row, item: dict) -> None:
	row.controlled_item_type = item.controlled_item_type
	meta = getattr(row, "meta", None)
	if meta and meta.has_field("controlled_procurement_catalog"):
		# _catalog_item only returns catalogue members; default 1 when payload omits the flag.
		flag = (
			item.get("controlled_procurement_catalog")
			if isinstance(item, dict)
			else getattr(item, "controlled_procurement_catalog", None)
		)
		row.controlled_procurement_catalog = cint(1 if flag is None else flag)


def _append_item(doc, row_data: dict, defaults: dict, warehouses: bool = False) -> None:
	item = _catalog_item(row_data.get("item_code"))
	if doc.doctype == "Purchase Order" and item.controlled_item_type == SERVICE:
		_ensure_service_item_default_for_company(item.name, doc.company)
	row = doc.append("items", {
		"item_code": item.name,
		"qty": flt(row_data.get("qty")),
		"uom": item.stock_uom,
		"stock_uom": item.stock_uom,
		"conversion_factor": 1,
		"rate": flt(row_data.get("rate")),
		"description": row_data.get("notes") or "",
	})
	if warehouses:
		row.s_warehouse = row_data.get("source_warehouse")
		row.t_warehouse = row_data.get("target_warehouse")
	elif doc.doctype == "Purchase Order":
		row.warehouse = row_data.get("warehouse") or doc.set_warehouse
	if doc.doctype != "Stock Entry":
		_apply_allocation(row, row_data, defaults, doc.company)
	_set_row_type(row, item)


def _validate_company_links(company: str, project: str | None = None, warehouse: str | None = None) -> None:
	"""Projects and warehouses picked in the workspace must belong to the document's company."""
	for doctype, name in (("Project", project), ("Warehouse", warehouse)):
		if name and frappe.db.get_value(doctype, name, "company") not in (company, None, ""):
			frappe.throw(_("{0} {1} does not belong to company {2}.").format(_(doctype), name, company))


def _validate_line_data(items: list[dict]) -> None:
	if not items:
		frappe.throw(_("Add at least one item."))
	for row in items:
		if not row.get("item_code") or flt(row.get("qty")) <= 0:
			frappe.throw(_("Every item must have a controlled item and a quantity greater than zero."))


def _item_has_supplier(item_code: str, supplier: str) -> bool:
	return bool(supplier and frappe.db.exists("Item Supplier", {
		"parent": item_code, "parenttype": "Item", "supplier": supplier,
	}))


def _validate_order_supplier_items(data: dict, items: list[dict], existing=None) -> None:
	"""Every Standard/Open order line must be linked to its header supplier."""
	lpo_type = data.get("lpo_type") or "Standard"
	if lpo_type not in {"Standard", "Open"}:
		return
	supplier = data.get("supplier") or ""
	if not supplier:
		frappe.throw(_("Select a supplier before choosing an item."))
	previous_type = (existing.get("custom_lpo_type") or ("Open" if existing.get("custom_is_provisional_po") else "Standard")) if existing else ""
	legacy_items = Counter(row.item_code for row in existing.items) if existing and existing.supplier == supplier and previous_type in {"Standard", "Open"} else Counter()
	for row in items:
		item_code = row["item_code"]
		if legacy_items[item_code]:
			legacy_items[item_code] -= 1
		elif not _item_has_supplier(item_code, supplier):
			frappe.throw(_("Item {0} is not linked to supplier {1}.").format(item_code, supplier))


def _document_company(company: str | None) -> str:
	"""Resolve an explicit company or the active workspace company for a new document."""
	company = company or _active_company()
	if not company:
		frappe.throw(_("Select a company before creating a controlled procurement document."))
	if company not in _allowed_companies():
		frappe.throw(_("You do not have access to company {0}.").format(company), frappe.PermissionError)
	return company


def _allocation_defaults(data: dict, company: str | None = None) -> dict:
	defaults = _normalized_allocation(data)
	_validate_allocation(defaults, company or _document_company(data.get("company")))
	return defaults


def _apply_allocation(row, row_data: dict, defaults: dict, company: str = "") -> None:
	row_allocation = _normalized_allocation(row_data)
	allocation = {field: row_allocation[field] or defaults[field] for field in defaults}
	_validate_allocation(allocation, company)
	for field, value in allocation.items():
		setattr(row, field, value)


def _is_mrg_company(company: str | None) -> bool:
	if not company:
		return False
	abbr = frappe.db.get_value("Company", company, "abbr") or ""
	return "MRG" in company.upper() or abbr.upper() == "MRG"


def _boq_allocation_mode(company: str | None) -> str:
	"""How much of Project / Bill No / BOQ Item a company's procurement documents must carry.

	"required": all three (SKADA, and any company not configured otherwise).
	"project_only": Project is mandatory, Bill No and BOQ Item are optional (MRG).
	"""
	from construction_management.api.project_numbering import is_skada_company

	if is_skada_company(company):
		return "required"
	if _is_mrg_company(company):
		return "project_only"
	return "required"


def _validate_allocation(allocation: dict, company: str | None = None, mode: str | None = None) -> None:
	"""Check Project / Bill No / BOQ Item against the company's rule (or an explicit `mode`)."""
	allocation = _normalized_allocation(allocation)
	project = allocation.get("project")
	bill_no = allocation.get("bill_no")
	boq_item = allocation.get("boq_item")
	if not project:
		frappe.throw(_("Project is required for controlled procurement."))
	if (mode or _boq_allocation_mode(company)) == "required" and (not bill_no or not boq_item):
		frappe.throw(_("Project, Bill No, and BOQ Item are required for controlled procurement."))
	if boq_item and not bill_no:
		frappe.throw(_("Select the Bill No for BOQ Item {0}.").format(boq_item))

	frappe.has_permission("Project", "read", doc=frappe.get_doc("Project", project), throw=True)
	# whatever is filled in must belong together
	if bill_no:
		bill = frappe.db.get_value("BOQ Bill", bill_no, ["project"], as_dict=True)
		if not bill or bill.project != project:
			frappe.throw(_("Bill No {0} does not belong to Project {1}.").format(bill_no, project))
	if boq_item:
		item = frappe.db.get_value("BOQ Item", boq_item, ["project", "parent_bill"], as_dict=True)
		if not item or item.project != project or item.parent_bill != bill_no:
			frappe.throw(_("Bill No and BOQ Item must belong to the selected Project."))


CATALOG_SORT_FIELDS = {
	"item_name": "item.item_name",
	"name": "item.name",
	"item_group": "item.item_group",
	"controlled_item_type": "item.controlled_item_type",
	"controlled_catalog_source": "item.controlled_catalog_source",
	"actual_qty": "COALESCE(stock.actual_qty, 0)",
}
CATALOG_EXPORT_HEADERS = (
	"Item Code", "Description", "Item Group", "Type", "UOM", "Available Stock", "Source",
)


@frappe.whitelist(methods=["GET"])
def get_catalog(
	company: str = "", search: str = "", start: int = 0, page_length: int = 50,
	item_type: str = "", item_group: str = "", source: str = "", stock_status: str = "",
	sort_by: str = "item_name", sort_order: str = "asc", supplier: str = "", stock_only: int = 0,
	unique_items: int = 0,
) -> dict:
	"""Filter and sort the full approved catalog before pagination."""
	_require(CATALOG_ROLES)
	from_clause, where_clause, params, order_clause = _catalog_query(
		company, search, item_type, item_group, source, stock_status, sort_by, sort_order, supplier, stock_only,
	)
	start = max(cint(start), 0)
	page_length = min(max(cint(page_length) or 50, 1), 100)
	if cint(unique_items):
		rows = _unique_catalog_rows(_catalog_rows(from_clause, where_clause, params, order_clause))
		return {"rows": rows[start : start + page_length], "total_count": len(rows)}

	params.update({"start": start, "page_length": page_length})
	rows = _catalog_rows(from_clause, where_clause, params, order_clause, paginated=True)
	total_count = frappe.db.sql(
		f"SELECT COUNT(*) {from_clause} WHERE {where_clause}", params,
	)[0][0]
	return {"rows": rows, "total_count": total_count}


def _catalog_rows(
	from_clause: str, where_clause: str, params: dict, order_clause: str, paginated: bool = False,
) -> list[dict]:
	limit = " LIMIT %(page_length)s OFFSET %(start)s" if paginated else ""
	return frappe.db.sql(
		f"""SELECT item.name, item.item_name, item.item_group, item.stock_uom,
		item.controlled_item_type, item.controlled_catalog_source,
		COALESCE(stock.actual_qty, 0) AS actual_qty
		{from_clause} WHERE {where_clause}
		ORDER BY {order_clause}{limit}""",
		params, as_dict=True,
	)


def _unique_catalog_rows(rows: list[dict]) -> list[dict]:
	"""Keep the stable original item for each exact catalog identity."""
	unique = {}
	for position, row in enumerate(rows):
		identity = tuple(_normalise_name(row[field]) for field in (
			"item_name", "item_group", "stock_uom", "controlled_item_type",
		))
		candidate = (position, row)
		current = unique.get(identity)
		if not current or _catalog_item_preference(candidate[1]) < _catalog_item_preference(current[1]):
			unique[identity] = candidate
	return [row for _, row in sorted(unique.values(), key=lambda entry: entry[0])]


def _catalog_item_preference(row: dict) -> tuple[bool, str]:
	name = row["name"]
	return bool(re.search(r"-\d+$", name)), name


@frappe.whitelist(methods=["POST"])
def download_catalog_export(
	file_format: str = "csv", company: str = "", search: str = "", item_type: str = "",
	item_group: str = "", source: str = "", stock_status: str = "", sort_by: str = "item_name",
	sort_order: str = "asc", supplier: str = "", stock_only: int = 0,
) -> dict:
	"""Create a private CSV or XLSX export using the active catalog filters."""
	_require(CATALOG_ROLES)
	file_format = file_format.lower()
	if file_format not in {"csv", "xlsx"}:
		frappe.throw(_("Select CSV or XLSX export format."))
	from_clause, where_clause, params, order_clause = _catalog_query(
		company, search, item_type, item_group, source, stock_status, sort_by, sort_order, supplier, stock_only,
	)
	rows = frappe.db.sql(
		f"""SELECT item.name, item.item_name, item.item_group, item.stock_uom,
		item.controlled_item_type, item.controlled_catalog_source,
		COALESCE(stock.actual_qty, 0) AS actual_qty
		{from_clause} WHERE {where_clause} ORDER BY {order_clause}""",
		params,
		as_dict=True,
	)
	content = _catalog_export_content(rows, file_format)
	filename = f"controlled_catalog_{today()}.{file_format}"
	file_doc = save_file(filename, content, None, None, is_private=1)
	return {"file_url": file_doc.file_url, "filename": filename, "rows": len(rows)}


def _catalog_query(
	company: str, search: str, item_type: str, item_group: str, source: str, stock_status: str,
	sort_by: str, sort_order: str, supplier: str, stock_only: int,
) -> tuple[str, str, dict, str]:
	company = _document_company(company)
	if sort_by not in CATALOG_SORT_FIELDS or sort_order.lower() not in {"asc", "desc"}:
		frappe.throw(_("Invalid catalog sort option."))
	if stock_status not in {"", "in_stock", "out_of_stock"}:
		frappe.throw(_("Invalid stock filter."))
	if item_type and item_type not in STOCK_TYPES | {ASSET, SERVICE}:
		frappe.throw(_("Invalid catalog item type."))

	conditions = ["item.disabled = 0", "item.is_purchase_item = 1", "item.controlled_procurement_catalog = 1"]
	params = {"company": company}
	if search.strip():
		conditions.append("(item.item_name LIKE %(search)s OR item.name LIKE %(search)s)")
		params["search"] = f"%{search.strip()}%"
	for value, column, key in ((item_type, "controlled_item_type", "item_type"), (item_group, "item_group", "item_group"), (source, "controlled_catalog_source", "source")):
		if value:
			conditions.append(f"item.{column} = %({key})s")
			params[key] = value
	if stock_status or cint(stock_only):
		conditions.append("item.controlled_item_type IN ('Stockable', 'Consumable')")
	if stock_status:
		conditions.append("COALESCE(stock.actual_qty, 0) > 0" if stock_status == "in_stock" else "COALESCE(stock.actual_qty, 0) <= 0")
	if supplier:
		conditions.append("EXISTS (SELECT 1 FROM `tabItem Supplier` item_supplier WHERE item_supplier.parent = item.name AND item_supplier.parenttype = 'Item' AND item_supplier.supplier = %(supplier)s)")
		params["supplier"] = supplier

	from_clause = """FROM `tabItem` item
	LEFT JOIN (
		SELECT bin.item_code, SUM(bin.actual_qty) AS actual_qty
		FROM `tabBin` bin
		JOIN `tabWarehouse` warehouse ON warehouse.name = bin.warehouse
		WHERE warehouse.company = %(company)s AND warehouse.is_group = 0
		GROUP BY bin.item_code
	) stock ON stock.item_code = item.name"""
	return from_clause, " AND ".join(conditions), params, f"{CATALOG_SORT_FIELDS[sort_by]} {sort_order.upper()}, item.name ASC"


def _catalog_export_content(rows: list[dict], file_format: str) -> bytes:
	values = [_catalog_export_row(row) for row in rows]
	if file_format == "csv":
		output = StringIO()
		writer = csv.writer(output)
		writer.writerow(CATALOG_EXPORT_HEADERS)
		writer.writerows(values)
		return output.getvalue().encode("utf-8-sig")

	book = Workbook()
	sheet = book.active
	sheet.title = "Controlled Catalog"
	sheet.append(CATALOG_EXPORT_HEADERS)
	for row in values:
		sheet.append(row)
	for column in sheet.columns:
		sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 48)
	stream = BytesIO()
	book.save(stream)
	return stream.getvalue()


def _catalog_export_row(row: dict) -> tuple:
	return (
		_catalog_export_value(row.name), _catalog_export_value(row.item_name), _catalog_export_value(row.item_group),
		_catalog_export_value(row.controlled_item_type), _catalog_export_value(row.stock_uom), flt(row.actual_qty),
		_catalog_export_value(row.controlled_catalog_source),
	)


def _catalog_export_value(value) -> str:
	value = str(value or "")
	return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


@frappe.whitelist(methods=["GET"])
def get_catalog_supplier_match(item_code: str, supplier: str) -> dict:
	"""Check whether an already selected catalog item belongs to a new PO supplier."""
	_require(PURCHASE_ROLES)
	_catalog_item(item_code)
	return {"eligible": _item_has_supplier(item_code, supplier)}


@frappe.whitelist(methods=["GET"])
def get_catalog_item_suppliers(item_code: str, company: str = "") -> list[dict]:
	"""Return visible, active suppliers linked to an approved buying item."""
	_require(PURCHASE_ROLES)
	_document_company(company)
	_catalog_item(item_code)
	if frappe.db.get_value("Item", item_code, "disabled") or not frappe.db.get_value("Item", item_code, "is_purchase_item"):
		frappe.throw(_("Select an active buying item from the controlled catalog."))
	suppliers = frappe.get_all("Item Supplier", filters={"parent": item_code, "parenttype": "Item"}, pluck="supplier", limit_page_length=0)
	if not suppliers:
		return []
	return frappe.get_list("Supplier", filters={"name": ["in", sorted(set(suppliers))], "disabled": 0},
		fields=["name", "supplier_name"], order_by="supplier_name asc", limit_page_length=0)


@frappe.whitelist(methods=["GET"])
def get_catalog_filter_options() -> dict:
	"""Return selectable groups and sources used by approved catalog items."""
	_require(CATALOG_ROLES)
	filters = {"disabled": 0, "is_purchase_item": 1, "controlled_procurement_catalog": 1}
	rows = frappe.get_all("Item", filters=filters, fields=["item_group", "controlled_catalog_source"], limit_page_length=0)
	return {
		"item_groups": sorted({row.item_group for row in rows if row.item_group}),
		"sources": sorted({row.controlled_catalog_source for row in rows if row.controlled_catalog_source}),
	}


@frappe.whitelist(methods=["GET"])
def get_catalog_item_by_name(item_name: str) -> dict:
	"""Return an active controlled catalog Item that matches the description, if any."""
	_require(CATALOG_ROLES)
	item_name = " ".join(str(item_name or "").split())
	if not item_name:
		return {"matched": False}
	item = frappe.db.get_value(
		"Item",
		{
			"item_name": item_name,
			"disabled": 0,
			"controlled_procurement_catalog": 1,
			"is_purchase_item": 1,
		},
		["name", "item_name", "item_group", "stock_uom", "controlled_item_type"],
		as_dict=True,
	)
	if not item:
		# Case-insensitive match for typed descriptions.
		item = frappe.db.sql(
			"""
			SELECT name, item_name, item_group, stock_uom, controlled_item_type
			FROM `tabItem`
			WHERE LOWER(item_name) = LOWER(%s)
				AND IFNULL(disabled, 0) = 0
				AND controlled_procurement_catalog = 1
				AND is_purchase_item = 1
			ORDER BY name ASC
			LIMIT 1
			""",
			item_name,
			as_dict=True,
		)
		item = item[0] if item else None
	if not item:
		return {"matched": False}
	return {
		"matched": True,
		"item_code": item.name,
		"item_name": item.item_name,
		"item_group": item.item_group,
		"stock_uom": item.stock_uom,
		"item_type": item.controlled_item_type,
	}


def _catalog_stock_totals(item_codes: list[str], company: str) -> dict[str, float]:
	if not item_codes or not company:
		return {}
	rows = frappe.db.sql(
		"""select bin.item_code, sum(bin.actual_qty) as actual_qty
		from `tabBin` bin join `tabWarehouse` warehouse on warehouse.name = bin.warehouse
		where bin.item_code in %(items)s and warehouse.company = %(company)s and warehouse.is_group = 0
		group by bin.item_code""",
		{"items": item_codes, "company": company}, as_dict=True,
	)
	return {row.item_code: flt(row.actual_qty) for row in rows}


def _allowed_companies() -> list[str]:
	"""Companies the user can access, preferring explicit Company User Permissions."""
	permitted_companies = frappe.get_all(
		"User Permission",
		filters={"user": frappe.session.user, "allow": "Company"},
		pluck="for_value",
	)
	if permitted_companies:
		return sorted(set(permitted_companies))
	return frappe.get_list("Company", pluck="name", order_by="name asc", limit_page_length=0)


def _active_company() -> str:
	"""The user's default company, if they can still access it; scopes every workspace list."""
	company = frappe.defaults.get_user_default("Company")
	allowed = _allowed_companies()
	if company in allowed:
		return company
	return allowed[0] if len(allowed) == 1 else ""


def _scoped(filters: dict) -> dict:
	company = _active_company()
	return {**filters, "company": company} if company else filters


@frappe.whitelist(methods=["POST"])
def set_active_company(company: str) -> dict:
	"""Switch the workspace company; stored as the user's default like Desk does."""
	_require(CATALOG_ROLES)
	if company not in _allowed_companies():
		frappe.throw(_("You do not have access to company {0}.").format(company), frappe.PermissionError)
	# Desk stores the session default under the scrubbed key; get_user_default reads it from there
	frappe.defaults.set_user_default("company", company)
	return get_workspace_context()


@frappe.whitelist(methods=["GET"])
def get_workspace_context() -> dict:
	"""Return permissions, defaults, and safe dashboard totals for the Vue workspace."""
	roles = _roles()
	if not roles.intersection(CATALOG_ROLES):
		frappe.throw(_("Not permitted"), frappe.PermissionError)
	company = _active_company()
	return {
		"user": frappe.session.user,
		"full_name": frappe.utils.get_fullname(frappe.session.user),
		"user_image": frappe.db.get_value("User", frappe.session.user, "user_image"),
		"currency": (frappe.get_cached_value("Company", company, "default_currency") if company else None)
		or frappe.defaults.get_global_default("currency"),
		"default_company": company,
		"companies": _allowed_companies(),
		"can_purchase": bool(roles.intersection(PURCHASE_ROLES)),
		"can_transfer": bool(roles.intersection(STOCK_ROLES)),
		"can_manage_catalog": bool(roles.intersection(CATALOG_ROLES)),
		"can_create_catalog_as_accounts": _has_explicit_role("Accounts Manager"),
		"can_approve_catalog": _has_explicit_role("Accounts Manager") or _catalog_ceo(),
		"can_approve_catalog_accounts": _has_explicit_role("Accounts Manager"),
		"can_approve_catalog_ceo": _catalog_ceo(),
		"can_disable_catalog": _can_request_catalog_disable(),
		"can_request_price_change": bool(roles.intersection(CATALOG_ROLES)),
		"can_approve_price_change": _explicit_ceo(),
		"can_approve_po_price_accounts": any(frappe.db.exists("Has Role", {
			"parent": frappe.session.user, "parenttype": "User", "role": role,
		}) for role in ACCOUNT_ROLES),
		"can_opening_stock": bool(roles.intersection(OPENING_STOCK_ROLES)),
		"can_approve_lpo": bool(roles.intersection(ACCOUNT_ROLES | ADMIN_ROLES)),
		"can_approve_lpo_accounts": bool(roles.intersection(ACCOUNT_ROLES | {"System Manager"})),
		"can_approve_lpo_ceo": bool(roles.intersection(ADMIN_ROLES)),
		"can_override_lpo": "System Manager" in roles,
		"can_view_insights": bool(roles.intersection(CATALOG_ROLES)),
		"can_view_stock_ledger": bool(roles.intersection(ACCOUNT_ROLES | ADMIN_ROLES)),
		"boq_allocation_mode": _boq_allocation_mode(company),
		"can_email_purchase_order": frappe.has_permission("Purchase Order", "email"),
		"can_create_pdc": frappe.has_permission("Post Dated Cheques", "create") and frappe.has_permission("Post Dated Cheques", "submit"),
		"dashboard": {
			"purchase_orders": frappe.db.count("Purchase Order", _scoped({"is_subcontracted": 0, "docstatus": ["<", 2]})),
			"purchase_receipts": frappe.db.count("Purchase Receipt", _scoped({"is_subcontracted": 0, "docstatus": ["<", 2]})),
			"material_transfers": frappe.db.count("Stock Entry", _scoped({"purpose": "Material Transfer", "docstatus": ["<", 2]})),
			"catalog_items": frappe.db.count("Item", {"controlled_procurement_catalog": 1, "disabled": 0}),
			"draft_purchase_orders": frappe.db.count("Purchase Order", _scoped({"is_subcontracted": 0, "docstatus": 0})),
			"purchase_invoices": frappe.db.count("Purchase Invoice", _scoped(_workspace_filters("Purchase Invoice", {"docstatus": ["<", 2]}))),
			"unpaid_invoices": frappe.db.count("Purchase Invoice", _scoped(_workspace_filters("Purchase Invoice", {"docstatus": 1, "outstanding_amount": [">", 0]}))),
			"payable_outstanding": sum(frappe.get_all(
				"Purchase Invoice", filters=_scoped(_workspace_filters("Purchase Invoice", {"docstatus": 1, "outstanding_amount": [">", 0]})),
				pluck="outstanding_amount", limit_page_length=0,
			)),
			"open_lpos": 0 if not _has_open_po_field() else frappe.db.count("Purchase Order", _scoped({
				"is_subcontracted": 0, "docstatus": 1, "custom_is_provisional_po": 1,
				"status": ["not in", ("Closed", "On Hold")],
			})),
		},
	}


DOCUMENT_SORT_FIELDS = {
	"Purchase Order": {"modified": "modified", "name": "name", "supplier": "supplier_name", "project": "project", "date": "transaction_date", "status": "status", "value": "grand_total", "received": "per_received"},
	"Purchase Receipt": {"modified": "modified", "name": "name", "supplier": "supplier_name", "project": "project", "date": "posting_date", "status": "status", "value": "grand_total"},
	"Purchase Invoice": {"modified": "modified", "name": "name", "supplier": "supplier_name", "project": "project", "date": "posting_date", "status": "status", "value": "grand_total"},
	"Stock Entry": {"modified": "modified", "name": "name", "project": "project", "date": "posting_date", "status": "status", "value": "total_outgoing_value"},
}


@frappe.whitelist(methods=["GET"])
def get_controlled_documents(doctype: str, search: str = "", start: int = 0, page_length: int = 25, submitted_only: int = 0, docstatus: str = "", supplier: str = "", project: str = "", lpo_type: str = "", purchase_order: str = "", warehouse: str = "", from_warehouse: str = "", to_warehouse: str = "", item_code: str = "", date_from: str = "", date_to: str = "", sort_by: str = "modified", sort_order: str = "desc") -> dict:
	"""List documents that belong to the controlled workspace.

	`docstatus` also accepts "unpaid" and "overdue" for purchase invoices.
	"""
	if doctype not in DOCUMENT_ROLES:
		frappe.throw(_("Unsupported controlled document type."))
	_require(DOCUMENT_ROLES[doctype])
	if sort_by not in DOCUMENT_SORT_FIELDS[doctype] or sort_order.lower() not in {"asc", "desc"}:
		frappe.throw(_("Invalid document sort option."))
	company = _document_company(None)
	filters = _workspace_filters(doctype, {"company": company, "docstatus": ["<", 3 if doctype in {"Purchase Order", "Purchase Receipt", "Stock Entry"} else 2]})
	if doctype == "Stock Entry":
		filters["purpose"] = "Material Transfer"
	if cint(submitted_only):
		filters["docstatus"] = 1
	elif docstatus in ({"0", "1", "2"} if doctype in {"Purchase Order", "Purchase Receipt", "Stock Entry"} else {"0", "1"}):
		filters["docstatus"] = cint(docstatus)
	elif doctype == "Purchase Invoice" and docstatus in {"unpaid", "overdue"}:
		filters.update({"docstatus": 1, "outstanding_amount": [">", 0]})
		if docstatus == "overdue":
			filters["due_date"] = ["<", frappe.utils.today()]
	if supplier and doctype in {"Purchase Order", "Purchase Receipt"}:
		filters["supplier"] = supplier
	matching_names = None

	def narrow_to(names):
		nonlocal matching_names
		names = set(names)
		matching_names = names if matching_names is None else matching_names & names

	if project:
		if doctype in {"Purchase Order", "Purchase Receipt"}:
			header_names = frappe.get_all(doctype, filters={"company": company, "project": project}, pluck="name", limit_page_length=0)
			line_names = frappe.get_all(f"{doctype} Item", filters={"project": project}, pluck="parent", limit_page_length=0)
			narrow_to(header_names + line_names)
		else:
			filters["project"] = project
	if lpo_type and doctype == "Purchase Order":
		if lpo_type not in {"Standard", "Open", "Manual"}:
			frappe.throw(_("Invalid LPO type."))
		if lpo_type in {"Standard", "Open"}:
			narrow_to(row.name for row in frappe.get_all("Purchase Order", filters={"company": company},
				fields=["name", "custom_lpo_type", "custom_is_provisional_po"], limit_page_length=0)
				if row.custom_lpo_type == lpo_type or (not row.custom_lpo_type and ("Open" if cint(row.custom_is_provisional_po) else "Standard") == lpo_type))
		else:
			filters["custom_lpo_type"] = lpo_type
	if date_from or date_to:
		date_field = "transaction_date" if doctype == "Purchase Order" else "posting_date"
		if date_from and date_to and getdate(date_from) > getdate(date_to):
			frappe.throw(_("From date must be before To date."))
		filters[date_field] = ["between", [str(getdate(date_from or "1900-01-01")), str(getdate(date_to or "2999-12-31"))]]
	if doctype == "Purchase Receipt" and purchase_order:
		narrow_to(frappe.get_all("Purchase Receipt Item", filters={"purchase_order": purchase_order}, pluck="parent", limit_page_length=0))
	if doctype == "Purchase Receipt" and warehouse:
		narrow_to(frappe.get_all("Purchase Receipt Item", filters={"warehouse": warehouse}, pluck="parent", limit_page_length=0))
	if doctype == "Stock Entry":
		for field, value in (("from_warehouse", from_warehouse), ("to_warehouse", to_warehouse)):
			if value:
				filters[field] = value
		if item_code:
			narrow_to(frappe.get_all("Stock Entry Detail", filters={"item_code": item_code}, pluck="parent", limit_page_length=0))
	if matching_names is not None:
		filters["name"] = ["in", sorted(matching_names)]
	fields = _document_list_fields(doctype)
	or_filters = None
	if search:
		like = f"%{search.strip()}%"
		or_filters = [["name", "like", like], ["project", "like", like]]
		if doctype != "Stock Entry":
			or_filters.append(["supplier_name", "like", like])
		if doctype == "Purchase Invoice":
			or_filters.append(["custom_supplier_invoice_no", "like", like])
		if doctype == "Purchase Receipt":
			or_filters.append(["supplier_delivery_note", "like", like])
		if doctype == "Purchase Order":
			or_filters.append(["custom_lpo_reference", "like", like])
	start = max(cint(start), 0)
	page_length = min(max(cint(page_length) or 25, 1), 100)
	order_by = f"{DOCUMENT_SORT_FIELDS[doctype][sort_by]} {sort_order.lower()}, name asc"
	rows = frappe.get_list(doctype, filters=filters, or_filters=or_filters, fields=fields, order_by=order_by, limit_start=start, limit_page_length=page_length)
	if doctype == "Purchase Order":
		_annotate_po_list_catalog_flags(rows)
	return {
		"rows": rows,
		"total_count": len(frappe.get_list(doctype, filters=filters, or_filters=or_filters, pluck="name", limit_page_length=0)),
	}


@frappe.whitelist(methods=["GET"])
def get_controlled_document(doctype: str, name: str) -> dict:
	"""Return a controlled document and its standard email trail."""
	doc = _controlled_doc(doctype, name)
	result = doc.as_dict()
	for row in result.get("items") or []:
		row["vat"] = _vat_choice(row.get("item_tax_template"), doc.company)
	result["email_history"] = frappe.get_all(
		"Communication",
		filters={"reference_doctype": doctype, "reference_name": name, "communication_medium": "Email"},
		fields=["name", "subject", "sender", "recipients", "sent_or_received", "creation"],
		order_by="creation desc",
	)
	if doctype == "Purchase Order":
		result["supplier_email"] = frappe.db.get_value("Supplier", doc.supplier, "email_id") or ""
		from construction_management.api.lpo_workflow import lpo_type
		result["custom_lpo_type"] = lpo_type(doc)
		result["approval_history"] = json.loads(doc.get("custom_lpo_approval_history") or "[]")
	if doctype in {"Purchase Order", "Purchase Receipt"}:
		result["can_delete"] = _can_delete_controlled_document(doc)
	if doctype in {"Purchase Order", "Purchase Receipt", "Purchase Invoice", "Stock Entry"}:
		result["can_cancel"] = bool(frappe.has_permission(doctype, "cancel", doc=doc))
		result["can_amend"] = bool(frappe.has_permission(doctype, "amend", doc=doc))
		result["can_email"] = bool(frappe.has_permission(doctype, "email", doc=doc))
	if doctype == "Stock Entry":
		result["source_project"] = get_warehouse_project(doc.from_warehouse) if doc.from_warehouse else ""
		result["target_project"] = get_warehouse_project(doc.to_warehouse) if doc.to_warehouse else ""
	if doc.get("amended_from"):
		result["previous_revision"] = doc.amended_from
	result["next_revision"] = frappe.db.get_value(doctype, {"amended_from": name}, "name")
	result["attachments"] = frappe.get_all(
		"File", filters={"attached_to_doctype": doctype, "attached_to_name": name},
		fields=["name", "file_name", "file_url", "creation"], order_by="creation desc",
	)
	result["owner_name"] = frappe.utils.get_fullname(doc.owner)
	return result


TAX_MASTER = "Purchase Taxes and Charges Template"
LINE_VAT = ("standard", "zero", "exempt")


def _line_vat_templates(company: str) -> dict:
	"""Item Tax Templates behind the per-line VAT choices, resolved by the company's naming."""
	templates = frappe.get_all("Item Tax Template", filters={"company": company, "disabled": 0}, pluck="name")

	def find(word):
		return next((name for name in templates if "vat" in name.lower() and word in name.lower()), None)

	return {"standard": None, "zero": find("zero"), "exempt": find("exempt")}


def _vat_choice(item_tax_template: str | None, company: str) -> str:
	"""Map a stored Item Tax Template back to the line choice shown in the UI."""
	if not item_tax_template:
		return "standard"
	for choice, template in _line_vat_templates(company).items():
		if template and template == item_tax_template:
			return choice
	return "standard"


@frappe.whitelist(methods=["GET"])
def get_tax_options(company: str = "") -> dict:
	"""Purchase tax templates and per-line VAT choices for the active company."""
	_require(PURCHASE_ROLES | STOCK_ROLES)
	company = company or _active_company()
	return {
		"templates": frappe.get_all(TAX_MASTER, filters={"company": company, "disabled": 0}, pluck="name", order_by="is_default desc, name asc"),
		"default_template": frappe.db.get_value(TAX_MASTER, {"company": company, "is_default": 1, "disabled": 0}, "name"),
		"line_templates": _line_vat_templates(company),
	}


def _erpnext_get_taxes_and_charges():
	"""ERPNext moved this helper; try both locations for bench compatibility."""
	try:
		from erpnext.controllers.accounts_controller import get_taxes_and_charges
		return get_taxes_and_charges
	except ImportError:
		from erpnext.accounts.services.taxes import get_taxes_and_charges
		return get_taxes_and_charges


def _erpnext_update_child_qty_rate():
	"""ERPNext moved this helper; try both locations for bench compatibility."""
	try:
		from erpnext.controllers.accounts_controller import update_child_qty_rate
		return update_child_qty_rate
	except ImportError:
		from erpnext.accounts.services.child_item_update import update_child_qty_rate
		return update_child_qty_rate


def _apply_taxes(doc, taxes_and_charges: str | None, vat_by_row: list[str] | None = None) -> None:
	"""Load the header tax template and tag Zero-rated / Exempt lines with their Item Tax Template.

	ERPNext rebuilds each line's tax map from these on save, so the template is all a line needs.
	"""
	get_taxes_and_charges = _erpnext_get_taxes_and_charges()

	doc.set("taxes", [])
	doc.taxes_and_charges = taxes_and_charges or None
	if taxes_and_charges:
		if frappe.db.get_value(TAX_MASTER, taxes_and_charges, "company") != doc.company:
			frappe.throw(_("Tax template {0} does not belong to company {1}.").format(taxes_and_charges, doc.company))
		for tax in get_taxes_and_charges(TAX_MASTER, taxes_and_charges) or []:
			doc.append("taxes", tax)
	line_templates = _line_vat_templates(doc.company)
	for row, choice in zip(doc.items, vat_by_row or []):
		choice = choice if choice in LINE_VAT else "standard"
		if choice != "standard" and not line_templates.get(choice):
			frappe.throw(_("No {0} VAT Item Tax Template is set up for company {1}.").format(choice, doc.company))
		row.item_tax_template = line_templates.get(choice)


def _totals_payload(doc) -> dict:
	"""Per-line VAT and tax rows as ERPNext calculated them."""
	# calculate_taxes_and_totals keeps an unsaved per-item breakdown on the document
	position = {id(row): i for i, row in enumerate(doc.items)}
	line_tax = [0.0] * len(doc.items)
	for detail in doc.get("_item_wise_tax_details") or []:
		i = position.get(id(detail.item))
		if i is not None:
			line_tax[i] += flt(detail.amount)
	return {
		"net_total": flt(doc.net_total),
		"total_taxes_and_charges": flt(doc.total_taxes_and_charges),
		"grand_total": flt(doc.grand_total),
		"rounded_total": flt(doc.get("rounded_total")),
		"taxes": [{"description": tax.description or tax.account_head, "rate": flt(tax.rate), "amount": flt(tax.tax_amount)} for tax in doc.taxes],
		"line_tax": line_tax,
	}


@frappe.whitelist(methods=["POST"])
def preview_totals(doctype: str, data: str | dict) -> dict:
	"""Calculate taxes and totals for an unsaved form exactly as ERPNext will on save."""
	_require(PURCHASE_ROLES)
	if doctype not in ("Purchase Order", "Purchase Receipt", "Purchase Invoice"):
		frappe.throw(_("Unsupported document type."))
	data = _as_dict(data)
	doc = frappe.new_doc(doctype)
	doc.company = data.get("company") or _active_company()
	doc.supplier = data.get("supplier") or None
	doc.currency = frappe.get_cached_value("Company", doc.company, "default_currency")
	doc.conversion_rate = 1
	lines = [row for row in data.get("items") or [] if row.get("item_code")]
	for i, row in enumerate(lines, start=1):
		qty, rate = flt(row.get("qty")), flt(row.get("rate"))
		uom = row.get("uom") or row.get("stock_uom") or frappe.db.get_value("Item", row.get("item_code"), "stock_uom") or "Nos"
		doc.append("items", {"idx": i, "name": f"row-{i}", "item_code": row.get("item_code"), "qty": qty, "rate": rate,
			"price_list_rate": rate, "conversion_factor": 1, "uom": uom})
	_apply_taxes(doc, data.get("taxes_and_charges"), [row.get("vat") for row in lines])
	from erpnext.controllers.taxes_and_totals import calculate_taxes_and_totals

	calculate_taxes_and_totals(doc)
	payload = _totals_payload(doc)
	# map line tax back onto the submitted rows, including blank ones
	it = iter(payload["line_tax"])
	payload["line_tax"] = [next(it) if row.get("item_code") else 0 for row in data.get("items") or []]
	return payload


CONNECTION_DATE_FIELDS = ("posting_date", "transaction_date", "reference_date", "schedule_date")
CONNECTION_AMOUNT_FIELDS = ("grand_total", "paid_amount", "total_debit", "amount", "total")


def _connection_names(doctype: str, name: str, target: str, found: dict, links: dict) -> list[str]:
	"""Names of `target` documents linked to this one, following Desk's Connections rules."""
	from frappe.desk.notifications import get_child_doctypes_with_field, get_external_link_fieldname

	if found.get("names"):
		return list(found["names"])
	fieldname = get_external_link_fieldname(target, links)
	if not fieldname:
		return []
	child_doctypes = get_child_doctypes_with_field(target, fieldname)
	if child_doctypes:
		or_filters = [[child, fieldname, "=", name] for child in child_doctypes]
		return frappe.get_list(target, or_filters=or_filters, pluck="name", distinct=True, limit_page_length=50, order_by="creation desc")
	return frappe.get_list(target, filters={fieldname: name}, pluck="name", limit_page_length=50, order_by="creation desc")


def _connection_rows(doctype: str, names: list[str], company: str = "") -> list[dict]:
	"""Status, date and amount for each linked document (whatever fields that doctype has)."""
	if not names:
		return []
	meta = frappe.get_meta(doctype)
	fields = ["name", "docstatus"] + [field for field in ("status", "title") if meta.has_field(field)]
	date_field = next((field for field in CONNECTION_DATE_FIELDS if meta.has_field(field)), None)
	amount_field = next((field for field in CONNECTION_AMOUNT_FIELDS if meta.has_field(field)), None)
	fields += [f"{date_field} as date"] if date_field else []
	fields += [f"{amount_field} as amount"] if amount_field else []
	if doctype == "Purchase Order" and meta.has_field("project"):
		fields.append("project")
	if not meta.is_submittable:
		fields.remove("docstatus")
	filters = {"name": ["in", names]}
	if company and meta.has_field("company"):
		filters["company"] = company
	rows = frappe.get_list(doctype, filters=filters, fields=fields, order_by="creation desc", limit_page_length=50)
	if doctype == "Purchase Order":
		projects = {row.project for row in rows if row.get("project")}
		labels = {row.name: row.project_name for row in frappe.get_all("Project", filters={"name": ["in", list(projects)]}, fields=["name", "project_name"])} if projects else {}
		for row in rows:
			if row.get("project"):
				row.title = f"{row.project} - {labels.get(row.project) or row.project}"
	return rows


def _receipt_direct_connections(doc) -> list[tuple[str, str, list[str]]]:
	"""Find source and downstream documents even when Desk's dashboard omits them."""
	po_names = {row.purchase_order for row in doc.items if row.purchase_order}
	if doc.get("custom_purchase_order"):
		po_names.add(doc.custom_purchase_order)
	invoice_names = frappe.get_all(
		"Purchase Invoice Item", filters={"purchase_receipt": doc.name}, pluck="parent", distinct=True, limit_page_length=50,
	)
	return_names = frappe.get_all(
		"Purchase Receipt", filters={"return_against": doc.name, "company": doc.company}, pluck="name", limit_page_length=50,
	)
	return [
		(_("Source"), "Purchase Order", sorted(po_names)),
		(_("Related"), "Purchase Invoice", invoice_names),
		(_("Returns"), "Purchase Receipt", return_names),
	]


@frappe.whitelist(methods=["GET"])
def get_document_connections(doctype: str, name: str) -> list[dict]:
	"""Desk's Connections tab for a workspace document: groups of linked doctypes with their documents."""
	from frappe.desk.notifications import get_open_count

	doc = _controlled_doc(doctype, name)
	links = doc.meta.get_dashboard_data()
	counts = get_open_count(doctype, name).get("count") or {}
	found_by_doctype = {
		row["doctype"]: row
		for row in (counts.get("internal_links_found") or []) + (counts.get("external_links_found") or [])
		if row.get("count")
	}
	groups = []
	seen = defaultdict(set)
	if doctype == "Purchase Receipt":
		for label, target, names in _receipt_direct_connections(doc):
			if not names or not frappe.has_permission(target, "read"):
				continue
			rows = _connection_rows(target, names, doc.company)
			if rows:
				groups.append({"label": label, "items": [{"doctype": target, "count": len(rows), "documents": rows}]})
				seen[target].update(row.name for row in rows)
	for group in links.transactions:
		items = []
		for target in group.get("items") or []:
			found = found_by_doctype.get(target)
			if not found or not frappe.has_permission(target, "read"):
				continue
			names = [n for n in _connection_names(doctype, name, target, found, links) if n not in seen[target]]
			rows = _connection_rows(target, names, doc.company)
			if rows:
				items.append({"doctype": target, "count": len(rows), "documents": rows})
				seen[target].update(row.name for row in rows)
		if items:
			label = group.get("label") or _("Related")
			existing_group = next((entry for entry in groups if entry["label"] == label), None)
			if existing_group:
				existing_group["items"].extend(items)
			else:
				groups.append({"label": label, "items": items})

	# post-dated cheques reserve invoice amounts but are not in ERPNext's dashboard config
	if doctype == "Purchase Invoice" and frappe.db.table_exists("PDC Invoice Reference") and frappe.has_permission("Post Dated Cheques", "read"):
		pdcs = frappe.get_all("PDC Invoice Reference", filters={"reference_doctype": doctype, "reference_name": name}, pluck="parent", distinct=True)
		if pdcs:
			item = {"doctype": "Post Dated Cheques", "count": len(pdcs), "documents": _connection_rows("Post Dated Cheques", pdcs, doc.company)}
			payment = next((group for group in groups if group["label"] == _("Payment")), None)
			if payment:
				payment["items"].append(item)
			else:
				groups.insert(0, {"label": _("Payment"), "items": [item]})
	if doctype == "Purchase Order" and frappe.db.has_column("Post Dated Cheques", "custom_purchase_order") and frappe.has_permission("Post Dated Cheques", "read"):
		pdcs = frappe.get_all("Post Dated Cheques", filters={"custom_purchase_order": name}, pluck="name")
		rows = _connection_rows("Post Dated Cheques", pdcs, doc.company)
		if rows:
			payment = next((group for group in groups if group["label"] == _("Payment")), None)
			if not payment:
				payment = {"label": _("Payment"), "items": []}
				groups.append(payment)
			payment["items"].append({"doctype": "Post Dated Cheques", "count": len(rows), "documents": rows})
	return groups


@frappe.whitelist(methods=["GET"])
def get_boq_options(project: str, bill_no: str = "") -> dict:
	"""Return only BOQ bills and items belonging to one selected project."""
	_require(PURCHASE_ROLES | STOCK_ROLES)
	if not project:
		return {"bills": [], "items": []}
	frappe.has_permission("Project", "read", doc=frappe.get_doc("Project", project), throw=True)
	bills = frappe.get_all("BOQ Bill", filters={"project": project}, fields=["name", "bill_no"], order_by="sequence asc")
	item_filters = {"project": project}
	if bill_no:
		item_filters["parent_bill"] = bill_no
	items = frappe.get_all("BOQ Item", filters=item_filters, fields=["name", "description", "parent_bill", "unit"], order_by="idx asc")
	return {"bills": bills, "items": items}


@frappe.whitelist(methods=["GET"])
def get_procurement_projects(search: str = "", selected: str = "", company: str = "") -> list[dict]:
	"""Permission-aware Project choices with both the stored ID and display name."""
	_require(PURCHASE_ROLES | STOCK_ROLES)
	filters = {"company": _document_company(company)}
	if selected:
		filters["name"] = selected
	return frappe.get_list("Project", filters=filters,
		or_filters=[["name", "like", f"%{search.strip()}%"], ["project_name", "like", f"%{search.strip()}%"]] if search.strip() and not selected else None,
		fields=["name", "project_name"], order_by="name asc", limit_page_length=1 if selected else 20)


@frappe.whitelist(methods=["GET"])
def get_purchase_order_items(purchase_order: str) -> list[dict]:
	"""Return receivable PO lines; Open LPO lines remain available after full receipt."""
	_require(PURCHASE_ROLES)
	po = _controlled_doc("Purchase Order", purchase_order)
	if po.docstatus != 1:
		frappe.throw(_("Select a submitted Purchase Order."))
	open_lpo = _is_open_lpo(po)
	_validate_receivable_po(po, open_lpo)
	return [
		{
			"purchase_order_item": row.name,
			"item_code": row.item_code,
			"item_name": row.item_name,
			"remaining_qty": max(flt(row.qty) - flt(row.received_qty), 0),
			"is_open_lpo": open_lpo,
			"rate": row.rate,
			"project": row.project,
			"bill_no": row.bill_no,
			"boq_item": row.boq_item,
			"controlled_item_type": row.controlled_item_type,
			"vat": _vat_choice(row.item_tax_template, po.company),
			"po_taxes_and_charges": po.taxes_and_charges or "",
			"uom": row.uom,
			"warehouse": row.warehouse,
			# app-created POs need catalog items + BOQ allocation; Desk POs are received as they are
			"controlled": cint(po.controlled_procurement),
		}
		for row in po.items
		if not cint(row.get("closed")) and not cint(row.get("delivered_by_supplier"))
		and (open_lpo or flt(row.qty) > flt(row.received_qty))
	]


def _receipt_items_by_po_line(po, items: list[dict], open_lpo: bool) -> dict[str, list[dict]]:
	"""Validate source identity and the combined quantity of split receipt rows."""
	po_rows = {row.name: row for row in po.items}
	grouped = defaultdict(list)
	for line in items:
		po_row = po_rows.get(line.get("purchase_order_item"))
		if not po_row or po_row.item_code != line.get("item_code") or cint(po_row.get("closed")) or cint(po_row.get("delivered_by_supplier")):
			frappe.throw(_("Receipt items must come from the selected Purchase Order."))
		if flt(line.get("qty")) <= 0:
			frappe.throw(_("Every receipt row needs a positive quantity."))
		grouped[po_row.name].append(line)
	for row_name, lines in grouped.items():
		po_row = po_rows[row_name]
		pending = max(flt(po_row.qty) - flt(po_row.received_qty), 0)
		if not open_lpo and sum(flt(line.get("qty")) for line in lines) > pending + 0.0001:
			frappe.throw(_("{0}: only {1} {2} is left to receive across all rows.").format(
				po_row.item_name or po_row.item_code, pending, po_row.uom))
	return grouped


@frappe.whitelist(methods=["POST"])
def create_controlled_item(data: str | dict) -> dict:
	"""Create a manual catalog request for the recorded approval route."""
	data = _as_dict(data)
	return create_catalog_request({"company": data.get("company"), "source": "Manual",
		"request_origin": data.get("request_origin"), "items": [data]})


@frappe.whitelist(methods=["POST"])
def create_purchase_order(data: str | dict) -> dict:
	_require(PURCHASE_ROLES)
	data = _as_dict(data)
	company = _document_company(data.get("company"))
	raw_request_id = data.get("creation_request_id") or ""
	if not isinstance(raw_request_id, str):
		frappe.throw(_("Invalid Purchase Order creation request ID."))
	request_id = raw_request_id.strip()
	if request_id:
		try:
			if str(UUID(request_id)) != request_id:
				raise ValueError
		except (ValueError, TypeError):
			frappe.throw(_("Invalid Purchase Order creation request ID."))
		existing = _purchase_order_for_request(request_id, company)
		if existing:
			return existing
	data.update(_normalized_allocation(data))
	data.update(_order_discount(data))
	data["company"] = company
	items = data.get("items") or []
	for row in items:
		row.update(_normalized_allocation(row))
	_validate_line_data(items)
	_validate_order_supplier_items(data, items)
	from construction_management.api import po_price_approvals
	price_changes = po_price_approvals.prepare_rates(company, data.get("supplier"), data.get("order_date") or today(), items)
	defaults = _allocation_defaults(data, company)
	_validate_company_links(company, project=data.get("project"), warehouse=data.get("warehouse"))
	for row in items:
		_validate_company_links(company, project=row.get("project"), warehouse=row.get("warehouse"))
	doc = frappe.new_doc("Purchase Order")
	doc.update({
		"company": company, "supplier": data.get("supplier"), "project": data.get("project"),
		"bill_no": defaults["bill_no"], "boq_item": defaults["boq_item"],
		"schedule_date": data.get("required_by"), "set_warehouse": data.get("warehouse"),
		"contact_person": data.get("contact_person") or None,
		"payment_terms_template": _payment_terms_template(data.get("payment_terms_template")),
		"tc_name": data.get("tc_name") or None, "terms": data.get("terms") or None,
		"controlled_procurement": 1,
		"transaction_date": data.get("order_date") or today(),
		"custom_lpo_type": data.get("lpo_type") or "Standard",
		"custom_creation_request_id": request_id or None,
		"apply_discount_on": data.get("apply_discount_on") or "Net Total",
		"additional_discount_percentage": flt(data.get("additional_discount_percentage")),
		"discount_amount": flt(data.get("discount_amount")),
	})
	for row in items:
		_append_item(doc, row, defaults)
	_apply_taxes(doc, data.get("taxes_and_charges"), [row.get("vat") for row in items])
	doc.flags.controlled_procurement_api = True
	if request_id:
		frappe.db.savepoint("controlled_po_creation")
	message_count = len(getattr(frappe.local, "message_log", []))
	try:
		doc.insert()
		po_price_approvals.stage(doc, price_changes)
	except frappe.UniqueValidationError:
		if not request_id:
			raise
		frappe.db.rollback(save_point="controlled_po_creation")
		existing = _purchase_order_for_request(request_id, company)
		if not existing:
			raise
		frappe.local.message_log = frappe.local.message_log[:message_count]
		return existing
	return {"doctype": doc.doctype, "name": doc.name, "price_approval_status": frappe.db.get_value("Purchase Order", doc.name, "custom_po_price_status")}


def _purchase_order_for_request(request_id: str, company: str) -> dict | None:
	rows = frappe.db.sql(
		"""select name, company, owner, controlled_procurement, is_subcontracted
		from `tabPurchase Order` where custom_creation_request_id = %s for update""",
		request_id, as_dict=True,
	)
	if not rows:
		return None
	doc = rows[0]
	if doc.company != company or doc.owner != frappe.session.user or not cint(doc.controlled_procurement) or cint(doc.is_subcontracted):
		frappe.throw(_("This Purchase Order creation request belongs to another user or company."), frappe.PermissionError)
	return {"doctype": "Purchase Order", "name": doc.name}


@frappe.whitelist(methods=["POST"])
def update_purchase_order(name: str, data: str | dict) -> dict:
	"""Replace editable draft values using the same catalog and allocation rules as creation."""
	doc = _controlled_doc("Purchase Order", name, "write")
	if doc.docstatus != 0 or not cint(doc.controlled_procurement):
		frappe.throw(_("Only a controlled draft Purchase Order can be edited here."))
	if doc.get("custom_lpo_approval_status") in {"Pending Accounts", "Pending CEO"}:
		frappe.throw(_("This Manual LPO is awaiting approval."))
	if doc.get("custom_lpo_approval_status") == "Needs Correction" and doc.owner != frappe.session.user and "System Manager" not in _roles():
		frappe.throw(_("Only the creator can correct this Manual LPO."), frappe.PermissionError)
	data = _as_dict(data)
	data.update(_order_discount(data))
	if data.get("company") and data["company"] != doc.company:
		frappe.throw(_("Switch company in the workspace; the Purchase Order company cannot be changed."))
	items = data.get("items") or []
	for row in items:
		row.update(_normalized_allocation(row))
	_validate_line_data(items)
	existing_rows = {row.name: row for row in doc.items}
	kept_names = [row.get("docname") for row in items if row.get("docname")]
	if len(kept_names) != len(set(kept_names)) or any(name not in existing_rows for name in kept_names):
		frappe.throw(_("Invalid or repeated Purchase Order item row."))
	if any(row.get("docname") and existing_rows[row["docname"]].item_code != row.get("item_code") for row in items):
		frappe.throw(_("A Purchase Order item row cannot be reassigned to another catalog item."))
	_validate_order_supplier_items(data, items, doc)
	from construction_management.api import po_price_approvals
	if doc.get("custom_po_price_status") in {"Pending Accounts", "Pending CEO"}:
		po_price_approvals.supersede(doc, "Superseded by a newer Purchase Order edit.")
		doc.reload()
	price_changes = po_price_approvals.prepare_rates(doc.company, data.get("supplier"), data.get("order_date") or doc.transaction_date, items, doc)
	defaults = _allocation_defaults(data, doc.company)
	_validate_company_links(doc.company, project=data.get("project"), warehouse=data.get("warehouse"))
	for row in items:
		_validate_company_links(doc.company, project=row.get("project"), warehouse=row.get("warehouse"))
	doc.update({
		"supplier": data.get("supplier"), "transaction_date": data.get("order_date") or doc.transaction_date,
		"project": defaults["project"], "bill_no": defaults["bill_no"], "boq_item": defaults["boq_item"],
		"schedule_date": data.get("required_by"), "set_warehouse": data.get("warehouse"),
		"contact_person": data.get("contact_person") or None,
		"payment_terms_template": _payment_terms_template(data.get("payment_terms_template", doc.payment_terms_template)),
		"tc_name": data.get("tc_name") or None,
		"terms": data.get("terms") or None, "custom_lpo_type": data.get("lpo_type") or "Standard",
		"apply_discount_on": data.get("apply_discount_on") or "Net Total",
		"additional_discount_percentage": flt(data.get("additional_discount_percentage")),
		"discount_amount": flt(data.get("discount_amount")),
	})
	doc.set("items", [])
	for row in items:
		_append_item(doc, row, defaults)
		if row.get("docname"):
			doc.items[-1].name = row["docname"]
	doc.set("payment_schedule", [])
	_apply_taxes(doc, data.get("taxes_and_charges"), [row.get("vat") for row in items])
	doc.flags.controlled_procurement_api = True
	doc.save()
	po_price_approvals.stage(doc, price_changes)
	return {"doctype": doc.doctype, "name": doc.name, "price_approval_status": frappe.db.get_value("Purchase Order", doc.name, "custom_po_price_status")}


@frappe.whitelist(methods=["GET"])
def get_lpo_approvals(status: str = "") -> dict:
	_require(ACCOUNT_ROLES | ADMIN_ROLES)
	company = _document_company(None)
	statuses = ["Pending Accounts", "Pending CEO", "Needs Correction"]
	if status and status in statuses:
		statuses = [status]
	rows = frappe.get_all("Purchase Order", filters={
		"company": company, "controlled_procurement": 1, "custom_lpo_type": "Manual",
		"docstatus": 0, "custom_lpo_approval_status": ["in", statuses],
	}, fields=["name", "supplier_name", "project", "grand_total", "currency", "owner", "modified", "custom_lpo_reference", "custom_lpo_approval_status"], order_by="modified desc", limit_page_length=100)
	return {"rows": rows}


@frappe.whitelist(methods=["POST"])
def request_lpo_approval(name: str) -> dict:
	from construction_management.api.lpo_workflow import request_approval
	return request_approval(name)


@frappe.whitelist(methods=["POST"])
def approve_lpo(name: str) -> dict:
	from construction_management.api.lpo_workflow import approve
	return approve(name)


@frappe.whitelist(methods=["POST"])
def reject_lpo(name: str, comment: str) -> dict:
	from construction_management.api.lpo_workflow import reject
	return reject(name, comment)


@frappe.whitelist(methods=["POST"])
def bulk_approve_lpos(names: str | list | None = None) -> dict:
	_require(ACCOUNT_ROLES | ADMIN_ROLES)
	return _run_bulk(_as_name_list(names), lambda name: approve_lpo(name))


@frappe.whitelist(methods=["POST"])
def bulk_reject_lpos(names: str | list | None = None, comment: str = "") -> dict:
	_require(ACCOUNT_ROLES | ADMIN_ROLES)
	if not (comment or "").strip():
		frappe.throw(_("Enter a rejection reason."))
	return _run_bulk(_as_name_list(names), lambda name: reject_lpo(name, comment))


@frappe.whitelist(methods=["POST"])
def cancel_purchase_order(name: str) -> dict:
	doc = _controlled_doc("Purchase Order", name, "cancel")
	if not cint(doc.get("controlled_procurement")) or doc.docstatus != 1:
		frappe.throw(_("Only a submitted Purchase Order can be cancelled."))
	doc.flags.controlled_procurement_api = True
	doc.cancel()
	return {"name": doc.name, "docstatus": doc.docstatus}


@frappe.whitelist(methods=["POST"])
def delete_purchase_order(name: str) -> dict:
	"""Delete a controlled draft or unlinked cancelled Purchase Order."""
	doc = _controlled_doc("Purchase Order", name, "delete")
	if not cint(doc.get("controlled_procurement")):
		frappe.throw(_("Only a controlled Purchase Order can be deleted."))
	_assert_deletable_controlled_document(doc)
	if doc.docstatus == 0 and doc.get("custom_po_price_status") in {"Pending Accounts", "Pending CEO"}:
		from construction_management.api import po_price_approvals
		proposal = json.loads(doc.get("custom_po_price_proposal") or "{}")
		po_price_approvals.supersede(doc, f"Superseded because draft Purchase Order {doc.name} was deleted.")
		for request_name in proposal.get("requests", []):
			# A deleted document cannot remain a Link target. The decision comment
			# retains its original PO number in the immutable price-request audit.
			frappe.db.set_value("Controlled Buying Price Request", request_name, "purchase_order", None, update_modified=False)
	doc.delete()
	return {"name": name, "deleted": True}


@frappe.whitelist(methods=["POST"])
def delete_purchase_receipt(name: str) -> dict:
	"""Delete a controlled draft or unlinked cancelled Receive Note."""
	doc = _controlled_doc("Purchase Receipt", name, "delete")
	_assert_deletable_controlled_document(doc)
	doc.delete()
	return {"name": name, "deleted": True}


@frappe.whitelist(methods=["POST"])
def amend_purchase_order(name: str) -> dict:
	doc = _controlled_doc("Purchase Order", name, "amend")
	if not cint(doc.get("controlled_procurement")) or doc.docstatus != 2:
		frappe.throw(_("Cancel the Purchase Order before amending it."))
	amended = frappe.copy_doc(doc)
	amended.amended_from = doc.name
	# copy_doc defaults to ignore_no_copy=True; clear the unique create idempotency key.
	amended.custom_creation_request_id = None
	amended.custom_lpo_approval_status = "Draft" if doc.get("custom_lpo_type") == "Manual" else None
	amended.custom_lpo_approval_history = None
	amended.flags.controlled_procurement_api = True
	amended.insert()
	return {"name": amended.name, "docstatus": amended.docstatus}


@frappe.whitelist(methods=["POST"])
def create_purchase_receipt(data: str | dict) -> dict:
	_require(PURCHASE_ROLES)
	data = _as_dict(data)
	po_name = data.get("purchase_order")
	po = _controlled_doc("Purchase Order", po_name)
	open_lpo = _is_open_lpo(po)
	_validate_receivable_po(po, open_lpo)
	if not cint(po.controlled_procurement):
		return _receive_desk_purchase_order(po, data)
	items = data.get("items") or []
	_validate_line_data(items)
	_receipt_items_by_po_line(po, items, open_lpo)
	po_rows = {row.name: row for row in po.items}
	# receipt-level allocation defaults to the PO's but can be changed on the receipt
	defaults = {field: data.get(field) or po.get(field) for field in ("project", "bill_no", "boq_item")}
	_validate_allocation(defaults, po.company)
	_validate_company_links(po.company, project=defaults["project"])
	doc = frappe.new_doc("Purchase Receipt")
	doc.update({
		"company": po.company, "supplier": po.supplier, "project": defaults["project"],
		"bill_no": defaults["bill_no"], "boq_item": defaults["boq_item"],
		"posting_date": data.get("received_on"),
		"set_posting_time": 1 if data.get("received_on") else 0,
		"custom_purchase_order": po.name,
		"supplier_delivery_note": (data.get("supplier_delivery_note") or "").strip(),
		"controlled_procurement": 1,
	})
	for data_row in items:
		po_row = po_rows.get(data_row.get("purchase_order_item"))
		item = _catalog_item(po_row.item_code)
		_validate_company_links(po.company, warehouse=data_row.get("warehouse"))
		line_project = (data_row if data_row.get("use_override") else defaults).get("project")
		warehouse_project = get_warehouse_project(data_row.get("warehouse")) if data_row.get("warehouse") else ""
		if warehouse_project and warehouse_project != line_project:
			frappe.throw(_("Warehouse {0} belongs to project {1}, so goods received there are booked to {1}. Select project {1} (and its Bill / BOQ Item) or a different warehouse.").format(
				data_row.get("warehouse"), warehouse_project))
		row = doc.append("items", {
			"item_code": po_row.item_code, "qty": flt(data_row.get("qty")), "uom": po_row.uom, "stock_uom": po_row.stock_uom,
			"conversion_factor": 1, "rate": po_row.rate, "purchase_order": po.name,
			"purchase_order_item": po_row.name, "warehouse": data_row.get("warehouse"),
		})
		_apply_allocation(row, data_row if data_row.get("use_override") else {}, defaults, po.company)
		_set_row_type(row, item)
	# VAT follows the order unless the receiver changes it
	taxes_and_charges = data.get("taxes_and_charges") if "taxes_and_charges" in data else po.taxes_and_charges
	_apply_taxes(doc, taxes_and_charges, [
		data_row.get("vat") or _vat_choice(po_rows[data_row.get("purchase_order_item")].item_tax_template, po.company)
		for data_row in items
	])
	doc.flags.controlled_procurement_api = True
	doc.insert()
	return {"doctype": doc.doctype, "name": doc.name}


def _receive_desk_purchase_order(po, data: dict) -> dict:
	"""Receive note for a PO raised in Desk, built with ERPNext's own PO → receipt mapping.

	Items and UOMs come from the PO as they are; Project / Bill / BOQ are optional here.
	"""
	from erpnext.buying.doctype.purchase_order.mapper import make_purchase_receipt
	open_lpo = _is_open_lpo(po)

	lines = [row for row in data.get("items") or [] if flt(row.get("qty")) > 0]
	if not lines:
		frappe.throw(_("Enter a receive qty on at least one line."))
	grouped = _receipt_items_by_po_line(po, lines, open_lpo)
	doc = make_purchase_receipt(po.name)
	if open_lpo:
		mapped = {row.purchase_order_item for row in doc.items}
		for po_row in po.items:
			if po_row.name in mapped or cint(po_row.get("closed")) or cint(po_row.get("delivered_by_supplier")):
				continue
			doc.append("items", {
				"item_code": po_row.item_code, "item_name": po_row.item_name,
				"description": po_row.description, "qty": 0, "received_qty": 0,
				"uom": po_row.uom, "stock_uom": po_row.stock_uom,
				"conversion_factor": po_row.conversion_factor, "rate": po_row.rate,
				"purchase_order": po.name, "purchase_order_item": po_row.name,
				"warehouse": po_row.warehouse, "project": po_row.project,
			})
	header = {field: data.get(field) or po.get(field) for field in ("project", "bill_no", "boq_item")}
	# receiving a Desk PO only needs a project (checked per line below); validate what was filled in
	if header.get("project"):
		_validate_allocation(header, po.company, mode="project_only")
	_validate_company_links(po.company, project=header["project"])
	templates = {row.purchase_order_item: row.as_dict() for row in doc.items}
	if set(grouped) - set(templates):
		frappe.throw(_("Receive note lines must come from Purchase Order {0}.").format(po.name))
	doc.set("items", [])
	vat_choices = []
	for line in lines:
		template = templates[line["purchase_order_item"]].copy()
		template.pop("name", None)
		row = doc.append("items", template)
		row.qty = flt(line.get("qty"))
		row.received_qty = row.qty
		row.warehouse = line.get("warehouse") or row.warehouse
		_validate_company_links(po.company, warehouse=row.warehouse)
		allocation = {field: (line if line.get("use_override") else header).get(field) for field in ("project", "bill_no", "boq_item")}
		if line.get("use_override") and allocation.get("project"):
			_validate_allocation(allocation, po.company, mode="project_only")
		for field, value in allocation.items():
			if value and row.meta.has_field(field):
				row.set(field, value)
		# receipts post to a project (site rule, enforced again on submit); a site warehouse decides it
		row.project = get_warehouse_project(row.warehouse) or row.project
		if not row.project:
			frappe.throw(_("{0}: select a Project for this receive note (warehouse {1} is not linked to one).").format(
				row.item_name or row.item_code, row.warehouse))
		vat_choices.append(line.get("vat") or _vat_choice(row.item_tax_template, po.company))
	doc.update({
		"project": header["project"], "bill_no": header["bill_no"], "boq_item": header["boq_item"],
		"posting_date": data.get("received_on") or doc.posting_date, "set_posting_time": 1 if data.get("received_on") else 0,
		"supplier_delivery_note": (data.get("supplier_delivery_note") or "").strip(),
		"custom_purchase_order": po.name if doc.meta.has_field("custom_purchase_order") else None,
		"controlled_procurement": 0,
	})
	taxes_and_charges = data.get("taxes_and_charges") if "taxes_and_charges" in data else po.taxes_and_charges
	_apply_taxes(doc, taxes_and_charges, vat_choices)
	doc.flags.controlled_procurement_api = True
	doc.insert()
	return {"doctype": doc.doctype, "name": doc.name}


@frappe.whitelist(methods=["POST"])
def create_material_transfer(data: str | dict) -> dict:
	_require(STOCK_ROLES)
	data = _as_dict(data)
	company = _document_company(data.get("company"))
	data["company"] = company
	items = data.get("items") or []
	_validate_line_data(items)
	source = data.get("source_warehouse")
	target = data.get("target_warehouse")
	_validate_transfer_warehouses(company, source, target)
	doc = frappe.new_doc("Stock Entry")
	doc.update({
		"company": company, "purpose": "Material Transfer", "from_warehouse": source, "to_warehouse": target,
		# mandatory in this ERPNext version; the standard type for the Material Transfer purpose
		"stock_entry_type": frappe.db.get_value("Stock Entry Type", {"purpose": "Material Transfer"}, "name", order_by="is_standard desc") or "Material Transfer",
		"posting_date": data.get("posting_date"), "controlled_procurement": 1,
	})
	for row_data in items:
		item = _catalog_item(row_data.get("item_code"))
		if item.controlled_item_type not in STOCK_TYPES or not cint(item.is_stock_item):
			frappe.throw(_("Only controlled Stockable or Consumable items can be transferred."))
		_append_item(doc, {**row_data, "source_warehouse": source, "target_warehouse": target}, {}, warehouses=True)
	doc.flags.controlled_procurement_api = True
	doc.insert()
	return {"doctype": doc.doctype, "name": doc.name}


def _validate_transfer_warehouses(company: str, source: str, target: str) -> None:
	if not source or not target or source == target:
		frappe.throw(_("Select different From and To warehouses."))
	for warehouse in (source, target):
		if not frappe.db.exists("Warehouse", {"name": warehouse, "company": company, "is_group": 0, "disabled": 0}):
			frappe.throw(_("Select an enabled leaf warehouse in company {0}.").format(company))


@frappe.whitelist(methods=["GET"])
def get_transfer_stock_preview(source_warehouse: str, item_codes: str | list[str]) -> dict:
	"""Current physical stock for controlled items in the active company's source warehouse."""
	_require(STOCK_ROLES)
	company = _document_company(None)
	if not frappe.db.exists("Warehouse", {"name": source_warehouse, "company": company, "is_group": 0, "disabled": 0}):
		frappe.throw(_("Select an enabled source warehouse in the active company."), frappe.PermissionError)
	if isinstance(item_codes, str):
		try:
			item_codes = frappe.parse_json(item_codes)
		except (ValueError, TypeError):
			frappe.throw(_("Invalid item list."))
	if not isinstance(item_codes, list) or len(item_codes) > 100 or any(not isinstance(code, str) or not code for code in item_codes):
		frappe.throw(_("Select up to 100 valid items."))
	codes = list(dict.fromkeys(item_codes))
	for code in codes:
		item = _catalog_item(code)
		if item.controlled_item_type not in STOCK_TYPES or not cint(item.is_stock_item):
			frappe.throw(_("Only controlled Stockable or Consumable items can be transferred."))
	balances = dict.fromkeys(codes, 0.0)
	if codes:
		for row in frappe.get_all("Bin", filters={"warehouse": source_warehouse, "item_code": ["in", codes]}, fields=["item_code", "actual_qty"]):
			balances[row.item_code] = flt(row.actual_qty)
	return {"warehouse": source_warehouse, "balances": balances}


@frappe.whitelist(methods=["POST"])
def update_material_transfer(name: str, data: str | dict) -> dict:
	doc = _controlled_doc("Stock Entry", name, "write")
	if doc.docstatus != 0 or not cint(doc.controlled_procurement):
		frappe.throw(_("Only a controlled draft Material Transfer can be edited here."))
	data = _as_dict(data)
	if data.get("company") and data["company"] != doc.company:
		frappe.throw(_("The transfer company cannot be changed."))
	source, target = data.get("source_warehouse"), data.get("target_warehouse")
	_validate_transfer_warehouses(doc.company, source, target)
	items = data.get("items") or []
	_validate_line_data(items)
	doc.update({"posting_date": data.get("posting_date"), "from_warehouse": source, "to_warehouse": target,
		"project": None, "bill_no": None, "boq_item": None})
	doc.set("items", [])
	for row in items:
		item = _catalog_item(row.get("item_code"))
		if item.controlled_item_type not in STOCK_TYPES or not cint(item.is_stock_item):
			frappe.throw(_("Only controlled Stockable or Consumable items can be transferred."))
		_append_item(doc, {**row, "source_warehouse": source, "target_warehouse": target}, {}, warehouses=True)
	doc.flags.controlled_procurement_api = True
	doc.save()
	return {"name": doc.name}


@frappe.whitelist(methods=["POST"])
def update_purchase_receipt(name: str, data: str | dict) -> dict:
	"""Edit a draft receipt while preserving its linked PO item identity."""
	doc = _controlled_doc("Purchase Receipt", name, "write")
	if doc.docstatus != 0:
		frappe.throw(_("Only a draft Receive Note can be edited."))
	data = _as_dict(data)
	po_names = {row.purchase_order for row in doc.items if row.purchase_order}
	if len(po_names) != 1 or data.get("purchase_order") != next(iter(po_names)):
		frappe.throw(_("The linked Purchase Order cannot be changed on a draft Receive Note."))
	po = _controlled_doc("Purchase Order", data["purchase_order"])
	if po.company != doc.company:
		frappe.throw(_("The Purchase Order must be submitted in the same company."))
	open_lpo = _is_open_lpo(po)
	_validate_receivable_po(po, open_lpo)
	items = [row for row in data.get("items") or [] if flt(row.get("qty")) > 0]
	if not items:
		frappe.throw(_("Enter a receive qty on at least one line."))
	_receipt_items_by_po_line(po, items, open_lpo)
	po_rows = {row.name: row for row in po.items}
	original_rows = {row.name: row for row in doc.items}
	original_by_po = {row.purchase_order_item: row for row in doc.items}
	mapped_rows = {}
	if not cint(doc.controlled_procurement):
		from erpnext.buying.doctype.purchase_order.mapper import make_purchase_receipt
		mapped_rows = {row.purchase_order_item: row for row in make_purchase_receipt(po.name).items}
	used_names = [row.get("name") for row in items if row.get("name")]
	if len(used_names) != len(set(used_names)) or any(name not in original_rows for name in used_names):
		frappe.throw(_("Invalid or repeated Receive Note item row."))
	defaults = {field: data.get(field) or po.get(field) for field in ALLOCATION_FIELDS}
	if cint(doc.controlled_procurement):
		_validate_allocation(defaults, doc.company)
	elif defaults.get("project"):
		_validate_allocation(defaults, doc.company, mode="project_only")
	doc.update({"posting_date": data.get("received_on"), "set_posting_time": 1,
		"supplier_delivery_note": (data.get("supplier_delivery_note") or "").strip(), **defaults})
	doc.set("items", [])
	for data_row in items:
		po_row = po_rows.get(data_row.get("purchase_order_item"))
		warehouse = data_row.get("warehouse") or data.get("warehouse")
		_validate_company_links(doc.company, warehouse=warehouse)
		if not warehouse:
			frappe.throw(_("Select a receiving warehouse for every item."))
		allocation = data_row if data_row.get("use_override") else defaults
		project = allocation.get("project") or get_warehouse_project(warehouse)
		warehouse_project = get_warehouse_project(warehouse)
		if warehouse_project and project != warehouse_project:
			frappe.throw(_("Warehouse {0} belongs to project {1}.").format(warehouse, warehouse_project))
		if not project:
			frappe.throw(_("Select a Project for this Receive Note."))
		original = original_rows.get(data_row.get("name"))
		if original and original.purchase_order_item != po_row.name:
			frappe.throw(_("Receive Note row does not match its Purchase Order line."))
		base_row = original or original_by_po.get(po_row.name) or mapped_rows.get(po_row.name)
		base_data = base_row.as_dict() if base_row else {}
		if not original:
			base_data.pop("name", None)
		row = doc.append("items", {**base_data,
			"item_code": po_row.item_code, "qty": flt(data_row.get("qty")),
			"uom": po_row.uom, "stock_uom": po_row.stock_uom, "conversion_factor": po_row.conversion_factor,
			"rate": po_row.rate, "purchase_order": po.name, "purchase_order_item": po_row.name,
			"warehouse": warehouse, "project": project})
		if cint(doc.controlled_procurement):
			_apply_allocation(row, data_row if data_row.get("use_override") else {}, defaults, doc.company)
			if warehouse_project and row.project != warehouse_project:
				frappe.throw(_("Warehouse {0} belongs to project {1}.").format(warehouse, warehouse_project))
			_set_row_type(row, _catalog_item(po_row.item_code))
		else:
			_validate_allocation({**allocation, "project": project}, doc.company, mode="project_only")
			for field in ("bill_no", "boq_item"):
				row.set(field, allocation.get(field) or None)
	_apply_taxes(doc, data.get("taxes_and_charges"), [row.get("vat") for row in items])
	doc.flags.controlled_procurement_api = True
	doc.save()
	return {"name": doc.name}


SUPPLIER_DO_CANCELLED_MARKER = "#CANCELLED#"


def original_supplier_delivery_note(value: str | None) -> str:
	"""Strip cancel tombstone so amend/new GRN can reuse the supplier's DO number."""
	value = (value or "").strip()
	if SUPPLIER_DO_CANCELLED_MARKER in value:
		return value.split(SUPPLIER_DO_CANCELLED_MARKER, 1)[0].strip()
	return value


def release_supplier_delivery_note_on_cancel(doc) -> None:
	"""Free the unique Supplier DO No. while keeping a required audit value on the cancelled PR."""
	if not getattr(doc, "meta", None) or not doc.meta.has_field("supplier_delivery_note"):
		return
	dn = (doc.get("supplier_delivery_note") or "").strip()
	if not dn or SUPPLIER_DO_CANCELLED_MARKER in dn:
		return
	tombstone = f"{dn}{SUPPLIER_DO_CANCELLED_MARKER}{doc.name}"
	doc.db_set("supplier_delivery_note", tombstone, update_modified=False)
	doc.supplier_delivery_note = tombstone


def _restore_supplier_delivery_note_for_amend(doc) -> None:
	if not getattr(doc, "meta", None) or not doc.meta.has_field("supplier_delivery_note"):
		return
	doc.supplier_delivery_note = original_supplier_delivery_note(doc.get("supplier_delivery_note"))


def _document_is_unlinked(doctype: str, name: str) -> bool:
	from frappe.model.delete_doc import get_linked_docs

	return not get_linked_docs(frappe.get_doc(doctype, name), method="Delete")


def _can_delete_controlled_document(doc) -> bool:
	if not frappe.has_permission(doc.doctype, "delete", doc=doc):
		return False
	if doc.docstatus == 0:
		return True
	if doc.docstatus == 2:
		return _document_is_unlinked(doc.doctype, doc.name)
	return False


def _assert_deletable_controlled_document(doc) -> None:
	if doc.docstatus == 0:
		return
	if doc.docstatus == 2:
		if not _document_is_unlinked(doc.doctype, doc.name):
			frappe.throw(
				_("{0} {1} is still linked to other documents and cannot be deleted.").format(
					_(doc.doctype), doc.name
				)
			)
		return
	frappe.throw(_("Cancel {0} {1} before deleting it, or delete only a draft.").format(_(doc.doctype), doc.name))


def _cancel_and_amend(doctype: str, name: str, amend: bool) -> dict:
	doc = _controlled_doc(doctype, name, "amend" if amend else "cancel")
	if amend:
		if doc.docstatus != 2:
			frappe.throw(_("Cancel the document before amending it."))
		amended = frappe.copy_doc(doc)
		amended.amended_from = doc.name
		# Cancel tombstoned the DO on the cancelled doc; restore the original for the amend.
		_restore_supplier_delivery_note_for_amend(amended)
		amended.flags.controlled_procurement_api = True
		amended.insert()
		return {"name": amended.name, "docstatus": amended.docstatus}
	if doc.docstatus != 1:
		frappe.throw(_("Only submitted documents can be cancelled."))
	doc.flags.controlled_procurement_api = True
	doc.cancel()
	return {"name": doc.name, "docstatus": doc.docstatus}


@frappe.whitelist(methods=["POST"])
def cancel_purchase_receipt(name: str) -> dict:
	return _cancel_and_amend("Purchase Receipt", name, False)


@frappe.whitelist(methods=["POST"])
def amend_purchase_receipt(name: str) -> dict:
	return _cancel_and_amend("Purchase Receipt", name, True)


@frappe.whitelist(methods=["POST"])
def cancel_purchase_invoice(name: str) -> dict:
	return _cancel_and_amend("Purchase Invoice", name, False)


@frappe.whitelist(methods=["POST"])
def amend_purchase_invoice(name: str) -> dict:
	return _cancel_and_amend("Purchase Invoice", name, True)


@frappe.whitelist(methods=["POST"])
def cancel_material_transfer(name: str) -> dict:
	return _cancel_and_amend("Stock Entry", name, False)


@frappe.whitelist(methods=["POST"])
def amend_material_transfer(name: str) -> dict:
	return _cancel_and_amend("Stock Entry", name, True)


@frappe.whitelist(methods=["POST"])
def update_purchase_order_items(name: str, items: str | list) -> dict:
	"""Change qty/rate or add/remove lines on a submitted PO (ERPNext "Update Items").

	Lines already received cannot go below their received qty or be removed; ERPNext enforces that.
	"""
	update_child_qty_rate = _erpnext_update_child_qty_rate()

	po = _controlled_doc("Purchase Order", name, "write")
	if po.get("custom_po_price_status") in {"Pending Accounts", "Pending CEO"}:
		frappe.throw(_("This Purchase Order is locked pending CEO price approval."))
	if po.docstatus != 1:
		frappe.throw(_("Only submitted purchase orders can be updated this way; edit the draft instead."))
	if po.status in ("Closed", "Completed", "Delivered"):
		frappe.throw(_("Purchase Order {0} is {1} and can no longer be changed.").format(name, po.status))
	items = frappe.parse_json(items) if isinstance(items, str) else items
	if not items:
		frappe.throw(_("A purchase order needs at least one item."))
	existing = {row.name: row for row in po.items}
	trans_items = []
	for row in items:
		if flt(row.get("qty")) <= 0:
			frappe.throw(_("Quantity must be greater than zero for {0}.").format(row.get("item_code")))
		current = existing.get(row.get("docname"))
		if not current and cint(po.controlled_procurement):
			_catalog_item(row.get("item_code"))  # new lines on controlled POs must be catalog items
		trans_items.append({
			"docname": current.name if current else None,
			"item_code": current.item_code if current else row.get("item_code"),
			"qty": flt(row.get("qty")),
			"rate": flt(row.get("rate")),
			"uom": current.uom if current else (row.get("uom") or frappe.db.get_value("Item", row.get("item_code"), "stock_uom")),
			"conversion_factor": current.conversion_factor if current else 1,
			"schedule_date": str(current.schedule_date if current else po.schedule_date),
			**({} if current else {"item_tax_template": _line_vat_templates(po.company).get(row.get("vat") or "standard")}),
		})
	if cint(po.controlled_procurement):
		from construction_management.api import po_price_approvals
		for row in items:
			row.setdefault("rate_edited", True)
		changes = po_price_approvals.prepare_rates(po.company, po.supplier, po.transaction_date, items, po)
		if changes:
			po_price_approvals.stage(po, changes, trans_items)
			return {"name": name, "price_approval_status": frappe.db.get_value("Purchase Order", name, "custom_po_price_status")}
	update_child_qty_rate("Purchase Order", json.dumps(trans_items), name)
	return {"name": name}


INVOICE_SOURCES = {"Purchase Order": "po_detail", "Purchase Receipt": "pr_detail"}


def _map_invoice(source_doctype: str, source_name: str):
	"""Unsaved Purchase Invoice for the still-unbilled part of a PO or receipt (ERPNext's own mapper)."""
	if source_doctype not in INVOICE_SOURCES:
		frappe.throw(_("Invoices can be created from a Purchase Order or Purchase Receipt."))
	source = _controlled_doc(source_doctype, source_name)
	if source.docstatus != 1:
		frappe.throw(_("{0} {1} must be submitted before it can be invoiced.").format(_(source_doctype), source_name))
	if flt(source.get("per_billed")) >= 100:
		frappe.throw(_("{0} {1} is already fully invoiced.").format(_(source_doctype), source_name))
	if source_doctype == "Purchase Order":
		from erpnext.buying.doctype.purchase_order.mapper import make_purchase_invoice
	else:
		from erpnext.stock.doctype.purchase_receipt.mapper import make_purchase_invoice
	return source, make_purchase_invoice(source_name)


@frappe.whitelist(methods=["GET"])
def get_invoice_sources(search: str = "") -> list[dict]:
	"""Submitted POs and receipts of the active company that still have something to invoice."""
	_require(PURCHASE_ROLES)
	like = f"%{(search or '').strip()}%"
	rows = []
	for doctype, date_field, extra in (
		("Purchase Order", "transaction_date", {"status": ["not in", ["Closed", "On Hold"]]}),
		("Purchase Receipt", "posting_date", {"is_return": 0}),
	):
		filters = _scoped(_workspace_filters(doctype, {"docstatus": 1, "per_billed": ["<", 100], **extra}))
		for row in frappe.get_list(
			doctype, filters=filters, or_filters=[["name", "like", like], ["supplier_name", "like", like]] if search else None,
			fields=["name", "supplier", "supplier_name", "grand_total", "per_billed", f"{date_field} as date", "project"],
			order_by="modified desc", limit_page_length=40,
		):
			rows.append({"doctype": doctype, **row})
	return sorted(rows, key=lambda row: str(row["date"]), reverse=True)


@frappe.whitelist(methods=["GET"])
def get_invoice_lines(source_doctype: str, source_name: str) -> dict:
	"""The lines, tax and supplier an invoice from this PO / receipt would start with."""
	_require(PURCHASE_ROLES)
	source, mapped = _map_invoice(source_doctype, source_name)
	key_field = INVOICE_SOURCES[source_doctype]
	return {
		"supplier": mapped.supplier,
		"supplier_name": mapped.supplier_name,
		"company": mapped.company,
		"currency": mapped.currency,
		"project": source.get("project"),
		# follow the PO / receipt; if it carried no tax, start from the company default (usually VAT 5%)
		"taxes_and_charges": mapped.taxes_and_charges or source.get("taxes_and_charges")
		or frappe.db.get_value(TAX_MASTER, {"company": mapped.company, "is_default": 1, "disabled": 0}, "name") or "",
		"due_date": mapped.get("due_date"),
		"lines": [
			{
				"key": row.get(key_field),
				"item_code": row.item_code,
				"item_name": row.item_name,
				"description": row.description,
				"unbilled_qty": flt(row.qty),
				"qty": flt(row.qty),
				"rate": flt(row.rate),
				"uom": row.uom,
				"project": row.get("project"),
				"bill_no": row.get("bill_no"),
				"boq_item": row.get("boq_item"),
				"purchase_order": row.get("purchase_order"),
				"purchase_receipt": row.get("purchase_receipt"),
				"vat": _vat_choice(row.get("item_tax_template"), mapped.company),
			}
			for row in mapped.items
			if flt(row.qty) > 0
		],
	}


@frappe.whitelist(methods=["POST"])
def create_purchase_invoice(data: str | dict) -> dict:
	"""Draft Purchase Invoice for chosen lines of a PO or receipt, never above the unbilled qty."""
	_require(PURCHASE_ROLES)
	frappe.has_permission("Purchase Invoice", "create", throw=True)
	data = _as_dict(data)
	if not (data.get("supplier_invoice_no") or "").strip():
		frappe.throw(_("Supplier invoice no is required."))
	source_doctype, source_name = data.get("source_doctype"), data.get("source_name")
	_source, doc = _map_invoice(source_doctype, source_name)
	key_field = INVOICE_SOURCES[source_doctype]
	mapped_rows = {row.get(key_field): row for row in doc.items}
	lines = [line for line in data.get("items") or [] if flt(line.get("qty")) > 0]
	if not lines:
		frappe.throw(_("Add at least one line to invoice."))
	if len({line.get("key") for line in lines}) != len(lines):
		frappe.throw(_("The same source line cannot be invoiced twice."))
	kept = []
	for line in lines:
		row = mapped_rows.get(line.get("key"))
		if not row:
			frappe.throw(_("Invoice lines must come from {0} {1}.").format(_(source_doctype), source_name))
		if flt(line.get("qty")) > flt(row.qty) + 0.0001:
			frappe.throw(_("{0}: only {1} is left to invoice.").format(row.item_name or row.item_code, flt(row.qty)))
		row.qty = flt(line.get("qty"))
		if line.get("rate") is not None:
			row.rate = flt(line.get("rate"))
		row.idx = len(kept) + 1
		# mandatory Warehouse Type dimension; every existing invoice books lines to "Transit"
		for field in ("site", "rejected_site"):
			if row.meta.has_field(field) and not row.get(field):
				row.set(field, "Transit")
		kept.append(row)
	doc.set("items", kept)
	doc.custom_supplier_invoice_no = data["supplier_invoice_no"].strip()
	doc.bill_date = data.get("supplier_invoice_date") or None
	# ERPNext resets posting_date to today unless set_posting_time is on.
	doc.set_posting_time = 1
	doc.posting_date = getdate(data.get("posting_date") or doc.posting_date)
	if data.get("due_date"):
		doc.due_date = data["due_date"]
		doc.payment_terms_template = None
		doc.set("payment_schedule", [])
	if data.get("remarks"):
		doc.remarks = data["remarks"]
	_apply_taxes(doc, data.get("taxes_and_charges"), [line.get("vat") for line in lines])
	doc.insert()
	return {"doctype": doc.doctype, "name": doc.name}


def _payable_invoice(invoice: str):
	doc = _controlled_doc("Purchase Invoice", invoice)
	if doc.docstatus != 1:
		frappe.throw(_("Submit the invoice before recording a payment."))
	if flt(doc.outstanding_amount) <= 0:
		frappe.throw(_("Purchase Invoice {0} is already fully paid.").format(invoice))
	return doc


def _remaining_payable(doc) -> float:
	"""Outstanding amount not already reserved by pending post-dated cheques."""
	try:
		from redtra_customisation.redtra_customisation.doctype.post_dated_cheques.post_dated_cheques import (
			get_remaining_pdc_allocatable,
		)
	except ImportError:
		return flt(doc.outstanding_amount)
	return flt(get_remaining_pdc_allocatable(doc.doctype, doc.name))


@frappe.whitelist(methods=["GET"])
def get_payment_defaults(invoice: str) -> dict:
	"""Amounts, payment modes and the company's bank / cash accounts for the Record payment form."""
	_require(PURCHASE_ROLES)
	doc = _payable_invoice(invoice)
	company = frappe.db.get_value("Company", doc.company, ["default_bank_account", "default_cash_account"], as_dict=True)
	modes = frappe.get_all("Mode of Payment", filters={"enabled": 1}, fields=["name", "type"], order_by="name asc")
	for mode in modes:
		mode["account"] = frappe.db.get_value("Mode of Payment Account", {"parent": mode.name, "company": doc.company}, "default_account") or (
			company.default_cash_account if mode.type == "Cash" else company.default_bank_account
		)
	return {
		"invoice": doc.name,
		"supplier": doc.supplier,
		"supplier_name": doc.supplier_name,
		"company": doc.company,
		"currency": doc.currency,
		"outstanding_amount": flt(doc.outstanding_amount),
		"payable_amount": _remaining_payable(doc),
		"supplier_invoice_no": doc.get("custom_supplier_invoice_no") or "",
		"modes": modes,
	}


@frappe.whitelist(methods=["GET"])
def get_po_pdc_defaults(purchase_order: str) -> dict:
	"""Options for an unallocated supplier cheque linked to a submitted PO."""
	_require(PURCHASE_ROLES)
	po = _controlled_doc("Purchase Order", purchase_order)
	if po.docstatus != 1 or po.status in ("Closed", "On Hold"):
		frappe.throw(_("Select a submitted, open Purchase Order."))
	if po.currency != frappe.db.get_value("Company", po.company, "default_currency"):
		frappe.throw(_("PO-level PDC creation currently requires the company currency."))
	frappe.has_permission("Post Dated Cheques", "create", throw=True)
	modes = frappe.get_all("Mode of Payment", filters={"enabled": 1}, fields=["name", "type"], order_by="name")
	modes = [mode for mode in modes if "cheque" in mode.name.lower() and mode.type == "Bank"]
	for mode in modes:
		mode["account"] = frappe.db.get_value("Mode of Payment Account", {"parent": mode.name,
			"company": po.company}, "default_account") or frappe.db.get_value("Company", po.company, "default_bank_account")
	return {"purchase_order": po.name, "company": po.company, "supplier": po.supplier,
		"supplier_name": po.supplier_name, "currency": po.currency, "order_total": flt(po.grand_total), "modes": modes}


@frappe.whitelist(methods=["POST"])
def create_po_pdc(purchase_order: str, data: str | dict) -> dict:
	"""Record a future-dated cheque against a PO without pretending an invoice exists."""
	_require(PURCHASE_ROLES)
	data = _as_dict(data)
	frappe.db.sql("select name from `tabPurchase Order` where name = %s for update", purchase_order)
	po = _controlled_doc("Purchase Order", purchase_order, "read")
	if po.docstatus != 1 or po.status in ("Closed", "On Hold"):
		frappe.throw(_("Select a submitted, open Purchase Order."))
	if po.currency != frappe.db.get_value("Company", po.company, "default_currency"):
		frappe.throw(_("PO-level PDC creation currently requires the company currency."))
	frappe.has_permission("Post Dated Cheques", "create", throw=True)
	frappe.has_permission("Post Dated Cheques", "submit", throw=True)
	amount = flt(data.get("amount"))
	if amount <= 0 or amount > flt(po.grand_total):
		frappe.throw(_("Cheque amount must be positive and no greater than the Purchase Order total."))
	existing_total = sum(flt(value) for value in frappe.get_all("Post Dated Cheques",
		filters={"custom_purchase_order": po.name, "docstatus": 1}, pluck="amount"))
	if existing_total + amount > flt(po.grand_total) + 0.001:
		frappe.throw(_("Active cheques for this Purchase Order would exceed its total value."))
	cheque_date = getdate(data.get("reference_date"))
	if cheque_date <= getdate(frappe.utils.today()):
		frappe.throw(_("A post-dated cheque needs a future cheque date."))
	cheque_no = (data.get("reference_no") or "").strip()
	if not cheque_no:
		frappe.throw(_("Enter the cheque number."))
	mode = data.get("mode_of_payment")
	if "cheque" not in (mode or "").lower() or frappe.db.get_value("Mode of Payment", mode, "type") != "Bank":
		frappe.throw(_("Select a bank cheque mode of payment."))
	account = data.get("bank_account")
	account_row = frappe.db.get_value("Account", account, ["company", "account_type", "is_group"], as_dict=True) if account else None
	if not account_row or account_row.company != po.company or account_row.account_type != "Bank" or account_row.is_group:
		frappe.throw(_("Select a bank account belonging to the Purchase Order company."))
	frappe.db.sql("select name from `tabSupplier` where name = %s for update", po.supplier)
	if frappe.db.exists("Post Dated Cheques", {"company": po.company, "party_type": "Supplier",
		"party": po.supplier, "reference_no": cheque_no, "docstatus": ["!=", 2]}):
		frappe.throw(_("This supplier already has an active PDC with cheque number {0}.").format(cheque_no))
	pdc = frappe.get_doc({"doctype": "Post Dated Cheques", "company": po.company,
		"custom_purchase_order": po.name, "project": po.project, "cost_center": po.cost_center,
		"posting_date": frappe.utils.today(), "payment_type": "Pay", "party_type": "Supplier",
		"party": po.supplier, "party_name": po.supplier_name, "mode_of_payment": mode,
		"reference_no": cheque_no, "reference_date": cheque_date, "amount": amount,
		"bank_account": account,
		"notes": _("Linked to Purchase Order {0}; no Purchase Invoice is allocated yet.").format(po.name),
	})
	pdc.insert()
	pdc.submit()
	return {"doctype": pdc.doctype, "name": pdc.name, "purchase_order": po.name}


@frappe.whitelist(methods=["POST"])
def record_invoice_payment(invoice: str, data: str | dict) -> dict:
	"""Pay a submitted invoice: a Payment Entry now, or a Post Dated Cheque for a future-dated cheque."""
	_require(PURCHASE_ROLES)
	data = _as_dict(data)
	doc = _payable_invoice(invoice)
	amount = flt(data.get("amount"))
	remaining = _remaining_payable(doc)
	if amount <= 0:
		frappe.throw(_("Enter an amount greater than zero."))
	if amount > remaining + 0.01:
		frappe.throw(_("Only {0} is left to pay on {1} after pending post-dated cheques.").format(
			frappe.format_value(remaining, {"fieldtype": "Currency", "options": doc.currency}), invoice))
	mode = data.get("mode_of_payment")
	mode_type = frappe.db.get_value("Mode of Payment", mode, "type") if mode else None
	if not mode_type:
		frappe.throw(_("Select a mode of payment."))
	account = data.get("account")
	account_row = frappe.db.get_value("Account", account, ["company", "account_type", "is_group"], as_dict=True) if account else None
	if not account_row or account_row.company != doc.company or account_row.is_group or account_row.account_type not in ("Bank", "Cash"):
		frappe.throw(_("Select a bank or cash account of {0}.").format(doc.company))
	today = getdate(frappe.utils.today())
	reference_date = getdate(data.get("reference_date") or today)
	reference_no = (data.get("reference_no") or "").strip()

	if "cheque" in mode.lower() and reference_date > today:
		if not reference_no:
			frappe.throw(_("Enter the cheque number."))
		frappe.has_permission("Post Dated Cheques", "create", throw=True)
		pdc = frappe.get_doc({
			"doctype": "Post Dated Cheques",
			"company": doc.company,
			"project": doc.get("project"),
			"cost_center": doc.get("cost_center"),
			"posting_date": reference_date,
			"payment_type": "Pay",
			"party_type": "Supplier",
			"party": doc.supplier,
			"party_name": doc.supplier_name,
			"mode_of_payment": mode,
			"reference_no": reference_no,
			"reference_date": reference_date,
			"amount": amount,
			"bank_account": account,
			"invoice_references": [{
				"reference_doctype": doc.doctype,
				"reference_name": doc.name,
				"total_amount": flt(doc.grand_total),
				"outstanding_amount": flt(doc.outstanding_amount),
				"allocated_amount": amount,
			}],
		})
		pdc.insert()
		pdc.submit()
		return {"doctype": pdc.doctype, "name": pdc.name, "post_dated": 1}

	from erpnext.accounts.doctype.payment_entry.payment_entry import get_payment_entry

	pe = get_payment_entry(doc.doctype, doc.name, party_amount=amount, bank_account=account)
	pe.mode_of_payment = mode
	pe.posting_date = getdate(data.get("posting_date") or today)
	pe.reference_no = reference_no or doc.get("custom_supplier_invoice_no") or doc.name
	pe.reference_date = reference_date
	if data.get("remarks"):
		pe.remarks = data["remarks"]
	pe.insert()
	pe.submit()
	return {"doctype": pe.doctype, "name": pe.name, "post_dated": 0}


@frappe.whitelist(methods=["POST"])
def submit_document(doctype: str, name: str) -> dict:
	"""Submit any workspace draft (controlled or raised in Desk) after normal permission checks."""
	doc = _controlled_doc(doctype, name, "submit")
	if doc.docstatus != 0:
		frappe.throw(_("{0} is not a draft.").format(name))
	if doctype == "Purchase Order":
		from construction_management.api.lpo_workflow import lpo_type
		if cint(doc.get("controlled_procurement")) and lpo_type(doc) == "Manual":
			frappe.throw(_("Send this Manual LPO for Accounts and CEO approval instead of submitting it directly."))
		# Defense in depth: do not rely only on before_submit when API flag skips validate.
		price_status = frappe.db.get_value("Purchase Order", name, "custom_po_price_status")
		if cint(doc.get("controlled_procurement")) and price_status in {"Pending Accounts", "Pending CEO"}:
			frappe.throw(_("This Purchase Order is pending CEO price approval and cannot be submitted."), frappe.PermissionError)
	doc.flags.controlled_procurement_api = True
	doc.submit()
	return {"doctype": doctype, "name": doc.name, "docstatus": doc.docstatus}


def _workbook_rows(path: Path) -> tuple[list[dict], list[dict]]:
	from openpyxl import load_workbook

	book = load_workbook(path, read_only=True, data_only=True)
	rows = []
	header_index = None
	for sheet in book.worksheets:
		candidate_rows = list(sheet.iter_rows(values_only=True))
		candidate_header = next((i for i, row in enumerate(candidate_rows) if "Item Group" in row and "Description" in row), None)
		if candidate_header is not None:
			rows, header_index = candidate_rows, candidate_header
			break
	if header_index is None:
		frappe.throw(_("The workbook must contain Item Group and Description columns."))
	headers = {str(value).strip(): position for position, value in enumerate(rows[header_index]) if value}
	valid, exceptions = [], []
	for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
		group = str(row[headers["Item Group"]] or "").strip()
		description = str(row[headers["Description"]] or "").strip()
		if not any(value is not None and str(value).strip() for value in row):
			continue
		if group == "Item Group" or not group or not description:
			exceptions.append({"row": row_number, "reason": "Missing item group or description"})
			continue
		valid.append({
			"row": row_number, "item_group": group, "item_name": description,
			"item_code": str(row[headers.get("Item Code", -1)] or "").strip() if "Item Code" in headers else "",
			"item_type": str(row[headers.get("Item Type", -1)] or STOCKABLE).strip() if "Item Type" in headers else STOCKABLE,
			"stock_uom": str(row[headers["Stock UOM"]] or "").strip() if "Stock UOM" in headers else "Nos",
			"asset_category": str(row[headers.get("Asset Category", -1)] or "").strip() if "Asset Category" in headers else "",
			"expense_account": str(row[headers.get("Expense Account", -1)] or "").strip() if "Expense Account" in headers else "",
			"supplier": str(row[headers.get("Supplier", -1)] or "").strip() if "Supplier" in headers else "",
			"rate": flt(row[headers.get("Rates", -1)] or 0) if "Rates" in headers else 0,
		})
	return valid, exceptions


def _item_code(item_name: str) -> str:
	base = "CP-" + frappe.scrub(item_name).upper()[:100]
	code = base
	index = 2
	while frappe.db.exists("Item", code):
		code = f"{base[:95]}-{index}"
		index += 1
	return code


def _ensure_group(group: str) -> None:
	if frappe.db.exists("Item Group", group):
		return
	frappe.get_doc({"doctype": "Item Group", "item_group_name": group, "parent_item_group": "All Item Groups", "is_group": 0}).insert()


def _normalise_name(value: str) -> str:
	return " ".join((value or "").split()).casefold()


def _supplier_name(supplier: str) -> str:
	"""Resolve harmless whitespace/case variations without creating suppliers."""
	if frappe.db.exists("Supplier", supplier):
		return supplier
	wanted = _normalise_name(supplier)
	for row in frappe.get_all("Supplier", fields=["name", "supplier_name"]):
		if wanted in {_normalise_name(row.name), _normalise_name(row.supplier_name)}:
			return row.name
	return ""


def _add_supplier_and_price(item, supplier: str, rate: float) -> str:
	if not supplier:
		return "missing supplier"
	requested_supplier = supplier
	supplier = _supplier_name(requested_supplier)
	if not supplier:
		return f"supplier not found: {requested_supplier}"
	if supplier not in {row.supplier for row in item.get("supplier_items") or []}:
		item.append("supplier_items", {"supplier": supplier})
	if rate > 0:
		from construction_management.api.controlled_price_requests import _price, approved_price_write

		current = _price(item.name, supplier, "Standard Buying", item.stock_uom)
		if not current or flt(current.price_list_rate) != rate:
			source = getattr(frappe.flags, "controlled_catalog_price_source", None)
			if not source:
				frappe.throw(_("A CEO-approved catalog request is required to change a buying price."), frappe.PermissionError)
			with approved_price_write(source):
				if current:
					price = frappe.get_doc("Item Price", current.name)
					price.price_list_rate = rate
					price.custom_controlled_catalog_request = source if source != "legacy-seed-migration" else None
					price.save(ignore_permissions=True)
				else:
					price = frappe.get_doc({"doctype": "Item Price", "item_code": item.name,
						"price_list": "Standard Buying", "price_list_rate": rate, "buying": 1,
						"supplier": supplier, "uom": item.stock_uom,
						"custom_controlled_catalog_request": source if source != "legacy-seed-migration" else None})
					price.insert(ignore_permissions=True)
			if source == "legacy-seed-migration":
				price.add_comment("Info", "Legacy packaged catalog seed: audited migration exception for buying price.")
	return ""


def _apply_catalog_supplier_mapping(item, supplier: str, rate: float, item_name: str = "") -> None:
	"""Apply an already-approved supplier mapping to a catalog Item."""
	error = _add_supplier_and_price(item, supplier, rate)
	if error:
		frappe.throw(
			_("Could not apply supplier details for {0}: {1}").format(item_name or item.item_name, error)
		)
	item.flags.controlled_procurement_api = True
	item.save(ignore_permissions=True)


def import_catalog_from_path(path: str) -> dict:
	"""Administrator-only development helper; production uses catalog file or Desk upload."""
	_require(ADMIN_ROLES)
	rows, exceptions = _workbook_rows(Path(path))
	return _import_catalog_rows(rows, exceptions, Path(path).name, "Workbook Import", "uploaded", _rows_checksum(rows))


@frappe.whitelist(methods=["POST"])
def preview_catalog_workbook(file_url: str) -> dict:
	_require(CATALOG_ROLES)
	rows, exceptions = _workbook_rows(_uploaded_file_path(file_url))
	company = _active_company()
	valid = []
	for row in rows:
		try:
			valid.append(_catalog_request_row(row, company))
		except frappe.ValidationError as error:
			exceptions.append({"row": row["row"], "item": row["item_name"], "reason": str(error)})
	if not rows:
		exceptions.append({"row": 0, "reason": "Add at least one catalog item row."})
	return {
		"source": Path(file_url).name, "checksum": _rows_checksum(rows), "valid_rows": len(valid),
		"create": sum(not row.get("result_item") for row in valid),
		"reuse": sum(bool(row.get("result_item")) for row in valid),
		"errors": exceptions, "rows": valid,
	}


@frappe.whitelist(methods=["POST"])
def import_catalog_workbook(file_url: str, request_origin: str = "") -> dict:
	_require(CATALOG_ROLES)
	rows, exceptions = _workbook_rows(_uploaded_file_path(file_url))
	if exceptions:
		frappe.throw(_("Correct workbook errors before submitting the import."))
	company = _active_company()
	checksum = _rows_checksum(rows)
	frappe.db.sql("select name from `tabCompany` where name = %s for update", company)
	existing = frappe.db.get_value("Controlled Catalog Request", {
		"company": company, "source": "Workbook Import", "checksum": checksum,
	}, ["name", "status"], as_dict=True)
	if existing:
		return {"name": existing.name, "status": existing.status, "already_imported": True}
	prepared = []
	for row in rows:
		try:
			prepared.append(_catalog_request_row(row, company))
		except frappe.ValidationError:
			frappe.throw(_("Correct workbook errors before submitting the import."))
	return create_catalog_request({"company": company, "source": "Workbook Import", "source_file": file_url,
		"request_origin": request_origin,
		"checksum": checksum, "items": prepared})


def import_packaged_catalog() -> dict:
	"""Import the versioned application asset during a production migration."""
	path = Path(frappe.get_app_path("construction_management", "data", "controlled_procurement_catalog_v1.json"))
	payload = json.loads(path.read_text())
	previous = getattr(frappe.flags, "controlled_catalog_price_source", None)
	frappe.flags.controlled_catalog_price_source = "legacy-seed-migration"
	try:
		return _import_catalog_rows(
			payload["items"], [], payload["source_filename"], "Workbook Import",
			payload["version"], payload["checksum"],
		)
	finally:
		frappe.flags.controlled_catalog_price_source = previous


def _import_catalog_rows(rows, exceptions, source_name, source_type, version, checksum) -> dict:
	by_name = defaultdict(list)
	for row in rows:
		by_name[row["item_name"]].append(row)
	created = reused = 0
	for item_name, source_rows in by_name.items():
		first = source_rows[0]
		_ensure_group(first["item_group"])
		existing = frappe.db.get_value("Item", {"item_name": item_name, "stock_uom": "Nos", "disabled": 0}, "name")
		if existing:
			item = frappe.get_doc("Item", existing)
			reused += 1
			item.item_group = first["item_group"]
		else:
			item = frappe.get_doc({
				"doctype": "Item", "item_code": _item_code(item_name), "item_name": item_name,
				"item_group": first["item_group"], "stock_uom": "Nos", "is_stock_item": 1,
				"is_purchase_item": 1, "is_fixed_asset": 0,
			})
			item.flags.controlled_procurement_api = True
			item.insert()
			created += 1
		item.controlled_procurement_catalog = 1
		item.controlled_item_type = STOCKABLE
		item.controlled_import_source = source_name
		item.controlled_imported_on = now_datetime()
		item.controlled_catalog_source = source_type
		item.controlled_catalog_version = version
		item.controlled_catalog_checksum = checksum
		for source in source_rows:
			error = _add_supplier_and_price(item, source["supplier"], source["rate"])
			if error:
				exceptions.append({"row": source["row"], "item": item_name, "reason": error})
		item.flags.controlled_procurement_api = True
		item.save()
	return {"created": created, "reused": reused, "exceptions": exceptions, "valid_rows": len(rows)}


def _catalog_preview(rows, exceptions, source_name, checksum) -> dict:
	preview = {"create": 0, "reuse": 0, "supplier_exceptions": list(exceptions)}
	by_name = defaultdict(list)
	for row in rows:
		by_name[row["item_name"]].append(row)
	for item_name, source_rows in by_name.items():
		if frappe.db.exists("Item", {"item_name": item_name, "stock_uom": "Nos", "disabled": 0}):
			preview["reuse"] += 1
		else:
			preview["create"] += 1
		for row in source_rows:
			if row.get("supplier") and not _supplier_name(row["supplier"]):
				preview["supplier_exceptions"].append({"row": row["row"], "item": item_name, "reason": f"supplier not found: {row['supplier']}"})
	return {"source": source_name, "checksum": checksum, "valid_rows": len(rows), **preview}


def _rows_checksum(rows) -> str:
	encoded = json.dumps(rows, ensure_ascii=False, separators=(",", ":"), sort_keys=True).encode()
	return hashlib.sha256(encoded).hexdigest()


def _uploaded_file_path(file_url: str) -> Path:
	file_doc = frappe.get_doc("File", {"file_url": file_url})
	if not file_doc.file_name.lower().endswith(".xlsx"):
		frappe.throw(_("Upload an Excel .xlsx workbook."))
	return Path(file_doc.get_full_path())


@frappe.whitelist(methods=["POST"])
def download_catalog_template() -> dict:
	"""Create the approved import template on demand so it always matches validation."""
	_require(CATALOG_ROLES)
	from openpyxl import Workbook

	book = Workbook()
	sheet = book.active
	sheet.title = "Catalog Items"
	sheet.append(["Item Code", "Description", "Item Group", "Item Type", "Stock UOM", "Asset Category", "Expense Account", "Supplier", "Rates"])
	sheet.freeze_panes = "A2"
	for column in sheet.columns:
		sheet.column_dimensions[column[0].column_letter].width = 22
	instructions = book.create_sheet("Instructions")
	for line in (
		"Add one item per row on Catalog Items; do not change the column headers.",
		"Description, Item Group, Item Type, and Supplier are required on every row.",
		"Supplier must be an existing Supplier ID; select it from the Procurement supplier list.",
		"Item Type must be Stockable, Consumable, Asset, or Service.",
		"Stock UOM must match an existing UOM; Nos is used if the cell is blank.",
		"Example Stockable row: Description = Example stock item; Item Group = Products; Item Type = Stockable; Supplier = an existing Supplier ID.",
		"Example Consumable row: Description = Example consumable; Item Group = Products; Item Type = Consumable; Supplier = an existing Supplier ID.",
		"Asset Category is needed for Asset rows unless a company default is configured.",
		"Expense Account is needed for Service rows unless a company default is configured.",
		"Item Code is optional; Rates may be zero.",
	):
		instructions.append([line])
	instructions.column_dimensions["A"].width = 100
	output = BytesIO()
	book.save(output)
	file_doc = save_file("controlled_catalog_import_template.xlsx", output.getvalue(), None, None, is_private=0)
	return {"file_url": file_doc.file_url}


def _explicit_ceo() -> bool:
	from construction_management.api.controlled_price_requests import has_explicit_ceo_role
	return has_explicit_ceo_role()


def _has_explicit_role(role: str, user: str | None = None) -> bool:
	return bool(frappe.db.exists("Has Role", {
		"parent": user or frappe.session.user, "parenttype": "User", "role": role,
	}))


def _catalog_ceo(user: str | None = None) -> bool:
	user = user or frappe.session.user
	return user in CATALOG_CEO_USERS and _has_explicit_role("CEO", user)


def _can_request_catalog_disable(user: str | None = None) -> bool:
	"""Accounts Manager, Purchase Manager, or catalog CEO may request disable."""
	return (
		_has_explicit_role("Accounts Manager", user)
		or _has_explicit_role("Purchase Manager", user)
		or _catalog_ceo(user)
	)


def _catalog_request_origin(requested: str | None = None) -> str:
	if requested and requested not in {"Purchase", "Accounts"}:
		frappe.throw(_("Select Purchase or Accounts as the request origin."))
	if requested == "Accounts" and not (_has_explicit_role("Accounts Manager") or _catalog_ceo()):
		frappe.throw(_("Only an Accounts Manager can create as Accounts."), frappe.PermissionError)
	if requested:
		return requested
	if _roles().intersection({"Purchase User", "Purchase Manager"}):
		return "Purchase"
	return "Accounts" if _has_explicit_role("Accounts Manager") else "Purchase"


def _catalog_request_status(origin: str) -> str:
	return "Pending CEO" if origin == "Accounts" else "Pending Accounts"


def _catalog_uom(value: str) -> str:
	value = str(value or "Nos").strip()
	aliases = {"nos.": "Nos", "ton": "Tonne", "hrs": "Hour", "monthly": "Months"}
	value = aliases.get(value.casefold(), value)
	if frappe.db.exists("UOM", value):
		return value
	for uom in frappe.get_all("UOM", pluck="name", limit_page_length=0):
		if uom.casefold() == value.casefold():
			return uom
	frappe.throw(_("Stock UOM {0} does not exist.").format(value))


def _catalog_request_row(data: dict, company: str, supplier_link_only: bool = False, allow_uom_edit: bool = False) -> dict:
	item_type = data.get("item_type") or STOCKABLE
	if item_type not in STOCK_TYPES | {ASSET, SERVICE}:
		frappe.throw(_("Select Stockable, Consumable, Asset, or Service."))
	if not data.get("item_name") or not data.get("item_group"):
		frappe.throw(_("Every catalog row needs a Description and Item Group."))
	stock_uom = _catalog_uom(data.get("stock_uom") or "Nos")
	if not frappe.db.exists("Item Group", data["item_group"]):
		frappe.throw(_("Item Group {0} does not exist.").format(data["item_group"]))
	supplier = str(data.get("supplier") or "").strip()
	if not supplier:
		frappe.throw(_("Supplier is required for every catalog item."))
	if not frappe.db.exists("Supplier", {"name": supplier, "disabled": 0}):
		frappe.throw(_("Supplier {0} is not active.").format(supplier))
	company_rules = frappe.db.get_value(
		"Company", company, ["controlled_asset_category", "controlled_service_expense_account"], as_dict=True,
	)
	asset_category = data.get("asset_category") or company_rules.controlled_asset_category
	expense_account = data.get("expense_account") or company_rules.controlled_service_expense_account
	if item_type == ASSET and not asset_category:
		frappe.throw(_("Configure or select an Asset Category for Asset items."))
	if item_type == ASSET and asset_category and not frappe.db.exists("Asset Category", asset_category):
		frappe.throw(_("Asset Category {0} does not exist.").format(data["asset_category"]))
	if item_type == SERVICE and not expense_account:
		frappe.throw(_("Configure or select a Service Expense Account for Service items."))
	if item_type == SERVICE and expense_account:
		from redtra_customisation.redtra_customisation.service_item_validator import ServiceItemValidator

		ServiceItemValidator.validate_expense_account(expense_account, company)
	item_code = (data.get("item_code") or "").strip()
	existing = frappe.db.get_value("Item", item_code, "name") if item_code else None
	# A supplied code deliberately identifies the catalog master.  This lets an
	# approved catalog row use its own UOM when a legacy Item shares the same
	# description with a different UOM.
	if not item_code:
		existing = frappe.db.get_value("Item", {"item_name": data["item_name"], "disabled": 0}, "name")
	if existing and frappe.db.get_value("Item", existing, "disabled"):
		frappe.throw(_("Item {0} is disabled. Create an approved replacement item instead.").format(existing))
	if existing:
		existing_meta = frappe.db.get_value(
			"Item",
			existing,
			["stock_uom", "controlled_procurement_catalog"],
			as_dict=True,
		)
		existing_uom = existing_meta.stock_uom
		if existing_uom != stock_uom:
			if allow_uom_edit:
				if frappe.db.exists("Stock Ledger Entry", {"item_code": existing, "is_cancelled": 0}):
					frappe.throw(
						_(
							"Existing item {0} has stock UOM {1}; requested {2}. "
							"Changing stock UOM requires an item edit with no stock history."
						).format(existing, existing_uom, stock_uom)
					)
			elif cint(existing_meta.controlled_procurement_catalog):
				# Already in the controlled catalogue: lock Stock UOM rather than fail the request.
				stock_uom = existing_uom
			else:
				frappe.throw(
					_(
						"Existing item {0} has stock UOM {1}; requested {2}. "
						"Its UOM cannot be changed by import. Supply a new Item Code to create a separate catalog item."
					).format(existing, existing_uom, stock_uom)
				)
	rate = flt(data.get("rate"))
	if rate < 0:
		frappe.throw(_("Buying rate cannot be negative."))
	return {
		"item_code": item_code,
		"item_name": data["item_name"].strip(),
		"item_group": data["item_group"],
		"item_type": item_type,
		"stock_uom": stock_uom,
		"asset_category": asset_category or "",
		"expense_account": expense_account or "",
		"supplier": supplier,
		"rate": rate,
		"action": "Link Supplier" if supplier_link_only and existing else "Update" if existing else "Create",
		"result_item": existing or "",
	}


@frappe.whitelist(methods=["POST"])
def create_catalog_supplier_link_request(data: str | dict) -> dict:
	"""Create an approval-routed request to link one supplier to selected catalog Items."""
	_require(CATALOG_ROLES)
	data = _as_dict(data)
	company = _document_company(data.get("company"))
	supplier = str(data.get("supplier") or "").strip()
	if not supplier or not frappe.db.exists("Supplier", {"name": supplier, "disabled": 0}):
		frappe.throw(_("Select an active Supplier."))

	item_codes = _unique_catalog_item_codes(data.get("item_codes"))
	if not item_codes:
		frappe.throw(_("Select at least one controlled catalog item."))
	items = _selected_catalog_items(item_codes)
	if len(items) != len(item_codes):
		frappe.throw(_("Every selected item must be an active controlled catalog item."))

	linked_items = set(
		frappe.get_all(
			"Item Supplier",
			filters={"parent": ["in", item_codes], "parenttype": "Item", "supplier": supplier},
			pluck="parent",
			limit_page_length=0,
		)
	)
	request_items = [
		{
			"item_code": item.name,
			"item_name": item.item_name,
			"item_group": item.item_group,
			"item_type": item.controlled_item_type,
			"stock_uom": item.stock_uom,
			"asset_category": item.asset_category or "",
			"supplier": supplier,
			"rate": 0,
		}
		for item in items
		if item.name not in linked_items
	]
	if not request_items:
		return {"name": "", "status": "", "requested": 0, "already_linked": len(item_codes)}

	result = create_catalog_request({
		"company": company,
		"source": "Manual",
		"request_origin": data.get("request_origin"),
		"supplier_link_only": 1,
		"items": request_items,
	})
	return {**result, "requested": len(request_items), "already_linked": len(linked_items)}


@frappe.whitelist(methods=["POST"])
def create_catalog_item_group_change_request(data: str | dict) -> dict:
	"""Create an approval-routed request to move selected catalog Items to one Item Group."""
	_require(CATALOG_ROLES)
	data = _as_dict(data)
	company = _document_company(data.get("company"))
	target_group = _catalog_leaf_item_group(data.get("item_group"))
	item_codes = _unique_catalog_item_codes(data.get("item_codes"))
	if not item_codes:
		frappe.throw(_("Select at least one controlled catalog item."))
	items = _selected_catalog_items(item_codes)
	if len(items) != len(item_codes):
		frappe.throw(_("Every selected item must be an active controlled catalog item."))

	request_items = [
		{"item_code": item.name, "item_group": target_group}
		for item in items
		if item.item_group != target_group
	]
	if not request_items:
		return {"name": "", "status": "", "requested": 0, "already_correct": len(item_codes)}

	result = create_catalog_request({
		"company": company,
		"source": "Manual",
		"request_origin": data.get("request_origin"),
		"item_group_change_only": 1,
		"items": request_items,
	})
	return {**result, "requested": len(request_items), "already_correct": len(item_codes) - len(request_items)}


def _unique_catalog_item_codes(item_codes) -> list[str]:
	if not isinstance(item_codes, list):
		frappe.throw(_("Selected item codes must be a list."))
	return list(dict.fromkeys(str(item_code).strip() for item_code in item_codes if str(item_code).strip()))


def _selected_catalog_items(item_codes: list[str]) -> list:
	return frappe.get_all(
		"Item",
		filters={
			"name": ["in", item_codes],
			"disabled": 0,
			"is_purchase_item": 1,
			"controlled_procurement_catalog": 1,
		},
		fields=[
			"name",
			"item_name",
			"item_group",
			"stock_uom",
			"controlled_item_type",
			"asset_category",
		],
		order_by="name asc",
		limit_page_length=0,
	)


def _catalog_leaf_item_group(item_group: str) -> str:
	item_group = str(item_group or "").strip()
	group = frappe.db.get_value("Item Group", item_group, ["name", "is_group"], as_dict=True)
	if not group or cint(group.is_group):
		frappe.throw(_("Select an active leaf Item Group."))
	return group.name


def _catalog_group_change_request_row(data: dict) -> dict:
	item_code = str(data.get("item_code") or "").strip()
	if not item_code:
		frappe.throw(_("Every Item Group change needs a catalog Item."))
	target_group = _catalog_leaf_item_group(data.get("item_group"))
	item = frappe.db.get_value(
		"Item",
		item_code,
		["name", "item_name", "item_group", "stock_uom", "controlled_item_type", "disabled",
			"is_purchase_item", "controlled_procurement_catalog"],
		as_dict=True,
	)
	if not item or item.disabled or not item.is_purchase_item or not item.controlled_procurement_catalog:
		frappe.throw(_("Item {0} is not an active controlled catalog item.").format(item_code))
	if item.item_group == target_group:
		frappe.throw(_("Item {0} already belongs to {1}.").format(item.name, target_group))
	return {
		"item_code": item.name,
		"item_name": item.item_name,
		"item_group": target_group,
		"item_type": item.controlled_item_type,
		"stock_uom": item.stock_uom,
		"asset_category": "",
		"expense_account": "",
		"supplier": "",
		"rate": 0,
		"action": "Change Item Group",
		"result_item": item.name,
	}


def _item_expense_account(item_name: str, company: str) -> str:
	"""Service expense accounts belong to Item Defaults, not every Item schema."""
	if not frappe.get_meta("Item").has_field("item_defaults"):
		return ""
	return frappe.db.get_value("Item Default", {"parent": item_name, "company": company}, "expense_account") or ""


@frappe.whitelist(methods=["GET"])
def get_catalog_item_edit_form(item_code: str) -> dict:
	"""Load the approved business fields for an approval-routed catalogue edit."""
	_require(CATALOG_ROLES)
	item = _catalog_item(item_code)
	if frappe.db.get_value("Item", item.name, "disabled"):
		frappe.throw(_("Disabled catalogue items cannot be edited."))
	fields = ["name", "item_name", "item_group", "stock_uom", "controlled_item_type", "asset_category"]
	values = frappe.db.get_value("Item", item.name, fields, as_dict=True)
	values.expense_account = _item_expense_account(item.name, _active_company())
	supplier = frappe.db.get_value("Item Supplier", {"parent": item.name, "parenttype": "Item"}, "supplier") or ""
	values.supplier = supplier
	values.rate = _catalog_buying_rate(item.name, supplier)
	return values


def _catalog_buying_rate(item_code: str, supplier: str | None) -> float:
	"""Active Standard Buying rate for a supplier; Item Edit must not invent a new price."""
	if not supplier:
		return 0.0
	return flt(frappe.db.get_value("Item Price", {
		"item_code": item_code, "supplier": supplier, "price_list": "Standard Buying", "buying": 1,
	}, "price_list_rate"))


@frappe.whitelist(methods=["POST"])
def create_catalog_item_edit_request(data: str | dict) -> dict:
	"""Request an approved change to an existing catalogue Item without changing its code."""
	_require(CATALOG_ROLES)
	payload = _as_dict(data)
	item_code = str(payload.get("item_code") or "").strip()
	if not item_code:
		frappe.throw(_("Select the catalogue item to edit."))
	item = _catalog_item(item_code)
	if frappe.db.get_value("Item", item.name, "disabled"):
		frappe.throw(_("Disabled catalogue items cannot be edited."))
	if payload.get("replacement_item_code"):
		frappe.throw(_("Item Code is immutable. Create an approved replacement mapping instead."))
	current = frappe.db.get_value(
		"Item", item.name,
		["item_name", "item_group", "stock_uom", "controlled_item_type", "asset_category"],
		as_dict=True,
	)
	current.expense_account = _item_expense_account(item.name, payload.get("company") or _active_company())
	supplier = (payload.get("supplier")
		or frappe.db.get_value("Item Supplier", {"parent": item.name, "parenttype": "Item"}, "supplier")
		or "")
	row = {
		"item_code": item.name,
		"item_name": payload.get("item_name") or current.item_name,
		"item_group": payload.get("item_group") or current.item_group,
		"item_type": payload.get("item_type") or current.controlled_item_type,
		"stock_uom": payload.get("stock_uom") or current.stock_uom,
		"asset_category": payload.get("asset_category") or current.asset_category,
		"expense_account": payload.get("expense_account") or current.expense_account,
		"supplier": supplier,
		# Price changes go through Propose price / Price requests, never Item Edit.
		"rate": _catalog_buying_rate(item.name, supplier),
	}
	if row["stock_uom"] != current.stock_uom and frappe.db.exists("Stock Ledger Entry", {"item_code": item.name, "is_cancelled": 0}):
		frappe.throw(_("Stock UOM cannot change after stock history exists. Create a replacement item and mapping instead."))
	return create_catalog_request({"company": payload.get("company"), "source": "Item Edit",
		"request_origin": payload.get("request_origin"), "items": [row]})


def _append_catalog_log(doc, action: str, comment: str = "") -> None:
	doc.flags.controlled_catalog_request_api = True
	doc.append("approval_log", {
		"action": action,
		"actor": frappe.session.user,
		"actioned_on": now_datetime(),
		"comment": comment,
	})


def _queue_catalog_notices(doc) -> None:
	"""Raven-only stage notices (no catalogue approval email)."""
	if doc.status not in {"Pending Accounts", "Pending CEO"}:
		return
	from construction_management.api.controlled_price_requests import _enqueue_catalog_notice

	_enqueue_catalog_notice(doc.name, doc.status)


def _queue_catalog_rejection_notice(doc) -> None:
	"""Tell the requester on Raven that Accounts/CEO sent the request back."""
	from construction_management.api.controlled_price_requests import _enqueue_catalog_notice

	_enqueue_catalog_notice(doc.name, "Needs Correction")


def _is_catalog_disable_request(doc) -> bool:
	return len(doc.items) == 1 and doc.items[0].action == "Disable"


def _catalog_disable_row(item_code: str) -> dict:
	item = frappe.db.get_value("Item", item_code, [
		"name", "item_name", "item_group", "stock_uom", "controlled_item_type", "modified",
		"disabled", "is_purchase_item", "controlled_procurement_catalog",
	], as_dict=True)
	if not item or item.disabled or not item.is_purchase_item or not item.controlled_procurement_catalog:
		frappe.throw(_("Select an active controlled catalogue item."))
	return {
		"item_code": item.name, "item_name": item.item_name, "item_group": item.item_group,
		"item_type": item.controlled_item_type, "stock_uom": item.stock_uom,
		"action": "Disable", "target_modified": item.modified, "result_item": item.name,
	}


def _validate_catalog_disable_row(row) -> None:
	current = _catalog_disable_row(row.item_code)
	if str(current["target_modified"]) != str(row.target_modified):
		frappe.throw(_("Item {0} changed since this disable request was created. Create a new request.").format(row.item_code))


@frappe.whitelist(methods=["POST"])
def create_catalog_disable_request(item_code: str, reason: str, company: str = "") -> dict:
	if not _can_request_catalog_disable():
		frappe.throw(_("Only Accounts Manager, Purchase Manager, or CEO can request item disabling."), frappe.PermissionError)
	company = company or _active_company()
	if not company or company not in _allowed_companies():
		frappe.throw(_("You do not have access to this company."), frappe.PermissionError)
	reason = (reason or "").strip()
	if not reason:
		frappe.throw(_("A reason is required to disable an item."))
	row = _catalog_disable_row(item_code)
	origin = "Accounts" if (_has_explicit_role("Accounts Manager") or _catalog_ceo()) else "Purchase"
	doc = frappe.get_doc({
		"doctype": "Controlled Catalog Request", "company": company, "source": "Manual",
		"request_origin": origin, "requested_by": frappe.session.user,
		"status": "Pending CEO", "disable_reason": reason, "items": [row],
	})
	_append_catalog_log(doc, "Created disable request", reason)
	doc.insert(ignore_permissions=True)
	if _catalog_ceo():
		_materialize_catalog_request(doc)
		doc.status = "Approved"
		_append_catalog_log(doc, "Approved by CEO", reason)
		doc.save(ignore_permissions=True)
	else:
		_queue_catalog_notices(doc)
	return {"name": doc.name, "status": doc.status}


@frappe.whitelist(methods=["POST"])
def create_catalog_request(data: str | dict) -> dict:
	"""Create a catalog change request; only approval materializes Item records."""
	_require(CATALOG_ROLES)
	data = _as_dict(data)
	company = data.get("company") or _active_company()
	if company not in _allowed_companies():
		frappe.throw(_("You do not have access to company {0}.").format(company), frappe.PermissionError)
	rows = data.get("items") or []
	if not rows:
		frappe.throw(_("Add at least one catalog item."))
	if data.get("supplier_link_only") and data.get("item_group_change_only"):
		frappe.throw(_("A catalog request can only perform one special action."))
	prepared = (
		[_catalog_group_change_request_row(row) for row in rows]
		if data.get("item_group_change_only")
		else [_catalog_request_row(row, company, supplier_link_only=bool(data.get("supplier_link_only")),
			allow_uom_edit=data.get("source") == "Item Edit") for row in rows]
	)
	if len({row["item_code"] or row["item_name"].casefold() for row in prepared}) != len(prepared):
		frappe.throw(_("The request contains duplicate item codes or descriptions."))
	origin = _catalog_request_origin(data.get("request_origin"))
	status = _catalog_request_status(origin)
	doc = frappe.get_doc({
		"doctype": "Controlled Catalog Request", "company": company, "status": status,
		"source": data.get("source") or "Manual", "request_origin": origin,
		"requested_by": frappe.session.user,
		"source_file": data.get("source_file") or "", "checksum": data.get("checksum") or "", "items": prepared,
	})
	_append_catalog_log(doc, "Created", f"Origin: {origin}")
	doc.insert(ignore_permissions=True)
	_queue_catalog_notices(doc)
	return {"name": doc.name, "status": doc.status}


@frappe.whitelist(methods=["GET"])
def get_catalog_requests(
	status: str = "",
	request_origin: str = "",
	source: str = "",
	kind: str = "",
) -> list[dict]:
	_require(CATALOG_ROLES)
	company = _active_company()
	filters = {"company": company} if company else {"company": ["in", _allowed_companies()]}
	if status:
		filters["status"] = status
	if request_origin:
		filters["request_origin"] = request_origin
	if source:
		filters["source"] = source
	requests = frappe.get_list("Controlled Catalog Request", filters=filters,
		fields=["name", "company", "status", "source", "request_origin", "requested_by", "creation", "modified", "rejection_reason"],
		order_by="modified desc", limit_page_length=100)
	if requests:
		priced = set(frappe.get_all("Controlled Catalog Request Item", filters={
			"parent": ["in", [row.name for row in requests]], "rate": [">", 0],
		}, pluck="parent", limit_page_length=0))
		disabling = set(frappe.get_all("Controlled Catalog Request Item", filters={
			"parent": ["in", [row.name for row in requests]], "action": "Disable",
		}, pluck="parent", limit_page_length=0))
		for row in requests:
			row.has_price_change = row.name in priced
			row.is_disable = row.name in disabling
		kind = (kind or "").strip().lower()
		if kind == "disable":
			requests = [row for row in requests if row.is_disable]
		elif kind == "change":
			requests = [row for row in requests if not row.is_disable]
	return requests


@frappe.whitelist(methods=["GET"])
def get_catalog_request(name: str) -> dict:
	_require(CATALOG_ROLES)
	doc = frappe.get_doc("Controlled Catalog Request", name)
	if doc.company not in _allowed_companies():
		frappe.throw(_("Not permitted"), frappe.PermissionError)
	result = doc.as_dict()
	result["currency"] = (frappe.db.get_value("Price List", "Standard Buying", "currency")
		or frappe.db.get_value("Company", doc.company, "default_currency") or "")
	result["review"] = _catalog_review_values(doc)
	queue_name = getattr(doc, "approval_email_queue", None)
	if queue_name:
		queue = frappe.db.get_value("Email Queue", queue_name, ["status", "error"], as_dict=True)
		if queue:
			result["approval_email_status"] = "Queued" if queue.status in {"Not Sent", "Sending"} else queue.status
			if queue.status in {"Error", "Partially Sent"}:
				result["approval_email_error"] = (queue.error or "Delivery failed. Check Email Queue.")[-400:]
	return result


def _catalog_review_values(doc) -> dict:
	"""Current master values are review context, never a replacement for the request snapshot."""
	codes = {row.result_item or row.item_code for row in doc.items
		if row.action != "Create" and (row.result_item or row.item_code)}
	if not codes:
		return {}
	items = {item.name: item for item in frappe.get_all("Item", filters={"name": ["in", list(codes)]},
		fields=["name", "item_name", "item_group", "controlled_item_type", "stock_uom", "asset_category",
			"is_stock_item", "is_fixed_asset"],
		limit_page_length=0)}
	defaults = {row.parent: row.expense_account for row in frappe.get_all("Item Default",
		filters={"parent": ["in", list(codes)], "company": doc.company},
		fields=["parent", "expense_account"], limit_page_length=0)}
	suppliers = defaultdict(list)
	for row in frappe.get_all("Item Supplier", filters={"parent": ["in", list(codes)], "parenttype": "Item"},
		fields=["parent", "supplier"], limit_page_length=0):
		suppliers[row.parent].append(row.supplier)
	prices = {}
	as_of = getdate(today())
	for row in frappe.get_all("Item Price", filters={"item_code": ["in", list(codes)],
		"price_list": "Standard Buying", "buying": 1},
		fields=["item_code", "supplier", "uom", "price_list_rate", "valid_from", "valid_upto"],
		order_by="valid_from desc, creation desc", limit_page_length=0):
		if (row.valid_from and getdate(row.valid_from) > as_of) or (row.valid_upto and getdate(row.valid_upto) < as_of):
			continue
		prices.setdefault((row.item_code, row.supplier or "", row.uom or ""), row.price_list_rate)
	review = {}
	for row in doc.items:
		code = row.result_item or row.item_code
		item = items.get(code)
		if not item:
			continue
		current_uom = item.stock_uom or row.stock_uom or ""
		rate = prices.get((code, row.supplier or "", current_uom))
		if rate is None:
			rate = prices.get((code, row.supplier or "", ""))
		if rate is None:
			rate = prices.get((code, "", current_uom))
		if rate is None:
			rate = prices.get((code, "", ""))
		current_type = item.controlled_item_type or (
			"Asset" if item.get("is_fixed_asset") else "Stockable" if item.get("is_stock_item") else "Service")
		review[row.name] = {"item_name": item.item_name, "item_group": item.item_group,
			"item_type": current_type, "stock_uom": item.stock_uom,
			"asset_category": item.asset_category or "", "expense_account": defaults.get(code) or "",
			"suppliers": suppliers.get(code, []), "rate": rate}
	return review


def _can_approve_catalog(doc) -> bool:
	if doc.status == "Pending Accounts":
		return _has_explicit_role("Accounts Manager") and doc.requested_by != frappe.session.user
	return doc.status == "Pending CEO" and _catalog_ceo()


@frappe.whitelist(methods=["POST"])
def retry_catalog_approval_email(name: str) -> dict:
	_require(CATALOG_ROLES)
	frappe.throw(_("Catalogue approval emails are disabled. Use Raven notification retry instead."))


@frappe.whitelist(methods=["POST"])
def approve_catalog_request(name: str, comment: str = "") -> dict:
	_require(CATALOG_ROLES)
	frappe.db.sql("select name from `tabControlled Catalog Request` where name = %s for update", name)
	doc = frappe.get_doc("Controlled Catalog Request", name)
	if doc.company not in _allowed_companies():
		frappe.throw(_("Not permitted"), frappe.PermissionError)
	if not _can_approve_catalog(doc):
		frappe.throw(_("You cannot approve this catalog request."), frappe.PermissionError)
	if doc.status == "Pending Accounts":
		doc.status = "Pending CEO"
		doc.price_notified_users = ""
		doc.price_notification_status = ""
		doc.price_notification_error = ""
		_append_catalog_log(doc, "Approved by Accounts", comment)
	else:
		_materialize_catalog_request(doc)
		doc.status = "Approved"
		_append_catalog_log(doc, "Approved by CEO", comment)
	doc.rejection_reason = ""
	_clear_catalog_approval_email(doc)
	doc.save(ignore_permissions=True)
	_queue_catalog_notices(doc)
	return {"name": doc.name, "status": doc.status}


@frappe.whitelist(methods=["POST"])
def reject_catalog_request(name: str, comment: str) -> dict:
	_require(CATALOG_ROLES)
	if not (comment or "").strip():
		frappe.throw(_("A rejection reason is required."))
	frappe.db.sql("select name from `tabControlled Catalog Request` where name = %s for update", name)
	doc = frappe.get_doc("Controlled Catalog Request", name)
	if doc.company not in _allowed_companies():
		frappe.throw(_("Not permitted"), frappe.PermissionError)
	if not _can_approve_catalog(doc):
		frappe.throw(_("You cannot reject this catalog request."), frappe.PermissionError)
	doc.status = "Needs Correction"
	doc.rejection_reason = comment.strip()
	doc.price_notified_users = ""
	doc.price_notification_status = ""
	doc.price_notification_error = ""
	_clear_catalog_approval_email(doc)
	_append_catalog_log(doc, "Rejected", doc.rejection_reason)
	doc.save(ignore_permissions=True)
	_queue_catalog_rejection_notice(doc)
	return {"name": doc.name, "status": doc.status}


@frappe.whitelist(methods=["POST"])
def bulk_approve_catalog_requests(names: str | list | None = None) -> dict:
	"""Approve many catalog requests; each uses the same rules as approve_catalog_request."""
	_require(CATALOG_ROLES)
	return _run_bulk(_as_name_list(names), lambda name: approve_catalog_request(name))


@frappe.whitelist(methods=["POST"])
def bulk_reject_catalog_requests(names: str | list | None = None, comment: str = "") -> dict:
	_require(CATALOG_ROLES)
	if not (comment or "").strip():
		frappe.throw(_("A rejection reason is required."))
	return _run_bulk(
		_as_name_list(names),
		lambda name: reject_catalog_request(name, comment),
	)


@frappe.whitelist(methods=["POST"])
def resubmit_catalog_request(name: str, data: str | dict | None = None) -> dict:
	_require(CATALOG_ROLES)
	frappe.db.sql("select name from `tabControlled Catalog Request` where name = %s for update", name)
	doc = frappe.get_doc("Controlled Catalog Request", name)
	if doc.company not in _allowed_companies():
		frappe.throw(_("Not permitted"), frappe.PermissionError)
	if doc.status != "Needs Correction" or doc.requested_by != frappe.session.user:
		frappe.throw(_("Only the requester can correct and resubmit this request."), frappe.PermissionError)
	payload = _as_dict(data)
	if str(payload.get("expected_modified") or "") != str(doc.modified):
		frappe.throw(_("This request changed since you opened it. Reload before resubmitting."))
	if _is_catalog_disable_request(doc):
		if not _can_request_catalog_disable():
			frappe.throw(_("Only Accounts Manager, Purchase Manager, or CEO can resubmit a disable request."), frappe.PermissionError)
		old_reason = doc.disable_reason or ""
		doc.disable_reason = (payload.get("disable_reason") or "").strip()
		if not doc.disable_reason:
			frappe.throw(_("A reason is required to disable an item."))
		_validate_catalog_disable_row(doc.items[0])
		changes = f"Disable reason: {old_reason} → {doc.disable_reason}" if old_reason != doc.disable_reason else "No field changes"
	else:
		rows, changes = _catalog_correction_rows(doc, payload.get("items"))
		doc.set("items", rows)
	doc.status = "Pending CEO" if _is_catalog_disable_request(doc) else _catalog_request_status(doc.request_origin or "Purchase")
	doc.rejection_reason = ""
	doc.price_notified_users = ""
	doc.price_notification_status = ""
	doc.price_notification_error = ""
	_clear_catalog_approval_email(doc)
	_append_catalog_log(doc, "Resubmitted", changes[:1000])
	doc.save(ignore_permissions=True)
	_queue_catalog_notices(doc)
	return {"name": doc.name, "status": doc.status}


def _clear_catalog_approval_email(doc) -> None:
	if hasattr(doc, "approval_email_queue"):
		doc.approval_email_queue = ""
	if hasattr(doc, "approval_email_error"):
		doc.approval_email_error = ""


CATALOG_CORRECTION_FIELDS = (
	"item_name", "item_group", "item_type", "stock_uom", "supplier", "rate",
	"asset_category", "expense_account",
)


def _catalog_correction_rows(doc, submitted: list[dict] | None) -> tuple[list[dict], str]:
	if not isinstance(submitted, list) or len(submitted) != len(doc.items):
		frappe.throw(_("Correct every original row; rows cannot be added or removed."))
	prepared, changes = [], []
	for original, edited in zip(doc.items, submitted, strict=True):
		if not isinstance(edited, dict) or str(edited.get("name")) != str(original.name):
			frappe.throw(_("Request rows changed. Reload before resubmitting."))
		for field in ("item_code", "action", "result_item"):
			if str(edited.get(field) or "") != str(original.get(field) or ""):
				frappe.throw(_("Item Code, action and target item cannot be changed."))
		allowed = {"item_group"} if original.action == "Change Item Group" else (
			{"supplier"} if original.action == "Link Supplier" else set(CATALOG_CORRECTION_FIELDS))
		if original.action not in {"Create", "Update", "Adopt", "Link Supplier", "Change Item Group"}:
			frappe.throw(_("This request action cannot be corrected."))
		for field in CATALOG_CORRECTION_FIELDS:
			if field not in allowed and field in edited and str(edited[field] or "") != str(original.get(field) or ""):
				frappe.throw(_("{0} cannot change for this request action.").format(field))
		candidate = original.as_dict()
		candidate.update({field: edited[field] for field in allowed if field in edited})
		validated = _catalog_corrected_row(doc, original, candidate)
		prepared.append(validated)
		for field in allowed:
			before, after = original.get(field), validated.get(field)
			if str(before or "") != str(after or ""):
				changes.append(f"{original.item_code or original.result_item or original.item_name}: {field} {before or '—'} → {after or '—'}")
	identities = [row["item_code"] or row["item_name"].casefold() for row in prepared]
	if len(set(identities)) != len(identities):
		frappe.throw(_("The corrected request contains duplicate item codes or descriptions."))
	return prepared, "; ".join(changes) if changes else "No field changes"


def _catalog_corrected_row(doc, original, candidate: dict) -> dict:
	if original.action == "Change Item Group":
		row = _catalog_group_change_request_row(candidate)
		if row["result_item"] != original.result_item:
			frappe.throw(_("The corrected row no longer targets the original item."))
		return row
	if candidate.get("item_group"):
		_catalog_leaf_item_group(candidate["item_group"])
	target = original.result_item or original.item_code
	if original.action != "Create" and target:
		candidate["item_code"] = target
	row = _catalog_request_row(candidate, doc.company,
		supplier_link_only=original.action == "Link Supplier", allow_uom_edit=doc.source == "Item Edit")
	if original.action != "Create" and candidate.get("stock_uom") and row["stock_uom"] != candidate["stock_uom"]:
		frappe.throw(_("The target item's Stock UOM changed. Reload and submit a new item-edit request."))
	if original.action == "Create":
		if row["action"] != "Create" or frappe.db.exists("Item", {"item_name": row["item_name"], "disabled": 0}):
			frappe.throw(_("The corrected name already belongs to an item. Start an item-edit request instead."))
	elif row["result_item"] != target:
		frappe.throw(_("The corrected row no longer targets the original item."))
	if original.action == "Link Supplier" and frappe.db.exists("Item Supplier", {
		"parent": target, "parenttype": "Item", "supplier": row["supplier"]}):
		frappe.throw(_("This supplier is already linked to the item."))
	row.update({"item_code": original.item_code or "", "action": original.action,
		"result_item": original.result_item or ""})
	return row


def _materialize_catalog_request(doc) -> None:
	if not _catalog_ceo():
		frappe.throw(_("Only Faraaz or Salman with the CEO role can approve catalogue changes."), frappe.PermissionError)
	previous_source = getattr(frappe.flags, "controlled_catalog_price_source", None)
	frappe.flags.controlled_catalog_price_source = doc.name
	try:
		frappe.db.sql("select name from `tabCompany` where name = %s for update", doc.company)
		_materialize_catalog_request_items(doc)
	finally:
		frappe.flags.controlled_catalog_price_source = previous_source


def _materialize_catalog_request_items(doc) -> None:
	materialized = {}
	for row in doc.items:
		if row.action == "Disable":
			_disable_catalog_item(row)
			continue
		if getattr(row, "action", "") == "Adopt":
			_adopt_catalog_item(doc, row)
			continue
		if getattr(row, "action", "") == "Link Supplier":
			item = frappe.get_doc("Item", row.result_item)
			_apply_catalog_supplier_mapping(item, row.supplier, flt(row.rate), row.item_name)
			continue
		if getattr(row, "action", "") == "Change Item Group":
			_change_catalog_item_group(row)
			continue

		identity = _catalog_row_identity(row)
		item = materialized.get(identity)
		if item:
			if row.supplier:
				_apply_catalog_supplier_mapping(item, row.supplier, flt(row.rate), row.item_name)
			row.result_item = item.name
			continue

		item = frappe.get_doc("Item", row.result_item) if row.result_item else _existing_catalog_item(row)
		if item and item.stock_uom != (row.stock_uom or "Nos"):
			frappe.throw(_("Item {0} stock UOM changed since this request was submitted.").format(item.name))
		if not item:
			item = frappe.new_doc("Item")
			item.item_code = row.item_code or _item_code(row.item_name)
		item.update({
			"item_name": row.item_name, "item_group": row.item_group, "stock_uom": row.stock_uom or "Nos",
			"is_purchase_item": 1, "controlled_procurement_catalog": 1,
			"controlled_item_type": row.item_type,
			"controlled_catalog_source": "Workbook Import" if doc.source == "Workbook Import" else "Manual Request",
			"controlled_catalog_version": doc.name,
		})
		if row.item_type in STOCK_TYPES:
			item.update({"is_stock_item": 1, "is_fixed_asset": 0})
		elif row.item_type == ASSET:
			item.update({"is_stock_item": 0, "is_fixed_asset": 1, "asset_category": row.asset_category or None})
		else:
			# Service items stay company-agnostic; expense account lives in Item Defaults per company.
			item.update({"is_stock_item": 0, "is_fixed_asset": 0})
			_set_item_expense_account(item, doc.company, row.expense_account or frappe.db.get_value("Company", doc.company, "controlled_service_expense_account"))
		item.flags.controlled_procurement_api = True
		if item.is_new():
			item.insert(ignore_permissions=True)
		else:
			item.save(ignore_permissions=True)
		if row.supplier:
			_apply_catalog_supplier_mapping(item, row.supplier, flt(row.rate), row.item_name)
		row.result_item = item.name
		materialized[identity] = item


def _set_item_expense_account(item, company: str, expense_account: str) -> None:
	"""Store a service account in Item Defaults on ERPNext schemas without Item.expense_account."""
	if not expense_account or not item.meta.has_field("item_defaults"):
		return
	default = next((row for row in item.item_defaults if row.company == company), None)
	if default:
		default.expense_account = expense_account
	else:
		item.append("item_defaults", {"company": company, "expense_account": expense_account})


def _ensure_service_item_default_for_company(item_code: str, company: str) -> None:
	"""Guarantee a Service catalog item has an Item Default expense row for the PO company."""
	if not company or not item_code:
		return
	if frappe.db.get_value("Item", item_code, "controlled_item_type") != SERVICE:
		return
	if not frappe.get_meta("Item").has_field("item_defaults"):
		return
	if frappe.db.exists("Item Default", {"parent": item_code, "company": company}):
		return
	expense_account = frappe.db.get_value("Company", company, "controlled_service_expense_account")
	if not expense_account:
		frappe.throw(
			_("Set Controlled Service Expense Account on Company {0} before using Service catalog items.").format(company)
		)
	from redtra_customisation.redtra_customisation.service_item_validator import ServiceItemValidator

	ServiceItemValidator.validate_expense_account(expense_account, company)
	item = frappe.get_doc("Item", item_code)
	_set_item_expense_account(item, company, expense_account)
	item.flags.controlled_procurement_api = True
	item.save(ignore_permissions=True)


def _change_catalog_item_group(row) -> None:
	"""Apply only the approved Item Group change; all other Item fields remain intact."""
	target_group = _catalog_leaf_item_group(row.item_group)
	item = frappe.get_doc("Item", row.result_item)
	if item.disabled or not item.is_purchase_item or not item.controlled_procurement_catalog:
		frappe.throw(_("Item {0} is no longer an active controlled catalog item.").format(item.name))
	if item.item_group != target_group:
		item.item_group = target_group
		item.flags.controlled_procurement_api = True
		item.save(ignore_permissions=True)
	row.result_item = item.name


def _disable_catalog_item(row) -> None:
	frappe.db.sql("select name from `tabItem` where name = %s for update", row.item_code)
	_validate_catalog_disable_row(row)
	item = frappe.get_doc("Item", row.item_code)
	item.disabled = 1
	item.flags.controlled_procurement_api = True
	item.save(ignore_permissions=True)
	row.result_item = item.name


def _catalog_row_identity(row) -> tuple[str, str, str]:
	return (_normalise_name(row.item_name), row.item_type, _normalise_name(row.stock_uom or "Nos"))


def _existing_catalog_item(row):
	items = frappe.get_all(
		"Item",
		filters={
			"disabled": 0,
			"controlled_procurement_catalog": 1,
			"item_name": row.item_name,
			"item_group": row.item_group,
			"stock_uom": row.stock_uom or "Nos",
			"controlled_item_type": row.item_type,
		},
		pluck="name",
		limit_page_length=0,
	)
	if not items:
		return None
	return frappe.get_doc("Item", min(items, key=_catalog_item_name_preference))


def _catalog_item_name_preference(name: str) -> tuple[bool, str]:
	return bool(re.search(r"-\d+$", name)), name


def _adopt_catalog_item(doc, row) -> None:
	item = frappe.get_doc("Item", row.result_item)
	if item.stock_uom != (row.stock_uom or "Nos"):
		frappe.throw(_("Item {0} stock UOM changed since this request was submitted.").format(item.name))
	item.update({
		"item_group": row.item_group,
		"controlled_procurement_catalog": 1,
		"controlled_item_type": row.item_type,
		"controlled_catalog_source": "Workbook Import" if doc.source == "Workbook Import" else "Manual Request",
		"controlled_catalog_version": doc.name,
	})
	item.flags.controlled_procurement_api = True
	item.save(ignore_permissions=True)


@frappe.whitelist(methods=["GET"])
def get_catalog_stock(item_code: str, company: str = "") -> dict:
	_require(CATALOG_ROLES)
	_catalog_item(item_code)
	company = company or _active_company()
	_validate_company_links(company)
	rows = frappe.db.sql(
		"""select bin.warehouse, bin.actual_qty, bin.valuation_rate
		from `tabBin` bin join `tabWarehouse` warehouse on warehouse.name = bin.warehouse
		where bin.item_code = %(item_code)s and warehouse.company = %(company)s and warehouse.is_group = 0
		order by bin.warehouse""", {"item_code": item_code, "company": company}, as_dict=True,
	)
	return {"rows": rows, "total_qty": sum(flt(row.actual_qty) for row in rows)}


@frappe.whitelist(methods=["POST"])
def record_opening_stock(data: str | dict) -> dict:
	_require(OPENING_STOCK_ROLES)
	data = _as_dict(data)
	company = data.get("company") or _active_company()
	if company not in _allowed_companies():
		frappe.throw(_("Not permitted"), frappe.PermissionError)
	rows = data.get("items") or []
	if not rows:
		frappe.throw(_("Add at least one opening-stock line."))
	doc = frappe.new_doc("Stock Reconciliation")
	doc.company, doc.posting_date = company, getdate(data.get("posting_date") or today())
	for row in rows:
		item = _catalog_item(row.get("item_code"))
		warehouse = row.get("warehouse")
		if item.controlled_item_type not in STOCK_TYPES or not cint(item.is_stock_item):
			frappe.throw(_("Opening stock accepts Stockable or Consumable catalog items only."))
		_validate_company_links(company, warehouse=warehouse)
		if not warehouse or flt(row.get("qty")) < 0 or flt(row.get("valuation_rate")) < 0:
			frappe.throw(_("Warehouse, quantity, and valuation rate must be non-negative."))
		doc.append("items", {"item_code": item.name, "warehouse": warehouse, "qty": flt(row.get("qty")),
			"valuation_rate": flt(row.get("valuation_rate"))})
	doc.flags.ignore_permissions = True
	doc.insert()
	doc.submit()
	return {"doctype": doc.doctype, "name": doc.name}


@frappe.whitelist()
def get_controlled_catalog_stock_reconciliation_items(
	warehouse: str,
	posting_date: str,
	posting_time: str,
	company: str,
	item_code: str | None = None,
	ignore_empty_stock: int | bool = 0,
) -> list:
	"""Warehouse stock rows limited to active controlled-catalogue stock Items."""
	_require(STOCK_ROLES | CATALOG_ROLES | OPENING_STOCK_ROLES)
	if company not in _allowed_companies():
		frappe.throw(_("Not permitted"), frappe.PermissionError)
	item_code = (item_code or "").strip() or None
	if item_code:
		meta = frappe.db.get_value(
			"Item", item_code,
			["controlled_procurement_catalog", "is_stock_item", "disabled"],
			as_dict=True,
		)
		if not meta or not cint(meta.controlled_procurement_catalog) or not cint(meta.is_stock_item) or cint(meta.disabled):
			frappe.throw(_("Select an active controlled catalogue stock item."))
	from erpnext.stock.doctype.stock_reconciliation.stock_reconciliation import get_items

	rows = get_items(
		warehouse, posting_date, posting_time, company,
		item_code=item_code, ignore_empty_stock=ignore_empty_stock,
	) or []
	return _filter_stock_reco_rows_to_catalog(rows)


def _filter_stock_reco_rows_to_catalog(rows: list) -> list:
	codes = {row.get("item_code") for row in rows if row.get("item_code")}
	if not codes:
		return []
	catalog = set(frappe.get_all(
		"Item",
		filters={"name": ["in", list(codes)], "controlled_procurement_catalog": 1, "disabled": 0, "is_stock_item": 1},
		pluck="name",
		limit_page_length=0,
	))
	return [row for row in rows if row.get("item_code") in catalog]


@frappe.whitelist(methods=["POST"])
def run_mrg_opening_stock_snapshot(company: str) -> dict:
	"""Run the approval-gated MRG opening-stock upload from the catalog workspace."""
	_require(OPENING_STOCK_ROLES)
	from construction_management.api import mrg_stock_snapshot

	if company != mrg_stock_snapshot.COMPANY or company not in _allowed_companies():
		frappe.throw(_("Not permitted"), frappe.PermissionError)
	return mrg_stock_snapshot.run()


# Master fields owned by Controlled Catalog Request. Stock ops may still save Item
# when only operational fields (valuation, rates, etc.) change.
CATALOG_ITEM_MASTER_FIELDS = (
	"item_name", "item_group", "stock_uom",
	"controlled_procurement_catalog", "controlled_item_type",
	"is_stock_item", "is_fixed_asset", "asset_category",
	"disabled", "is_purchase_item",
)
CATALOG_ITEM_MASTER_TABLES = ("supplier_items",)


def validate_item(doc, method=None) -> None:
	api = getattr(doc.flags, "controlled_procurement_api", False)
	roles = _roles()
	if doc.is_new() and not roles.intersection(ADMIN_ROLES) and not api:
		frappe.throw(_("Only CEO and System Manager can create items."), frappe.PermissionError)
	# System Manager / CEO may fix catalog Items in Desk; everyone else uses catalog requests.
	if _catalog_item_master_guard(doc) and not api and not roles.intersection(ADMIN_ROLES):
		frappe.throw(_("Create controlled catalog items through a catalog request."), frappe.PermissionError)
	if cint(doc.get("controlled_procurement_catalog")) and not frappe.db.exists("UOM", doc.stock_uom):
		frappe.throw(_("Controlled catalog items require a valid stock UOM."))


def _catalog_item_master_guard(doc) -> bool:
	"""True when this save must go through a Controlled Catalog Request."""
	if not cint(doc.get("controlled_procurement_catalog")) and not _was_controlled_catalog_item(doc):
		return False
	if doc.is_new():
		return True
	return _catalog_master_fields_changed(doc)


def _was_controlled_catalog_item(doc) -> bool:
	before = doc.get_doc_before_save()
	return bool(before and cint(before.get("controlled_procurement_catalog")))


def _catalog_master_fields_changed(doc) -> bool:
	for field in CATALOG_ITEM_MASTER_FIELDS:
		if doc.meta.has_field(field) and doc.has_value_changed(field):
			return True
	for field in CATALOG_ITEM_MASTER_TABLES:
		if doc.meta.has_field(field) and doc.has_value_changed(field):
			return True
	return False


def _workspace_only() -> bool:
	"""Site config `controlled_procurement_workspace_only: 1` forces new POs, receipts and transfers
	through the Procurement app. Off by default, so Desk creation works as usual."""
	return bool(cint(frappe.conf.get("controlled_procurement_workspace_only")))


def validate_transaction(doc, method=None) -> None:
	from_workspace = getattr(doc.flags, "controlled_procurement_api", False)
	if doc.is_new() and not from_workspace:
		if _workspace_only() and not _roles().intersection(ADMIN_ROLES):
			frappe.throw(_("Create procurement documents from the Controlled Procurement workspace."), frappe.PermissionError)
		# Desk duplicates and mapped documents (e.g. a receipt made from a controlled PO) copy the
		# flag; a document created in Desk is a normal one, not a catalog-controlled one
		doc.controlled_procurement = 0
	if not cint(doc.get("controlled_procurement")):
		return
	if not getattr(doc.flags, "controlled_procurement_api", False) and not _roles().intersection(ADMIN_ROLES):
		frappe.throw(_("Edit or submit this document from the Controlled Procurement workspace."), frappe.PermissionError)
	transfer = doc.doctype == "Stock Entry" and doc.purpose == "Material Transfer"
	if transfer:
		_validate_transfer_warehouses(doc.company, doc.from_warehouse, doc.to_warehouse)
	else:
		_validate_allocation({field: doc.get(field) for field in ALLOCATION_FIELDS}, doc.company)
	for row in doc.get("items") or []:
		item = _catalog_item(row.item_code)
		if row.uom != item.stock_uom or row.stock_uom != item.stock_uom:
			frappe.throw(_("Controlled procurement item {0} must use its stock UOM {1}.").format(row.item_code, item.stock_uom))
		_set_row_type(row, item)
		if doc.doctype == "Purchase Order" and item.controlled_item_type == SERVICE:
			_ensure_service_item_default_for_company(item.name, doc.company)
		if doc.doctype == "Stock Entry" and (item.controlled_item_type not in STOCK_TYPES or not cint(item.is_stock_item)):
			frappe.throw(_("Material Transfer accepts Stockable or Consumable controlled items only."))
		if transfer:
			if row.s_warehouse != doc.from_warehouse or row.t_warehouse != doc.to_warehouse:
				frappe.throw(_("Every transfer item must use the header From and To warehouses."))
		else:
			_validate_allocation({field: row.get(field) for field in ALLOCATION_FIELDS}, doc.company)
	if doc.doctype == "Purchase Receipt":
		_validate_receipt_allocations(doc)


def _validate_receipt_allocations(doc) -> None:
	for row in doc.get("items") or []:
		if not row.purchase_order_item:
			frappe.throw(_("Controlled Purchase Receipt items must come from a Purchase Order."))
		po_values = frappe.db.exists("Purchase Order Item", row.purchase_order_item)
		if not po_values:
			frappe.throw(_("Controlled Purchase Receipt items must come from a Purchase Order."))
		# allocation may be re-pointed at receipt time; it only has to be a valid Project/Bill/BOQ
		# set (checked by the caller) inside the receipt's company
		_validate_company_links(doc.company, project=row.get("project"))


def _document_list_fields(doctype: str) -> list[str]:
	fields = ["name", "project", "bill_no", "boq_item", "grand_total", "status", "docstatus", "modified"]
	if doctype == "Purchase Order":
		return fields + ["supplier", "supplier_name", "transaction_date", "schedule_date", "per_received", "per_billed", "controlled_procurement", "custom_lpo_type", "custom_is_provisional_po", "custom_lpo_reference", "custom_lpo_approval_status"]
	if doctype == "Purchase Receipt":
		return fields + ["supplier", "supplier_name", "posting_date", "supplier_delivery_note", "controlled_procurement"]
	if doctype == "Purchase Invoice":
		return fields + ["supplier", "supplier_name", "posting_date", "due_date", "outstanding_amount", "custom_supplier_invoice_no", "is_paid"]
	return fields + ["posting_date", "purpose", "controlled_procurement", "total_outgoing_value"]


def _annotate_po_list_catalog_flags(rows: list) -> None:
	"""Mark list rows that include at least one controlled-catalog item line."""
	for row in rows:
		row["has_catalog_items"] = 0
	names = [row.name for row in rows if row.get("name")]
	if not names or not frappe.db.has_column("Purchase Order Item", "controlled_procurement_catalog"):
		return
	catalog_parents = {
		parent
		for parent in frappe.get_all(
			"Purchase Order Item",
			filters={"parent": ["in", names], "controlled_procurement_catalog": 1},
			pluck="parent",
			distinct=True,
			limit_page_length=0,
		)
	}
	for row in rows:
		if row.name in catalog_parents:
			row["has_catalog_items"] = 1


def _workspace_filters(doctype: str, filters: dict) -> dict:
	if doctype in ("Purchase Order", "Purchase Receipt"):
		return {**filters, "is_subcontracted": 0}
	if doctype == "Purchase Invoice":
		return {**filters, "is_subcontracted": 0, "is_return": 0}
	return filters


def _has_open_po_field() -> bool:
	return frappe.get_meta("Purchase Order").has_field("custom_is_provisional_po")


def _is_open_lpo(po) -> bool:
	from construction_management.api.lpo_workflow import lpo_type

	return lpo_type(po) == "Open"


def _validate_receivable_po(po, open_lpo: bool) -> None:
	if po.docstatus != 1 or po.status in {"Closed", "On Hold"}:
		frappe.throw(_("Select a submitted Purchase Order that is not closed or on hold."))
	if not open_lpo and (po.status in {"Completed", "Delivered"} or flt(po.per_received) >= 100):
		frappe.throw(_("This Purchase Order has no quantity left to receive."))


DOCUMENT_ROLES = {
	"Purchase Order": PURCHASE_ROLES,
	"Purchase Receipt": PURCHASE_ROLES,
	"Purchase Invoice": PURCHASE_ROLES,
	"Stock Entry": STOCK_ROLES,
}
CLOSED_PO_STATUSES = ("Closed", "Completed", "Delivered", "On Hold")


def _in_workspace(doc) -> bool:
	"""Every non-subcontracted PO is visible; receipts and transfers only when created here."""
	if doc.doctype in ("Purchase Order", "Purchase Receipt"):
		return not cint(doc.get("is_subcontracted"))
	if doc.doctype == "Purchase Invoice":
		return not cint(doc.get("is_subcontracted")) and not cint(doc.get("is_return"))
	return doc.get("purpose") == "Material Transfer"


def _controlled_doc(doctype: str, name: str, ptype: str = "read"):
	if doctype not in DOCUMENT_ROLES:
		frappe.throw(_("Unsupported controlled document type."))
	accounts_approval_view = doctype == "Purchase Order" and ptype == "read" and bool(_roles().intersection(ACCOUNT_ROLES | ADMIN_ROLES))
	if not accounts_approval_view:
		_require(DOCUMENT_ROLES[doctype])
	doc = frappe.get_doc(doctype, name)
	if doc.company != _document_company(None):
		frappe.throw(_("Switch to the document's company in Procurement before opening it."), frappe.PermissionError)
	if accounts_approval_view and not _roles().intersection(PURCHASE_ROLES):
		if doc.company not in _allowed_companies() or not cint(doc.get("controlled_procurement")) or doc.get("custom_lpo_type") != "Manual":
			frappe.throw(_("Not permitted"), frappe.PermissionError)
	else:
		frappe.has_permission(doctype, ptype, doc=doc, throw=True)
	if not _in_workspace(doc):
		frappe.throw(_("This document is not part of Controlled Procurement."))
	return doc


@frappe.whitelist(methods=["GET"])
def get_open_lpos(search: str = "", supplier: str = "", project: str = "", start: int = 0, page_length: int = 50, controlled_only: int = 0, open_po_only: int = 0) -> dict:
	"""Submitted receivable LPOs, including fully received Open LPOs."""
	_require(PURCHASE_ROLES)
	provisional = "po.custom_is_provisional_po" if _has_open_po_field() else "0"
	conditions = [
		"po.is_subcontracted = 0", "po.docstatus = 1",
		"po.status not in %(closed)s",
		f"({provisional} = 1 or (po.per_received < 100 and po.status not in %(completed)s))",
	]
	if cint(open_po_only):
		# the Open LPOs tab lists only POs ticked "Provisional / Open PO"
		if not _has_open_po_field():
			return {"rows": [], "total_count": 0, "summary": {"open_lpos": 0, "pending_value": 0, "overdue": 0}}
		conditions.append("po.custom_is_provisional_po = 1")
	if cint(controlled_only):
		# the workspace receipt form can only receive catalog-controlled orders
		conditions.append("po.controlled_procurement = 1")
	values = {"closed": ("Closed", "On Hold"), "completed": ("Completed", "Delivered"), "today": frappe.utils.today()}
	company = _document_company(None)
	conditions.append("po.company = %(company)s")
	values["company"] = company
	if supplier:
		conditions.append("po.supplier = %(supplier)s")
		values["supplier"] = supplier
	if project:
		conditions.append("po.project = %(project)s")
		values["project"] = project
	if search:
		conditions.append("(po.name like %(search)s or po.supplier_name like %(search)s or po.project like %(search)s or po.custom_lpo_reference like %(search)s)")
		values["search"] = f"%{search.strip()}%"
	# only list POs the user can read, then aggregate pending lines in SQL
	permitted = frappe.get_list(
		"Purchase Order", filters=_scoped({"is_subcontracted": 0, "docstatus": 1}),
		pluck="name", limit_page_length=0,
	)
	if not permitted:
		return {"rows": [], "total_count": 0, "summary": {"open_lpos": 0, "pending_value": 0, "overdue": 0}}
	conditions.append("po.name in %(permitted)s")
	values["permitted"] = tuple(permitted)
	where = " and ".join(conditions)
	# "Provisional / Open PO" tick (redtra_customisation); receipts may exceed the ordered qty
	rows = frappe.db.sql(
		f"""
		select po.name, po.supplier, po.supplier_name, po.project, po.transaction_date, po.schedule_date, po.custom_lpo_reference,
			po.grand_total, po.per_received, po.per_billed, po.status, po.currency, po.controlled_procurement,
			{provisional} as is_open_po,
			sum(greatest(poi.qty - poi.received_qty, 0)) as pending_qty,
			sum(greatest(poi.qty - poi.received_qty, 0) * poi.rate) as pending_value,
			count(poi.name) as line_count,
			(po.schedule_date < %(today)s and sum(greatest(poi.qty - poi.received_qty, 0)) > 0) as is_overdue,
			datediff(%(today)s, po.transaction_date) as age_days
		from `tabPurchase Order` po
		join `tabPurchase Order Item` poi on poi.parent = po.name and ifnull(poi.closed, 0) = 0 and ifnull(poi.delivered_by_supplier, 0) = 0
		where {where}
		group by po.name
		having pending_qty > 0 or is_open_po = 1
		order by po.schedule_date asc, po.name asc
		""",
		values,
		as_dict=True,
	)
	summary = {
		"open_lpos": len(rows),
		"pending_value": sum(flt(row.pending_value) for row in rows),
		"overdue": sum(1 for row in rows if cint(row.is_overdue)),
	}
	start = max(cint(start), 0)
	page_length = min(max(cint(page_length) or 50, 1), 200)
	return {"rows": rows[start : start + page_length], "total_count": len(rows), "summary": summary}


@frappe.whitelist(methods=["GET"])
def get_document_activity(doctype: str, name: str) -> list[dict]:
	"""Comments, lifecycle events, emails, and readable tracked edits for this revision."""
	doc = _controlled_doc(doctype, name)
	comments = frappe.get_all(
		"Comment",
		filters={"reference_doctype": doctype, "reference_name": name, "comment_type": ["in", ["Comment", "Info", "Edit", "Submitted", "Cancelled"]]},
		fields=["name", "content", "owner", "comment_email", "comment_type", "creation"],
	)
	versions = frappe.get_all("Version", filters={"ref_doctype": doctype, "docname": name},
		fields=["name", "data", "owner", "creation"], order_by="creation desc", limit_page_length=100)
	emails = frappe.get_all(
		"Communication",
		filters={"reference_doctype": doctype, "reference_name": name, "communication_medium": "Email"},
		fields=["name", "subject", "content", "sender", "recipients", "cc", "sent_or_received", "delivery_status", "has_attachment", "creation"],
	)
	queues = frappe.get_all("Email Queue", filters={"communication": ["in", [row.name for row in emails]]},
		fields=["communication", "status", "error", "creation"], order_by="creation desc") if emails else []
	queue_by_communication = {row.communication: row for row in reversed(queues)}
	users = {row.owner for row in comments} | {row.sender for row in emails} | {row.owner for row in versions} | {doc.owner}
	names = dict(frappe.get_all("User", filters={"name": ["in", list(users)]}, fields=["name", "full_name"], as_list=True)) if users else {}
	activity = [{"type": "event", "name": f"created-{name}", "content": _("Document created"),
		"by": names.get(doc.owner) or doc.owner, "creation": doc.creation}] + [
		{"type": "comment" if row.comment_type == "Comment" else "event", "name": row.name,
			"content": row.content, "by": names.get(row.owner) or row.owner, "creation": row.creation}
		for row in comments
	] + [
		{
			"type": "email", "name": row.name, "subject": row.subject, "content": row.content,
			"by": names.get(row.sender) or row.sender, "recipients": row.recipients, "cc": row.cc,
			"direction": row.sent_or_received,
			"status": _email_delivery_status(row, queue_by_communication.get(row.name)),
			"delivery_error": (queue_by_communication.get(row.name).error or "")[-400:]
				if queue_by_communication.get(row.name) and queue_by_communication.get(row.name).status == "Error" else "",
			"has_attachment": row.has_attachment, "creation": row.creation,
		}
		for row in emails
	] + [
		{"type": "event" if _version_status_change(row.data) else "edit", "name": row.name,
			"changes": _version_changes(doctype, row.data),
			"by": names.get(row.owner) or row.owner, "creation": row.creation}
		for row in versions if _version_changes(doctype, row.data)
	]
	return sorted(activity, key=lambda row: row["creation"], reverse=True)


def _email_delivery_status(communication, queue=None) -> str:
	if communication.sent_or_received == "Received":
		return communication.delivery_status or "Received"
	if queue:
		return "Queued" if queue.status in {"Not Sent", "Sending"} else queue.status
	return communication.delivery_status or "Queued"


def _version_status_change(raw: str) -> bool:
	try:
		return any(change[0] == "docstatus" for change in json.loads(raw or "{}").get("changed") or [])
	except (TypeError, ValueError):
		return False


AUDIT_FIELDS = {
	"supplier", "transaction_date", "schedule_date", "posting_date", "supplier_delivery_note",
	"project", "bill_no", "boq_item", "set_warehouse", "from_warehouse", "to_warehouse",
	"custom_lpo_type", "custom_lpo_reference", "custom_lpo_approval_status", "taxes_and_charges",
	"remarks", "items", "item_code", "qty", "rate", "warehouse", "s_warehouse", "t_warehouse", "description",
}


def _version_changes(doctype: str, raw: str) -> list[str]:
	try:
		data = json.loads(raw or "{}")
	except (ValueError, TypeError):
		return []
	meta = frappe.get_meta(doctype)
	changes = []
	for field, old, new in data.get("changed") or []:
		if field == "docstatus":
			changes.append({1: "Submitted", 2: "Cancelled", 0: "Returned to draft"}.get(cint(new), "Status changed"))
			continue
		if field in AUDIT_FIELDS and old != new:
			label = meta.get_label(field) or field.replace("_", " ").title()
			changes.append(f"{label}: {old or '—'} → {new or '—'}")
	for table, row in data.get("added") or []:
		if table == "items":
			changes.append(f"Added item {row.get('item_code', '—')} (qty {row.get('qty', '—')})")
	for table, row in data.get("removed") or []:
		if table == "items":
			changes.append(f"Removed item {row.get('item_code', '—')}")
	for table, row_name, _, fields in data.get("row_changed") or []:
		if table == "items":
			for field, old, new in fields:
				if field in AUDIT_FIELDS and old != new:
					changes.append(f"Item {row_name}: {field.replace('_', ' ').title()} {old or '—'} → {new or '—'}")
	return changes[:30]


def _eligible_tag_user(doc, user: str) -> bool:
	if not frappe.db.exists("User", {"name": user, "enabled": 1}) or not frappe.db.exists("Raven User", {"user": user, "enabled": 1}):
		return False
	roles = set(frappe.get_roles(user))
	special_approval_view = (doc.doctype == "Purchase Order" and bool(roles.intersection(ACCOUNT_ROLES | ADMIN_ROLES))
		and cint(doc.get("controlled_procurement")) and doc.get("custom_lpo_type") == "Manual")
	allowed_role = special_approval_view or bool(roles.intersection(DOCUMENT_ROLES[doc.doctype]))
	if not allowed_role or (not special_approval_view and not frappe.has_permission(doc.doctype, "read", doc=doc, user=user)):
		return False
	company_permissions = frappe.get_all("User Permission", filters={"user": user, "allow": "Company"}, pluck="for_value")
	return not company_permissions or doc.company in company_permissions


@frappe.whitelist(methods=["GET"])
def get_comment_tag_users(doctype: str, name: str, search: str = "") -> list[dict]:
	doc = _controlled_doc(doctype, name)
	users = frappe.get_all("Raven User", filters={"enabled": 1, "type": "User"}, fields=["user", "full_name"], limit_page_length=0)
	needle = search.strip().casefold()
	return [{"user": row.user, "full_name": row.full_name} for row in users
		if row.user and (not needle or needle in row.user.casefold() or needle in (row.full_name or "").casefold())
		and _eligible_tag_user(doc, row.user)][:20]


@frappe.whitelist(methods=["POST"])
def add_document_comment(doctype: str, name: str, content: str, tagged_users: str | list | None = None) -> dict:
	doc = _controlled_doc(doctype, name)
	if not (content or "").strip():
		frappe.throw(_("Comment cannot be empty."))
	tags = json.loads(tagged_users) if isinstance(tagged_users, str) else (tagged_users or [])
	if not isinstance(tags, list) or len(tags) > 20 or len(tags) != len(set(tags)):
		frappe.throw(_("Invalid tagged users."))
	for user in tags:
		if not isinstance(user, str) or not _eligible_tag_user(doc, user) or f"@{user}" not in content:
			frappe.throw(_("{0} cannot be tagged on this document.").format(user), frappe.PermissionError)
	if tags and not frappe.db.exists("Raven Bot", "procurement bot"):
		frappe.throw(_("The procurement Raven bot is not configured; tagged comment was not posted."))
	bot = frappe.get_doc("Raven Bot", "procurement bot") if tags else None
	comment = frappe.get_doc({
		"doctype": "Comment", "comment_type": "Comment", "reference_doctype": doctype, "reference_name": name,
		"comment_email": frappe.session.user, "comment_by": frappe.utils.get_fullname(frappe.session.user),
		"content": frappe.utils.sanitize_html(content),
	})
	comment.insert(ignore_permissions=True)
	for user in tags:
		try:
			message = bot.send_direct_message(user, text=f"{frappe.utils.get_fullname(frappe.session.user)} mentioned you on {doctype} {name}: {frappe.utils.strip_html(content)[:240]}", link_doctype=doctype, link_document=name)
			if not message:
				frappe.throw(_("Raven did not deliver the notification."))
		except Exception as exc:
			frappe.throw(_("Could not notify {0} through Raven: {1}").format(user, str(exc)))
	return {"name": comment.name}


@frappe.whitelist(methods=["GET", "POST"])
def search_user_emails(txt: str = "") -> list[dict]:
	"""Suggest enabled System Users by name or email for To/Cc autocomplete."""
	if frappe.session.user in {"", "Guest"}:
		frappe.throw(_("Not permitted"), frappe.PermissionError)
	txt = (txt or "").strip()
	if len(txt) < 1:
		return []
	like = f"%{txt}%"
	rows = frappe.db.sql(
		"""
		select name, full_name, email
		from `tabUser`
		where enabled = 1
			and user_type = 'System User'
			and name not in ('Administrator', 'Guest')
			and ifnull(email, '') != ''
			and (full_name like %(like)s or name like %(like)s or email like %(like)s)
		order by
			case when full_name like %(prefix)s then 0 when email like %(prefix)s then 1 else 2 end,
			full_name, email
		limit 15
		""",
		{"like": like, "prefix": f"{txt}%"},
		as_dict=True,
	)
	return [{"email": row.email, "full_name": row.full_name or row.name, "name": row.name} for row in rows]


@frappe.whitelist(methods=["GET"])
def get_email_draft(doctype: str, name: str) -> dict:
	"""Default recipients, subject, and print formats for the email composer."""
	doc = _controlled_doc(doctype, name)
	frappe.has_permission(doctype, "email", doc=doc, throw=True)
	from construction_management.api.catalog_email import outgoing_sender

	# the document's contact person first, then the supplier's own email
	recipients = doc.get("contact_email") or ""
	if not recipients and doc.get("supplier"):
		recipients = frappe.db.get_value("Supplier", doc.supplier, "email_id") or ""
	meta = frappe.get_meta(doctype)
	print_formats = ["Standard"] + frappe.get_all(
		"Print Format", filters={"doc_type": doctype, "disabled": 0}, pluck="name", order_by="name asc"
	)
	return {
		"sender": outgoing_sender(doc.company),
		"recipients": recipients,
		"subject": f"{doctype}: {doc.name}",
		"message": _("Dear {0},<br><br>Please find attached {1} {2}.<br><br>Regards,<br>{3}").format(
			(doc.get("supplier_name") or doc.get("supplier") or _("Sir/Madam")).strip(), doctype, doc.name,
			frappe.utils.get_fullname(frappe.session.user),
		),
		"print_formats": print_formats,
		"default_print_format": meta.default_print_format or "Standard",
		"email_templates": frappe.get_all(
			"Email Template",
			or_filters=[["reference_doctype", "=", doctype], ["reference_doctype", "is", "not set"], ["reference_doctype", "=", ""]],
			pluck="name", order_by="name asc",
		),
		"can_email": True,
	}


CONTACT_FIELDS = ["name", "first_name", "last_name", "full_name", "designation", "email_id", "phone", "mobile_no", "is_primary_contact"]


@frappe.whitelist(methods=["GET"])
def get_supplier_contacts(supplier: str) -> list[dict]:
	"""Contacts linked to a supplier, primary first."""
	_require(PURCHASE_ROLES)
	frappe.has_permission("Supplier", "read", doc=supplier, throw=True)
	names = frappe.get_all(
		"Dynamic Link", filters={"link_doctype": "Supplier", "link_name": supplier, "parenttype": "Contact"}, pluck="parent"
	)
	if not names:
		return []
	return frappe.get_all(
		"Contact", filters={"name": ["in", names]}, fields=CONTACT_FIELDS,
		order_by="is_primary_contact desc, full_name asc",
	)


@frappe.whitelist(methods=["GET"])
def get_supplier_payment_terms(supplier: str) -> dict:
	"""Read the default payment terms for a supplier visible to this user."""
	_require(PURCHASE_ROLES)
	if not frappe.db.exists("Supplier", supplier):
		frappe.throw(_("Supplier {0} does not exist.").format(supplier))
	frappe.has_permission("Supplier", "read", doc=supplier, throw=True)
	return {"payment_terms": frappe.db.get_value("Supplier", supplier, "payment_terms") or ""}


@frappe.whitelist(methods=["POST"])
def set_supplier_payment_terms(supplier: str, payment_terms: str = "") -> dict:
	"""Update only the supplier's payment-terms default from Procurement."""
	_require(PURCHASE_ROLES)
	doc = frappe.get_doc("Supplier", supplier)
	doc.payment_terms = _payment_terms_template(payment_terms)
	doc.save(ignore_permissions=True)
	return {"name": doc.name, "payment_terms": doc.payment_terms or ""}


@frappe.whitelist(methods=["GET"])
def get_procurement_suppliers(search: str = "") -> list[dict]:
	"""Supplier directory for purchasing users, independent of an LPO form."""
	_require(PURCHASE_ROLES)
	filters = {"disabled": 0}
	search = (search or "").strip()
	return frappe.get_all(
		"Supplier",
		filters=filters,
		or_filters=[[field, "like", f"%{search}%"] for field in ("name", "supplier_name", "tax_id")] if search else None,
		fields=["name", "supplier_name", "supplier_group", "supplier_type", "tax_id", "payment_terms"],
		order_by="supplier_name asc",
		limit_page_length=100,
	)


@frappe.whitelist(methods=["GET"])
def get_supplier_profile(supplier: str) -> dict:
	_require(PURCHASE_ROLES)
	doc = frappe.get_doc("Supplier", supplier)
	frappe.has_permission("Supplier", "read", doc=doc, throw=True)
	profile = {
		"supplier_name": doc.supplier_name or "", "supplier_group": doc.supplier_group or "",
		"supplier_type": doc.supplier_type or "Company", "tax_id": doc.tax_id or "",
		"trade_licence_no": doc.get("custom_trade_licence_no") or "",
		"po_box_no": doc.get("custom_po_box_no") or "",
		"payment_terms": doc.payment_terms or "",
	}
	contact_name = _supplier_linked_record(supplier, "Contact", doc.supplier_primary_contact)
	if contact_name:
		contact = frappe.get_doc("Contact", contact_name)
		profile.update({"contact_first_name": contact.first_name or "", "contact_last_name": contact.last_name or "",
			"contact_email": contact.email_id or "",
			"contact_mobile": contact.mobile_no or contact.phone or ""})
	address_name = _supplier_linked_record(supplier, "Address", doc.supplier_primary_address)
	if address_name:
		address = frappe.get_doc("Address", address_name)
		profile.update({"address_line1": address.address_line1 or "", "address_line2": address.address_line2 or "",
			"city": address.city or "", "state": address.state or "", "pincode": address.pincode or "",
			"country": address.country or ""})
	return profile


def _supplier_linked_record(supplier: str, parenttype: str, primary: str = "") -> str:
	filters = {"link_doctype": "Supplier", "link_name": supplier, "parenttype": parenttype}
	if primary and frappe.db.exists("Dynamic Link", {**filters, "parent": primary}):
		return primary
	return frappe.db.get_value("Dynamic Link", filters, "parent") or ""


def _upsert_supplier_contact(supplier: str, data: dict) -> None:
	first_name = (data.get("contact_first_name") or data.get("contact_name") or "").strip()
	last_name = (data.get("contact_last_name") or "").strip()
	email, mobile = (data.get("contact_email") or "").strip(), (data.get("contact_mobile") or "").strip()
	if not any((first_name, last_name, email, mobile)):
		return
	contact_name = _supplier_linked_record(
		supplier, "Contact", frappe.db.get_value("Supplier", supplier, "supplier_primary_contact"))
	contact = frappe.get_doc("Contact", contact_name) if contact_name else frappe.new_doc("Contact")
	if first_name:
		contact.first_name = first_name
		contact.last_name = last_name
	elif contact.is_new():
		frappe.throw(_("Primary contact first name is required when adding contact details."))
	_set_primary_child(contact, "email_ids", "email_id", "is_primary", email)
	_set_primary_child(contact, "phone_nos", "phone", "is_primary_mobile_no", mobile)
	if not any(row.link_doctype == "Supplier" and row.link_name == supplier for row in contact.links):
		contact.append("links", {"link_doctype": "Supplier", "link_name": supplier})
	if contact.is_new():
		contact.is_primary_contact = 1
		contact.insert(ignore_permissions=True)
	else:
		contact.save(ignore_permissions=True)
	frappe.db.set_value("Supplier", supplier, "supplier_primary_contact", contact.name)


def _upsert_supplier_address(supplier: str, data: dict) -> None:
	fields = {field: (data.get(field) or "").strip() for field in
		("address_line1", "address_line2", "city", "state", "pincode", "country")}
	if not any(fields.values()):
		return
	if not all(fields[field] for field in ("address_line1", "city", "country")):
		frappe.throw(_("Address Line 1, City, and Country are required when adding an address."))
	address_name = _supplier_linked_record(
		supplier, "Address", frappe.db.get_value("Supplier", supplier, "supplier_primary_address"))
	address = frappe.get_doc("Address", address_name) if address_name else frappe.new_doc("Address")
	address.update({"address_title": frappe.db.get_value("Supplier", supplier, "supplier_name"),
		"address_type": "Billing", **fields})
	if not any(row.link_doctype == "Supplier" and row.link_name == supplier for row in address.links):
		address.append("links", {"link_doctype": "Supplier", "link_name": supplier})
	if address.is_new():
		address.is_primary_address = 1
		address.insert(ignore_permissions=True)
	else:
		address.save(ignore_permissions=True)
	frappe.db.set_value("Supplier", supplier, "supplier_primary_address", address.name)


@frappe.whitelist(methods=["POST"])
def update_supplier_profile(supplier: str, data: str | dict) -> dict:
	"""Maintain the purchasing supplier profile, contact and primary address."""
	_require(PURCHASE_ROLES)
	data = _as_dict(data)
	doc = frappe.get_doc("Supplier", supplier)
	group = data.get("supplier_group") or doc.supplier_group
	if group and not frappe.db.exists("Supplier Group", group):
		frappe.throw(_("Supplier Group {0} does not exist.").format(group))
	supplier_type = data.get("supplier_type") or doc.supplier_type
	if supplier_type not in SUPPLIER_TYPES:
		frappe.throw(_("Invalid supplier type."))
	doc.update({"supplier_name": (data.get("supplier_name") or doc.supplier_name).strip(), "supplier_group": group,
		"supplier_type": supplier_type, "tax_id": (data.get("tax_id") or "").strip(),
		"custom_trade_licence_no": (data.get("trade_licence_no") or "").strip(),
		"custom_po_box_no": (data.get("po_box_no") or "").strip(),
		"payment_terms": _payment_terms_template(data.get("payment_terms"))})
	doc.save(ignore_permissions=True)
	_upsert_supplier_contact(doc.name, data)
	_upsert_supplier_address(doc.name, data)
	return {"name": doc.name, "supplier_name": doc.supplier_name, "payment_terms": doc.payment_terms or ""}


@frappe.whitelist(methods=["GET"])
def get_procurement_terms_templates() -> list[dict]:
	_require(PURCHASE_ROLES)
	return frappe.get_all(
		"Terms and Conditions", filters={"buying": 1}, fields=["name", "title", "terms", "disabled"],
		order_by="title asc", limit_page_length=0,
	)


@frappe.whitelist(methods=["POST"])
def save_procurement_terms_template(name: str = "", data: str | dict | None = None) -> dict:
	"""Create or edit a buying terms template available to Procurement LPOs."""
	_require(PURCHASE_ROLES)
	payload = _as_dict(data or {})
	title = (payload.get("title") or "").strip()
	if not title:
		frappe.throw(_("Terms template title is required."))
	doc = frappe.get_doc("Terms and Conditions", name) if name else frappe.new_doc("Terms and Conditions")
	doc.update({"title": title, "terms": payload.get("terms") or "", "buying": 1, "selling": 0, "disabled": cint(payload.get("disabled"))})
	if doc.is_new():
		doc.insert(ignore_permissions=True)
	else:
		doc.save(ignore_permissions=True)
	return {"name": doc.name, "title": doc.title, "terms": doc.terms, "disabled": doc.disabled}


@frappe.whitelist(methods=["GET"])
def get_catalog_buying_price(item_code: str, supplier: str = "", company: str = "", order_date: str = "") -> dict:
	"""Return an applicable supplier price, then a general buying price, in company currency."""
	_require(PURCHASE_ROLES)
	company = _document_company(company)
	_catalog_item(item_code)
	if not frappe.db.get_value("Item", item_code, "is_purchase_item") or frappe.db.get_value("Item", item_code, "disabled"):
		frappe.throw(_("Item {0} is not an active buying item.").format(item_code))
	if supplier:
		if not frappe.db.exists("Supplier", supplier):
			frappe.throw(_("Supplier {0} does not exist.").format(supplier))
		frappe.has_permission("Supplier", "read", doc=supplier, throw=True)
	price_list = frappe.db.get_value("Supplier", supplier, "default_price_list") if supplier else None
	price_lists = list(dict.fromkeys([price_list, "Standard Buying"])) if price_list else ["Standard Buying"]
	rows = frappe.db.sql(
		"""select ip.name, ip.price_list_rate, ip.price_list, ip.supplier
		from `tabItem Price` ip
		join `tabPrice List` pl on pl.name = ip.price_list
		where ip.item_code = %(item_code)s and ip.buying = 1
		and ip.price_list in %(price_lists)s and pl.enabled = 1 and pl.buying = 1
		and pl.currency = %(currency)s
		and (ifnull(ip.supplier, '') = '' or ip.supplier = %(supplier)s)
		and (ifnull(ip.uom, '') = '' or ip.uom = %(uom)s)
		and (ip.valid_from is null or ip.valid_from <= %(order_date)s)
		and (ip.valid_upto is null or ip.valid_upto >= %(order_date)s)
		order by (ip.supplier = %(supplier)s and %(supplier)s != '') desc,
		(ip.price_list = %(preferred_list)s) desc, ip.valid_from desc, ip.creation desc
		limit 1""",
		{
			"item_code": item_code, "supplier": supplier or "", "price_lists": price_lists,
			"uom": frappe.db.get_value("Item", item_code, "stock_uom"),
			"currency": frappe.get_cached_value("Company", company, "default_currency"),
			"order_date": str(getdate(order_date or today())), "preferred_list": price_list or "Standard Buying",
		}, as_dict=True,
	)
	if not rows:
		return {"rate": None, "price_list": "", "item_price": "", "supplier": ""}
	return {"rate": flt(rows[0].price_list_rate), "price_list": rows[0].price_list,
		"item_price": rows[0].name, "supplier": rows[0].supplier or ""}


# Keep the Procurement API prefix stable for the Vue workspace.
from construction_management.api.controlled_price_requests import (
	approve_price_change,
	bulk_decide_price_requests,
	get_price_context,
	get_price_request,
	get_price_requests,
	reject_price_change,
	request_price_change,
	retry_catalog_price_notice,
	retry_price_notice,
)
from construction_management.api.controlled_price_history import get_price_history
from construction_management.api.po_price_approvals import (
	approve_accounts_po_price,
	approve_po_price,
	deliver_po_price_notice,
	get_po_price_requests,
	reject_accounts_po_price,
	reject_po_price,
	retry_po_price_notice,
)


@frappe.whitelist(methods=["POST"])
def create_supplier(data: str | dict) -> dict:
	"""Create a supplier from the controlled procurement workspace."""
	_require(PURCHASE_ROLES)
	data = _as_dict(data)
	supplier_name = (data.get("supplier_name") or "").strip()
	if not supplier_name:
		frappe.throw(_("Supplier name is required."))
	if frappe.db.exists("Supplier", {"supplier_name": supplier_name}):
		frappe.throw(_("Supplier {0} already exists.").format(supplier_name))

	supplier_group = data.get("supplier_group") or None
	if supplier_group and not frappe.db.exists("Supplier Group", supplier_group):
		frappe.throw(_("Supplier Group {0} does not exist.").format(supplier_group))

	supplier_type = data.get("supplier_type") or "Company"
	if supplier_type not in SUPPLIER_TYPES:
		frappe.throw(_("Invalid supplier type."))

	supplier = frappe.new_doc("Supplier")
	supplier.update({
		"supplier_name": supplier_name,
		"supplier_group": supplier_group,
		"supplier_type": supplier_type,
		"tax_id": (data.get("tax_id") or "").strip(),
		"custom_trade_licence_no": (data.get("trade_licence_no") or "").strip(),
		"custom_po_box_no": (data.get("po_box_no") or "").strip(),
		"payment_terms": _payment_terms_template(data.get("payment_terms")),
	})
	_set_supplier_naming_series(supplier)
	supplier.insert(ignore_permissions=True)
	_upsert_supplier_contact(supplier.name, data)
	_upsert_supplier_address(supplier.name, data)
	return {"name": supplier.name, "supplier_name": supplier.supplier_name, "payment_terms": supplier.payment_terms or ""}


def _set_supplier_naming_series(supplier) -> None:
	"""Provide a valid series only when Supplier naming is series-based."""
	if frappe.defaults.get_global_default("supp_master_name") != "Naming Series":
		return
	series = (supplier.meta.get_field("naming_series").options or "").splitlines()
	if series:
		supplier.naming_series = series[0]


def _set_primary_child(contact, table: str, value_field: str, primary_field: str, value: str) -> None:
	"""Make `value` the primary entry of a Contact child table, adding it if missing."""
	value = (value or "").strip()
	if not value:
		return
	found = False
	for row in contact.get(table):
		is_match = row.get(value_field) == value
		row.set(primary_field, 1 if is_match else 0)
		found = found or is_match
	if not found:
		contact.append(table, {value_field: value, primary_field: 1})


@frappe.whitelist(methods=["POST"])
def save_supplier_contact(supplier: str, data: str | dict) -> dict:
	"""Create a supplier contact, or correct an existing one's name, email and numbers."""
	_require(PURCHASE_ROLES)
	data = _as_dict(data)
	if not (data.get("first_name") or "").strip():
		frappe.throw(_("First name is required."))
	if data.get("name"):
		contact = frappe.get_doc("Contact", data["name"])
		contact.check_permission("write")
		if not any(link.link_doctype == "Supplier" and link.link_name == supplier for link in contact.links):
			frappe.throw(_("Contact {0} is not linked to supplier {1}.").format(contact.name, supplier))
	else:
		frappe.has_permission("Supplier", "read", doc=supplier, throw=True)
		contact = frappe.new_doc("Contact")
		contact.append("links", {"link_doctype": "Supplier", "link_name": supplier})
	for field in ("first_name", "last_name", "designation"):
		contact.set(field, (data.get(field) or "").strip())
	_set_primary_child(contact, "email_ids", "email_id", "is_primary", data.get("email_id"))
	_set_primary_child(contact, "phone_nos", "phone", "is_primary_phone", data.get("phone"))
	_set_primary_child(contact, "phone_nos", "phone", "is_primary_mobile_no", data.get("mobile_no"))
	contact.save()
	return frappe.db.get_value("Contact", contact.name, CONTACT_FIELDS, as_dict=True)


@frappe.whitelist(methods=["GET"])
def get_warehouse_project(warehouse: str) -> str:
	"""Project a site warehouse belongs to; receipts into it are always booked to that project."""
	from construction_management.api.purchase_receipt_utils import get_warehouse_project as resolve

	_require(PURCHASE_ROLES | STOCK_ROLES)
	company = frappe.db.get_value("Warehouse", warehouse, "company")
	return resolve(warehouse, company=company) or ""


@frappe.whitelist(methods=["GET"])
def get_terms_template(template: str) -> str:
	"""Terms text from a Terms and Conditions template."""
	_require(PURCHASE_ROLES)
	tnc = frappe.get_doc("Terms and Conditions", template)
	tnc.check_permission("read")
	return tnc.terms or ""


@frappe.whitelist(methods=["GET"])
def get_print_options(doctype: str, name: str) -> dict:
	"""Print formats and letterheads available for a workspace document."""
	doc = _controlled_doc(doctype, name)
	formats = frappe.get_all("Print Format", filters={"doc_type": doctype, "disabled": 0}, pluck="name", order_by="name asc")
	default = frappe.get_meta(doctype).default_print_format
	return {
		"print_formats": ["Standard", *formats],
		"default_print_format": default if default in formats else "Standard",
		"letterheads": frappe.get_all("Letter Head", filters={"disabled": 0}, pluck="name", order_by="is_default desc, name asc"),
		# the document's own letterhead, then its company's, then the site default
		"default_letterhead": doc.get("letter_head")
		or frappe.db.get_value("Company", doc.company, "default_letter_head")
		or frappe.db.get_value("Letter Head", {"is_default": 1, "disabled": 0}, "name"),
	}


@frappe.whitelist(methods=["GET"])
def get_rendered_email_template(doctype: str, name: str, template: str) -> dict:
	"""Render an Email Template's subject and body against a controlled document."""
	doc = _controlled_doc(doctype, name)
	email_template = frappe.get_doc("Email Template", template)
	email_template.check_permission("read")
	return email_template.get_formatted_email(doc.as_dict())


@frappe.whitelist(methods=["POST"])
def send_document_email(
	doctype: str, name: str, recipients: str, subject: str, message: str = "",
	cc: str = "", attach_print: int = 1, print_format: str = "", email_template: str = "",
) -> dict:
	"""Send a controlled document by email; the Communication appears in the activity feed."""
	from frappe.core.doctype.communication.email import make

	doc = _controlled_doc(doctype, name)
	frappe.has_permission(doctype, "email", doc=doc, throw=True)
	if not (recipients or "").strip():
		frappe.throw(_("Add at least one recipient."))
	from construction_management.api.catalog_email import outgoing_sender

	result = make(
		doctype=doctype, name=name, content=message, subject=subject, recipients=recipients, cc=cc or None,
		sender=outgoing_sender(doc.company),
		send_email=1, print_format=(print_format or "Standard") if cint(attach_print) else None,
		email_template=email_template or None,
	)
	return {"name": result.get("name"), "status": "Queued", "emails_not_sent_to": result.get("emails_not_sent_to")}
