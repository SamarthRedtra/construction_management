"""Apply the approved, limited corrections to the MRG Item Tree import source."""

from __future__ import annotations

import hashlib
from pathlib import Path
from tempfile import TemporaryDirectory

import frappe
from openpyxl import load_workbook

from construction_management.api import item_tree_part_2_expanded_import as expanded_import


SOURCE_PATH = Path("/Users/samarthupare/Downloads/Item Tree Part 2 (3).xlsx")
EXPECTED_HELD_ROWS = 100
CORRECTION_PROFILE = "mrg-item-tree-selected-corrections-v1"
SUPPLIERS = (
	"SYNUTHANE BUILDING MATERIALS L.L.C",
	"AIM INDUSTRIES",
	"WARQA STAR MACHINERY & EQUIPMENT RENTAL L.L.C",
	"DST CARGO TRANSPORT",
	"SITCO DIS TIGARET STI LLC",
	"A2A SAFETY CONSULTANTS CO.LLC",
	"QUALITY INTERNATIONAL CERTIFICATES ISSUING SERVICES",
)
SUPPLIER_ALIASES = {
	"AL FAJER TRADING L.L.C": "AL FAJER TRADING LLC",
	"AL BAITH AL RAFIEE FABRICATED HOUSE TR. L.L.C": "AL BAIT AL RAFLEE PREFABRICATED HOUSE TR LL",
	"AL RAHA PREFABRICATED HOUSE TR.": "AL RAHA PREFABRICATED HOUSE TRADING",
}
RATES = {312: 34.5, 313: 21.25}
GROUPS = {
	"ELECTRICAL ITEMS": (355, 361, 365, 366, 387, 412, 421, 423, 468, 473, 670),
	"GROUT AND TILES": (560, 561, 662, 663),
	"HARDWARE ITEMS": (695, 806),
}


def run(source_path: str = str(SOURCE_PATH)) -> dict:
	"""Stage only catalog rows made valid by the approved MRG corrections."""
	frappe.set_user("Administrator")
	source = Path(source_path)
	supplier_results = _ensure_suppliers()
	frappe.db.commit()
	profile = _profile(source)
	with TemporaryDirectory(prefix="mrg-item-tree-") as directory:
		corrected_path = Path(directory) / "Item Tree Part 2 (3) - Corrected.xlsx"
		_write_corrected_workbook(source, corrected_path)
		preview = expanded_import.preview(str(corrected_path))
		_validate_preview(preview)
		staged = _stage_once(corrected_path, profile)
	return {
		"suppliers": supplier_results,
		"corrected_source_held_rows": len(preview["held"]),
		"held_summary": _held_summary(preview["held"]),
		"catalog_request": staged,
		"stock_changes": 0,
	}


def _ensure_suppliers() -> dict[str, list[str]]:
	created, reused = [], []
	for supplier_name in SUPPLIERS:
		existing = _supplier_name(supplier_name)
		if existing:
			reused.append(existing)
			continue
		document = frappe.get_doc({"doctype": "Supplier", "supplier_name": supplier_name, "supplier_type": "Company"})
		document.insert(ignore_permissions=True)
		created.append(document.name)
	return {"created": created, "reused": reused, "aliases": SUPPLIER_ALIASES}


def _supplier_name(value: str) -> str:
	key = expanded_import._key(value)
	for supplier in frappe.get_all("Supplier", fields=["name", "supplier_name"], limit_page_length=0):
		if key in {expanded_import._key(supplier.name), expanded_import._key(supplier.supplier_name)}:
			return supplier.name
	return ""


def _write_corrected_workbook(source: Path, destination: Path) -> None:
	book = load_workbook(source, data_only=True)
	sheet = book["Sheet2"]
	for row_number, rate in RATES.items():
		sheet.cell(row_number, 5).value = rate
	for group, rows in GROUPS.items():
		for row_number in rows:
			sheet.cell(row_number, 1).value = group
	for row_number in range(3, sheet.max_row + 1):
		cell = sheet.cell(row_number, 3)
		cell.value = SUPPLIER_ALIASES.get(str(cell.value or "").strip(), cell.value)
	book.save(destination)


def _validate_preview(preview: dict) -> None:
	if len(preview["held"]) != EXPECTED_HELD_ROWS:
		frappe.throw(
			f"Corrected Item Tree preflight produced {len(preview['held'])} held rows; "
			f"expected {EXPECTED_HELD_ROWS}. No catalog request was staged."
		)


def _profile(source: Path) -> str:
	checksum = hashlib.sha256(source.read_bytes()).hexdigest()
	return f"{CORRECTION_PROFILE}:{checksum}"


def _stage_once(corrected_path: Path, profile: str) -> dict:
	existing = frappe.db.get_value(
		"Controlled Catalog Request", {"import_profile": profile}, ["name", "status"], as_dict=True
	)
	if existing:
		return {"name": existing.name, "status": existing.status, "already_imported": True}
	staged = expanded_import.submit(str(corrected_path))
	if staged.get("name"):
		frappe.db.set_value("Controlled Catalog Request", staged["name"], "import_profile", profile)
	return staged


def _held_summary(held: list[list]) -> dict[str, int]:
	summary = {}
	for _, _, reason in held:
		summary[reason] = summary.get(reason, 0) + 1
	return dict(sorted(summary.items()))
