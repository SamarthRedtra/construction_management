"""Approval-gated expanded Item Tree Part 2 catalog import."""

import csv
import hashlib
from collections import defaultdict
from io import BytesIO, StringIO
from pathlib import Path

import frappe
from frappe.utils import flt, now_datetime
from frappe.utils.file_manager import save_file
from openpyxl import Workbook, load_workbook

from construction_management.api import controlled_procurement as catalog


COMPANY = "M R G INSULATION WORKS L.L.C"
PACKAGE_SOURCE = Path(__file__).resolve().parents[1] / "data" / "controlled_procurement_catalog_item_tree_part_2_v2.xlsx"
PACKAGE_SHA256 = "bf62b0dc65c27336f99d9c1799bc646d2afea8296c51f888133fb9da3a37bea4"
IMPORT_PROFILE = "item-tree-part-2-expanded-v1"
HEADERS = ["Item Code", "Description", "Item Group", "Item Type", "Stock UOM",
	"Asset Category", "Expense Account", "Supplier", "Rates"]
ITEM_TYPES = {
	"material in stock": "Stockable",
	"consumable": "Consumable",
	"asset": "Asset",
	"service": "Service",
}
ASSET_CATEGORIES = {
	"equipment, transport & machinery": "Plant & Machinery- Office",
	"safety items": "Tools & Equipments",
	"hardware items": "Tools & Equipments",
	"waterproofing gass": "Tools & Equipments",
	"injection items": "Tools & Equipments",
	"scaffolding & formwork": "Scaffolding Item",
	"electrical items": "Office Equipments- Office",
	"fixed asset": "Office Equipments- Office",
}
SERVICE_EXPENSE_ACCOUNTS = {
	"equipment, transport & machinery": "EQUIPMENT RENTAL EXPENSE - MRG",
	"scaffolding & formwork": "EQUIPMENT RENTAL EXPENSE - MRG",
	"vehicle maintenance": "REPAIR & MAINTENANCE - MRG",
	"testing & certification": "Professional Charges - MRG",
}


def preview(path: str | None = None) -> dict:
	"""Return the expanded-import preflight without changing catalog records."""
	rows, held, groups, summary = _prepare(Path(path or PACKAGE_SOURCE))
	return {
		"selected": len(rows),
		"held": held,
		"missing_groups": groups,
		"summary": summary,
		"source_sha256": _source_checksum(Path(path or PACKAGE_SOURCE)),
	}


def submit(path: str | None = None) -> dict:
	"""Stage the expanded Sheet2 catalog request for Accounts and CEO approval."""
	frappe.set_user("Administrator")
	source = Path(path or PACKAGE_SOURCE)
	source_checksum = _source_checksum(source)
	_validate_packaged_source(source, source_checksum)
	existing = frappe.db.get_value(
		"Controlled Catalog Request",
		{"full_source_checksum": source_checksum, "import_profile": IMPORT_PROFILE},
		["name", "status"],
		as_dict=True,
	)
	if existing:
		return {"name": existing.name, "status": existing.status, "already_imported": True}

	rows, held, groups, summary = _prepare(source)
	if not rows:
		return {"name": "", "status": "", "already_imported": True, "reason": "No eligible catalog rows"}
	for group in groups:
		catalog._ensure_group(group)

	normalized_bytes = _normalized_workbook(rows)
	normalized = save_file("item_tree_part_2_expanded_controlled_import.xlsx", normalized_bytes, None, None, is_private=1)
	full_source = save_file(source.name, source.read_bytes(), None, None, is_private=1)
	request = frappe.get_doc({
		"doctype": "Controlled Catalog Request",
		"company": COMPANY,
		"status": "Pending Accounts",
		"source": "Workbook Import",
		"requested_by": "Administrator",
		"source_file": normalized.file_url,
		"full_source_file": full_source.file_url,
		"checksum": hashlib.sha256(normalized_bytes).hexdigest(),
		"full_source_checksum": source_checksum,
		"import_profile": IMPORT_PROFILE,
		"items": rows,
		"approval_log": [{
			"action": "Created",
			"actor": "Administrator",
			"actioned_on": now_datetime(),
			"comment": "Expanded Item Tree Part 2 import staged for Accounts approval.",
		}],
	})
	request.insert(ignore_permissions=True)
	for file_doc in (normalized, full_source, save_file(
		"item_tree_part_2_expanded_exceptions.csv", _exception_csv(held), None, None, is_private=1,
	)):
		file_doc.reload()
		file_doc.attached_to_doctype = "Controlled Catalog Request"
		file_doc.attached_to_name = request.name
		file_doc.save(ignore_permissions=True)
	request.add_comment(
		"Info",
		f"Source Sheet2 SHA-256: {source_checksum}. {len(rows)} catalog rows staged; "
		f"{len(held)} source rows held in the attached exception report. {summary}",
	)
	return {
		"name": request.name,
		"status": request.status,
		"already_imported": False,
		"staged_rows": len(rows),
		"held_rows": len(held),
		"missing_groups_created": groups,
		"summary": summary,
	}


def _prepare(path: Path) -> tuple[list[dict], list[tuple], list[str], dict]:
	suppliers = _supplier_ids()
	groups = {_key(name): name for name in frappe.get_all("Item Group", pluck="name", limit_page_length=0)}
	items = _items_by_name()
	valid, held = _valid_source_rows(path, suppliers)
	rows, missing_groups, summary = _select_catalog_rows(valid, held, groups, items)
	return rows, sorted(held), sorted(missing_groups), summary


def _valid_source_rows(path: Path, suppliers: dict[str, str]) -> tuple[list[dict], list[tuple]]:
	valid, held = [], []
	for number, values in _source_rows(path):
		item_type = ITEM_TYPES.get(_key(values[5]))
		name = str(values[1]).strip()
		if not item_type:
			held.append((number, name, "Unsupported allocation"))
			continue
		supplier = suppliers.get(_key(values[2]))
		if not supplier:
			held.append((number, name, "Supplier master not found"))
			continue
		try:
			uom = catalog._catalog_uom(values[3])
		except frappe.ValidationError as error:
			held.append((number, name, str(error)))
			continue
		rate = flt(values[4])
		if item_type != "Service" and rate <= 0:
			held.append((number, name, "Positive buying rate required"))
			continue
		group = str(values[0]).strip()
		asset_category, expense_account = _accounting_mapping(item_type, group)
		if item_type == "Asset" and not asset_category:
			held.append((number, name, "Asset Category mapping required"))
			continue
		if item_type == "Service" and not expense_account:
			held.append((number, name, "Service Expense Account mapping required"))
			continue
		valid.append({
			"source_row": number,
			"item_name": name,
			"item_group": group,
			"item_type": item_type,
			"stock_uom": uom,
			"asset_category": asset_category,
			"expense_account": expense_account,
			"supplier": supplier,
			"rate": rate,
		})
	return valid, held


def _select_catalog_rows(valid: list[dict], held: list[tuple], groups: dict, items: dict) -> tuple[list[dict], set[str], dict]:
	selected, missing_groups = [], set()
	summary = defaultdict(int)
	by_identity = defaultdict(list)
	for row in valid:
		by_identity[_identity(row)].append(row)

	for identity, entries in by_identity.items():
		usable = _deduplicate_identity(entries, held)
		if not usable:
			continue
		first = usable[0]
		matches = [item for item in items.get(_key(first["item_name"]), []) if _key(item.stock_uom) == _key(first["stock_uom"])]
		action, result_item = _item_action(matches, first, held, usable)
		if action == "Reuse":
			summary["reused"] += 1
			continue
		if not action:
			continue
		if _key(first["item_group"]) not in groups:
			missing_groups.add(first["item_group"])
		for row in usable:
			selected.append({
				"item_code": "",
				"item_name": row["item_name"],
				"item_group": groups.get(_key(row["item_group"]), row["item_group"]),
				"item_type": row["item_type"],
				"stock_uom": row["stock_uom"],
				"asset_category": row["asset_category"],
				"expense_account": row["expense_account"],
				"supplier": row["supplier"],
				"rate": row["rate"],
				"action": action,
				"result_item": result_item,
			})
			summary[action.lower()] += 1
	return selected, missing_groups, dict(summary)


def _deduplicate_identity(entries: list[dict], held: list[tuple]) -> list[dict]:
	if len({(_key(row["item_group"]), row["asset_category"], row["expense_account"]) for row in entries}) != 1:
		held.extend((row["source_row"], row["item_name"], "Conflicting item group or accounting mapping") for row in entries)
		return []
	by_supplier = defaultdict(list)
	for row in entries:
		by_supplier[row["supplier"]].append(row)
	usable = []
	for supplier_rows in by_supplier.values():
		if len({row["rate"] for row in supplier_rows}) != 1:
			held.extend((row["source_row"], row["item_name"], "Conflicting buying rates for the same supplier") for row in supplier_rows)
			continue
		usable.append(supplier_rows[0])
	return usable


def _item_action(matches: list, row: dict, held: list[tuple], entries: list[dict]) -> tuple[str, str]:
	controlled = [item for item in matches if item.controlled_procurement_catalog]
	if controlled:
		if len(controlled) == 1 and controlled[0].controlled_item_type == row["item_type"]:
			return "Reuse", controlled[0].name
		held.extend((entry["source_row"], entry["item_name"], "Existing controlled Item conflicts with this type/UOM") for entry in entries)
		return "", ""
	if len(matches) == 1:
		return "Adopt", matches[0].name
	if len(matches) > 1:
		held.extend((entry["source_row"], entry["item_name"], "Multiple existing Items have this UOM") for entry in entries)
		return "", ""
	return "Create", ""


def _accounting_mapping(item_type: str, group: str) -> tuple[str, str]:
	group_key = _key(group)
	if item_type == "Asset":
		return ASSET_CATEGORIES.get(group_key, ""), ""
	if item_type == "Service":
		return "", SERVICE_EXPENSE_ACCOUNTS.get(group_key, "")
	return "", ""


def _items_by_name() -> dict[str, list]:
	items = defaultdict(list)
	for item in frappe.get_all(
		"Item",
		fields=["name", "item_name", "stock_uom", "controlled_procurement_catalog", "controlled_item_type", "disabled"],
		limit_page_length=0,
	):
		if not item.disabled:
			items[_key(item.item_name)].append(item)
	return items


def _supplier_ids() -> dict[str, str]:
	return {
		_key(value): row.name
		for row in frappe.get_all("Supplier", fields=["name", "supplier_name"], limit_page_length=0)
		for value in (row.name, row.supplier_name)
		if value
	}


def _identity(row: dict) -> tuple[str, str, str]:
	return _key(row["item_name"]), row["item_type"], _key(row["stock_uom"])


def _source_rows(path: Path):
	book = load_workbook(path, read_only=True, data_only=True)
	if "Sheet2" not in book:
		frappe.throw("Sheet2 is missing from the source workbook.")
	return [
		(number, values)
		for number, values in enumerate(book["Sheet2"].values, 1)
		if number > 2 and len(values) > 5 and values[0] and values[1] and _key(values[0]) != "item group"
	]


def _normalized_workbook(rows: list[dict]) -> bytes:
	book = Workbook()
	sheet = book.active
	sheet.title = "Catalog Items"
	sheet.append(HEADERS)
	for row in rows:
		sheet.append([row[key] for key in ("item_code", "item_name", "item_group", "item_type", "stock_uom",
			"asset_category", "expense_account", "supplier", "rate")])
	stream = BytesIO()
	book.save(stream)
	return stream.getvalue()


def _exception_csv(held: list[tuple]) -> bytes:
	stream = StringIO()
	writer = csv.writer(stream)
	writer.writerow(["Sheet", "Source row", "Description", "Reason"])
	for number, description, reason in held:
		writer.writerow(["Sheet2", number, description, reason])
	return stream.getvalue().encode("utf-8-sig")


def _source_checksum(path: Path) -> str:
	return hashlib.sha256(path.read_bytes()).hexdigest()


def _validate_packaged_source(path: Path, checksum: str) -> None:
	if path.resolve() == PACKAGE_SOURCE.resolve() and checksum != PACKAGE_SHA256:
		frappe.throw("The packaged Item Tree Part 2 workbook checksum is invalid.")


def _key(value) -> str:
	return " ".join(str(value or "").split()).casefold()
