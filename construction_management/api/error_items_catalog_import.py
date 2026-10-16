# Copyright (c) 2026, Construction Management
# License: MIT

"""Import Error Items.xlsx rows missing from the controlled catalogue."""

from __future__ import annotations

import csv
import hashlib
from collections import defaultdict
from io import BytesIO, StringIO
from pathlib import Path

import frappe
from frappe.utils import cint
from frappe.utils.file_manager import save_file
from openpyxl import Workbook, load_workbook

from construction_management.api import controlled_procurement as catalog


COMPANY = "M R G INSULATION WORKS L.L.C"
PACKAGE_SOURCE = (
	Path(__file__).resolve().parents[1] / "data" / "error_items_controlled_catalogue.xlsx"
)
PACKAGE_SHA256 = "906c71bd5974565e2938f18f395daa568e931152465548395a615d114820d0de"
HEADERS = [
	"Item Code",
	"Description",
	"Item Group",
	"Item Type",
	"Stock UOM",
	"Asset Category",
	"Expense Account",
	"Supplier",
	"Rates",
]
UOM_ALIASES = {
	"nos.": "Nos",
	"nos": "Nos",
	"ton": "Tonne",
	"day": "Day",
	"set": "Set",
	"hrs": "Hour",
	"monthly": "Months",
}


def _key(value) -> str:
	return " ".join(str(value or "").split()).casefold()


def _source_rows(path: Path) -> list[tuple[int, tuple]]:
	book = load_workbook(path, read_only=True, data_only=True)
	sheet = book.active
	rows = []
	for number, values in enumerate(sheet.values, 1):
		if number == 1 or not values or not values[1]:
			continue
		if _key(values[0]) == "item group":
			continue
		rows.append((number, values))
	return rows


def _resolve_uom(raw) -> str:
	value = str(raw or "Nos").strip()
	value = UOM_ALIASES.get(value.casefold(), value)
	return catalog._catalog_uom(value)


def _controlled_names() -> set[str]:
	return {
		_key(row.item_name)
		for row in frappe.get_all(
			"Item",
			filters={"controlled_procurement_catalog": 1, "disabled": 0},
			fields=["item_name"],
			limit_page_length=0,
		)
	}


def _legacy_by_name(description: str) -> list[dict]:
	return frappe.db.sql(
		"""
		SELECT name, item_name, stock_uom, is_stock_item, controlled_procurement_catalog
		FROM `tabItem`
		WHERE LOWER(item_name) = LOWER(%s) AND IFNULL(disabled, 0) = 0
		""",
		description,
		as_dict=True,
	)


def _allocation_of(row) -> str:
	return _key(row[5] if len(row) > 5 else "")


def _prepare(path: Path):
	controlled = _controlled_names()
	suppliers = frappe.get_all("Supplier", fields=["name", "supplier_name"], limit_page_length=0)
	supplier_ids = {
		_key(value): row.name for row in suppliers for value in (row.name, row.supplier_name)
	}
	groups = {_key(name): name for name in frappe.get_all("Item Group", pluck="name", limit_page_length=0)}
	by_name: dict[str, list] = defaultdict(list)
	held = []

	for number, values in _source_rows(path):
		description = str(values[1]).strip()
		name_key = _key(description)
		if name_key in controlled:
			continue
		allocation = _key(values[5] if len(values) > 5 else "")
		if allocation not in {"material in stock", "consumable"}:
			held.append((number, description, "Asset/Service: accounting mapping required"))
			continue
		by_name[name_key].append((number, values))

	selected = []
	missing_groups = set()
	for entries in by_name.values():
		identity = {
			(
				_key(row[0]),
				_key(row[2]),
				_key(row[3]),
				row[4],
				_key(row[5] if len(row) > 5 else ""),
			)
			for _, row in entries
		}
		if len(identity) != 1:
			held.extend(
				(number, row[1], "Conflicting duplicate description or rate") for number, row in entries
			)
			continue

		number, row = entries[0]
		description = str(row[1]).strip()
		supplier = supplier_ids.get(_key(row[2])) or catalog._supplier_name(str(row[2] or "").strip())
		if not supplier:
			held.append((number, description, "Supplier master not found"))
			continue

		group = str(row[0]).strip()
		if _key(group) not in groups:
			missing_groups.add(group)

		try:
			uom = _resolve_uom(row[3])
		except frappe.ValidationError as error:
			held.append((number, description, str(error)))
			continue

		if not isinstance(row[4], (int, float)) or row[4] <= 0:
			held.append((number, description, "Positive buying rate required"))
			continue

		legacy = _legacy_by_name(description)
		item_code = ""
		item_name = description
		if legacy:
			matching = [item for item in legacy if item.stock_uom == uom]
			if not matching:
				held.append(
					(
						number,
						description,
						f"Existing item has another stock UOM ({legacy[0].stock_uom}); cannot change by import",
					)
				)
				continue
			chosen = matching[0]
			# Catalog Consumable/Stockable require Maintain Stock on. Do not flip
			# is_stock_item on legacy Items that already have submitted vouchers.
			if cint(chosen.is_stock_item):
				item_code = chosen.name
				item_name = chosen.item_name
			else:
				item_code = catalog._item_code(description)
				item_name = description

		item_type = "Stockable" if _allocation_of(row) == "material in stock" else "Consumable"
		selected.append(
			(
				number,
				[
					item_code,
					item_name,
					groups.get(_key(group), group),
					item_type,
					uom,
					"",
					"",
					supplier,
					row[4],
				],
			)
		)

	return selected, sorted(held), sorted(missing_groups)


def _normalized_workbook(selected) -> bytes:
	book = Workbook()
	sheet = book.active
	sheet.title = "Catalog Items"
	sheet.append(HEADERS)
	for _, values in selected:
		sheet.append(values)
	stream = BytesIO()
	book.save(stream)
	return stream.getvalue()


def _exception_csv(held) -> bytes:
	stream = StringIO()
	writer = csv.writer(stream)
	writer.writerow(["Sheet", "Source row", "Description", "Reason"])
	for number, description, reason in held:
		writer.writerow(["Sheet1", number, description, reason])
	return stream.getvalue().encode("utf-8-sig")


def _source_checksum(path: Path) -> str:
	return hashlib.sha256(path.read_bytes()).hexdigest()


def _existing_request(source_checksum: str):
	return frappe.db.get_value(
		"Controlled Catalog Request",
		{"full_source_checksum": source_checksum},
		["name", "status"],
		as_dict=True,
	)


def _validate_packaged_source(path: Path) -> None:
	if path.resolve() != PACKAGE_SOURCE.resolve():
		return
	if _source_checksum(path) != PACKAGE_SHA256:
		frappe.throw("The packaged Error Items workbook checksum is invalid.")


def preview(path: str | None = None) -> dict:
	"""Read-only preflight for the Error Items workbook."""
	path = Path(path or PACKAGE_SOURCE)
	selected, held, groups = _prepare(path)
	return {
		"selected": len(selected),
		"stockable": sum(row[3] == "Stockable" for _, row in selected),
		"consumable": sum(row[3] == "Consumable" for _, row in selected),
		"missing_groups": groups,
		"held": held,
		"source_sha256": _source_checksum(path),
	}


def submit(path: str | None = None) -> dict:
	"""Create missing groups, one catalog request, and materialize Items."""
	frappe.set_user("Administrator")
	source = Path(path or PACKAGE_SOURCE)
	source_checksum = _source_checksum(source)
	existing = _existing_request(source_checksum)
	if existing:
		return {"name": existing.name, "status": existing.status, "already_imported": True}

	selected, held, groups = _prepare(source)
	_validate_packaged_source(source)
	if not selected:
		return {
			"name": "",
			"status": "",
			"already_imported": True,
			"reason": "No missing catalog items",
			"held_rows": len(held),
		}

	for group in groups:
		if not frappe.db.exists("Item Group", group):
			frappe.get_doc(
				{
					"doctype": "Item Group",
					"item_group_name": group,
					"parent_item_group": "All Item Groups",
					"is_group": 0,
				}
			).insert()

	frappe.defaults.set_user_default("company", COMPANY)
	normalized = save_file(
		"error_items_controlled_import.xlsx",
		_normalized_workbook(selected),
		None,
		None,
		is_private=1,
	)
	preview_result = catalog.preview_catalog_workbook(normalized.file_url)
	if preview_result["errors"] or preview_result["valid_rows"] != len(selected):
		frappe.throw("Normalized workbook preview did not confirm the expected valid items.")

	result = catalog.import_catalog_workbook(normalized.file_url)
	request_name = result["name"]

	full_source = save_file(source.name, source.read_bytes(), None, None, is_private=1)
	for file_doc in (
		normalized,
		full_source,
		save_file("error_items_exceptions.csv", _exception_csv(held), None, None, is_private=1),
	):
		file_doc.reload()
		file_doc.attached_to_doctype = "Controlled Catalog Request"
		file_doc.attached_to_name = request_name
		file_doc.save(ignore_permissions=True)

	request = frappe.get_doc("Controlled Catalog Request", request_name)
	request.full_source_file = full_source.file_url
	request.full_source_checksum = source_checksum
	request.flags.controlled_catalog_request_api = True
	request.save(ignore_permissions=True)
	request.add_comment(
		"Info",
		f"Error Items SHA-256: {source_checksum}. "
		f"Selected {len(selected)} Stockable/Consumable items; "
		f"{len(held)} source rows held in the attached exception report.",
	)
	request.reload()
	return {
		**result,
		"status": request.status,
		"preview": {
			key: preview_result[key] for key in ("valid_rows", "create", "reuse", "errors")
		},
		"held_rows": len(held),
		"missing_groups_created": groups,
		"source_file": normalized.file_url,
	}


