# Copyright (c) 2026, Construction Management
# License: MIT

"""API helpers for BOQ-style Quotation HTML content."""

from __future__ import annotations

import html

import frappe
from frappe.utils import flt


BOQ_TABLE_STYLE = (
	"width:100%;border-collapse:collapse;table-layout:fixed;"
	"font-family:Arial,Helvetica,sans-serif;font-size:11px;"
)
BOQ_CELL_STYLE = (
	"border:1px solid #333;padding:5px 7px;vertical-align:top;"
	"word-break:break-word;overflow-wrap:anywhere;white-space:normal;box-sizing:border-box;"
)
BOQ_HEAD_STYLE = BOQ_CELL_STYLE + "background:#e8e8e8;font-weight:bold;text-align:center;"
BOQ_SECTION_STYLE = BOQ_CELL_STYLE + "font-weight:bold;text-transform:uppercase;background:#f0f0f0;"
BOQ_NUM_STYLE = (
	"border:1px solid #333;padding:5px 7px;vertical-align:top;"
	"text-align:right;white-space:nowrap;box-sizing:border-box;"
)


@frappe.whitelist()
def get_boq_html_template() -> str:
	"""Return a starter BOQ HTML table users can edit in the Text Editor."""
	return _render_boq_html([], include_totals=True, template_only=True)


@frappe.whitelist()
def preview_boq_html_from_lines(
	lines: str | list | None = None,
	company: str | None = None,
	include_vat: int | str = 1,
) -> str:
	"""Build bordered BOQ HTML from draft child rows for the Quotation BOQ tab preview."""
	import json

	from frappe.utils import cint

	if isinstance(lines, str):
		lines = json.loads(lines or "[]")
	return lines_to_boq_html(
		lines or [],
		include_totals=True,
		company=company,
		include_vat=bool(cint(include_vat)),
	)


@frappe.whitelist()
def get_quotation_boq_hierarchy(company: str | None = None) -> str:
	"""Return configured Quotation BOQ hierarchy for a company."""
	from construction_management.quotation_boq_hierarchy import get_quotation_boq_hierarchy as _get_hierarchy

	return _get_hierarchy(company)


@frappe.whitelist()
def import_boq_lines_from_estimation(estimation_name: str) -> list[dict]:
	"""Build Quotation BOQ Line rows from Project Estimation (grouped by task)."""
	if not estimation_name:
		frappe.throw("Project Estimation is required")

	if not frappe.db.exists("Project Estimation", estimation_name):
		frappe.throw(f"Project Estimation {estimation_name} not found")

	estimation = frappe.get_doc("Project Estimation", estimation_name)
	source_items = _get_estimation_source_items(estimation)
	two_level = _is_two_level_import(estimation.company)

	lines: list[dict] = []
	parent_counter = 0
	current_task = object()

	for item in source_items:
		task_name = (item.get("task") or item.get("activity_type") or "General").strip()
		if task_name != current_task:
			current_task = task_name
			if not two_level:
				lines.append(
					{
						"line_type": "Section",
						"section_title": task_name,
						"description": task_name,
					}
				)
			parent_counter = 0

		parent_counter += 1
		parent_row = {
			"line_type": "Parent",
			"parent_no": str(parent_counter),
			"description": _estimation_item_description(item),
		}
		if two_level:
			parent_row["section_title"] = task_name
		lines.append(parent_row)

	return lines


def _is_two_level_import(company: str | None) -> bool:
	from construction_management.quotation_boq_hierarchy import is_two_level_hierarchy

	return is_two_level_hierarchy(company)


@frappe.whitelist()
def import_boq_html_from_estimation(estimation_name: str) -> str:
	"""Build BOQ HTML from Project Estimation (for print / legacy)."""
	lines = import_boq_lines_from_estimation(estimation_name)
	return lines_to_boq_html(lines, include_totals=True)


@frappe.whitelist()
def parse_pasted_boq_text(text: str, company: str | None = None) -> list[dict]:
	"""Parse Excel/clipboard text into Quotation BOQ Line dicts."""
	return _parse_pasted_boq_text(text, company)


QUOTATION_BOQ_EXCEL_COLUMNS = [
	"type",
	"section",
	"no",
	"description",
	"unit",
	"qty",
	"rate",
	"display",
]

QUOTATION_BOQ_EXCEL_SAMPLE_ROWS = [
	{
		"type": "Section",
		"section": "",
		"no": "",
		"description": "THERMAL AND MOISTURE PROTECTION",
		"unit": "",
		"qty": "",
		"rate": "",
		"display": "",
	},
	{
		"type": "Parent",
		"section": "",
		"no": "1",
		"description": "1000 gauge polythene sheets; all in accordance with drawings and specification",
		"unit": "",
		"qty": "",
		"rate": "",
		"display": "",
	},
	{
		"type": "Sub",
		"section": "",
		"no": "A",
		"description": "To Raft Slab",
		"unit": "m2",
		"qty": 2259,
		"rate": "",
		"display": "N/A",
	},
	{
		"type": "Sub",
		"section": "",
		"no": "B",
		"description": "Below Grade Slab",
		"unit": "m2",
		"qty": 1800,
		"rate": "",
		"display": "N/A",
	},
	{
		"type": "Parent",
		"section": "",
		"no": "2",
		"description": "2 layers of 4mm SBS waterproofing membrane 180gm/m2 from Awazel or Equivalent",
		"unit": "",
		"qty": "",
		"rate": "",
		"display": "",
	},
	{
		"type": "Sub",
		"section": "",
		"no": "A",
		"description": "Horizontally to Raft Slab",
		"unit": "m2",
		"qty": 2184,
		"rate": 62,
		"display": "Normal",
	},
	{
		"type": "Sub",
		"section": "",
		"no": "B",
		"description": "Vertically to Retaining Wall",
		"unit": "m2",
		"qty": 450,
		"rate": 75,
		"display": "Normal",
	},
]


@frappe.whitelist()
def get_quotation_boq_excel_format() -> dict:
	"""Return column definitions and sample rows for Quotation BOQ Excel import."""
	return {
		"columns": [
			{"key": "type", "label": "Type", "required": True, "values": "Section / Parent / Sub"},
			{"key": "section", "label": "Section", "required": False, "values": "Optional section title for 2-level BOQ"},
			{"key": "no", "label": "No", "required": False, "values": "Parent: 1, 2... | Sub: A, B..."},
			{"key": "description", "label": "Description", "required": True, "values": "Work description"},
			{"key": "unit", "label": "Unit", "required": False, "values": "m2, nos, etc."},
			{"key": "qty", "label": "Qty", "required": False, "values": "Quantity"},
			{"key": "rate", "label": "Rate", "required": False, "values": "Unit rate (leave blank for N/A)"},
			{"key": "display", "label": "Display", "required": False, "values": "Normal / N/A / Rate Only"},
		],
		"sample_rows": QUOTATION_BOQ_EXCEL_SAMPLE_ROWS,
		"notes": [
			"Use one row per Section, Parent, or Sub line.",
			"For 3-level BOQ, add Section rows. For 2-level BOQ, Section column on Parent is optional.",
			"You can also paste tab-separated data in Paste from Excel mode using #, 1, A patterns.",
		],
	}


@frappe.whitelist()
def download_quotation_boq_example_template(company: str | None = None) -> str:
	"""Generate downloadable example Excel template for Quotation BOQ import."""
	try:
		import openpyxl
		from io import BytesIO
		from openpyxl.styles import Font, PatternFill
	except ImportError:
		frappe.throw("openpyxl is required for Excel template download.")

	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = "BOQ Import"

	header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
	header_font = Font(color="FFFFFF", bold=True)

	headers = [col.replace("_", " ").title() for col in QUOTATION_BOQ_EXCEL_COLUMNS]
	for col_idx, header in enumerate(headers, 1):
		cell = ws.cell(row=1, column=col_idx, value=header)
		cell.fill = header_fill
		cell.font = header_font

	for row_idx, row_data in enumerate(QUOTATION_BOQ_EXCEL_SAMPLE_ROWS, 2):
		for col_idx, key in enumerate(QUOTATION_BOQ_EXCEL_COLUMNS, 1):
			ws.cell(row=row_idx, column=col_idx, value=row_data.get(key, ""))

	for col in ws.columns:
		max_length = 12
		column = col[0].column_letter
		for cell in col:
			value = str(cell.value or "")
			max_length = max(max_length, min(len(value) + 2, 60))
		ws.column_dimensions[column].width = max_length

	help_ws = wb.create_sheet("Instructions")
	help_ws["A1"] = "Quotation BOQ Excel Import Format"
	help_ws["A1"].font = Font(bold=True, size=12)
	instructions = [
		"",
		"Required columns: Type, Description",
		"Type values: Section, Parent, Sub",
		"Parent No: 1, 2, 3...",
		"Sub No: A, B, C...",
		"Display: Normal (default), N/A, Rate Only",
		"",
		"Fill rows in BOQ Import sheet, save file, then use Import Excel File in Easy BOQ Entry.",
	]
	for idx, line in enumerate(instructions, 2):
		help_ws.cell(row=idx, column=1, value=line)

	output = BytesIO()
	wb.save(output)
	output.seek(0)

	file_doc = frappe.get_doc(
		{
			"doctype": "File",
			"file_name": "quotation_boq_import_example.xlsx",
			"content": output.getvalue(),
			"is_private": 0,
		}
	)
	file_doc.insert(ignore_permissions=True)
	return file_doc.file_url


@frappe.whitelist()
def import_boq_lines_from_excel(file_url: str, company: str | None = None) -> list[dict]:
	"""Parse uploaded Excel file into Quotation BOQ Line dicts."""
	rows = _read_quotation_boq_excel_rows(file_url)
	return _parse_excel_boq_rows(rows, company)


def _read_quotation_boq_excel_rows(file_url: str) -> list[dict]:
	if not file_url:
		frappe.throw("Excel file is required")

	try:
		import openpyxl
		from io import BytesIO
	except ImportError:
		frappe.throw("openpyxl is required for Excel import.")

	file_doc = frappe.get_doc("File", {"file_url": file_url})
	content = file_doc.get_content()
	wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
	ws = wb.active

	headers: list[str] = []
	for cell in ws[1]:
		headers.append(_normalize_excel_header(cell.value))

	if not headers or not any(headers):
		frappe.throw("Excel file has no header row.")

	rows: list[dict] = []
	for row in ws.iter_rows(min_row=2, values_only=True):
		if not row or not any(cell not in (None, "") for cell in row):
			continue
		row_data = {}
		for idx, header in enumerate(headers):
			if not header:
				continue
			value = row[idx] if idx < len(row) else ""
			row_data[header] = value
		if any(str(v).strip() for v in row_data.values() if v is not None):
			rows.append(row_data)

	if not rows:
		frappe.throw("No data rows found in Excel file.")

	return rows


def _normalize_excel_header(value) -> str:
	label = str(value or "").strip().lower()
	mapping = {
		"type": "type",
		"row type": "type",
		"line type": "type",
		"section": "section",
		"section title": "section",
		"no": "no",
		"no.": "no",
		"number": "no",
		"sl no": "no",
		"description": "description",
		"description of works": "description",
		"unit": "unit",
		"uom": "unit",
		"qty": "qty",
		"quantity": "qty",
		"approx. qty": "qty",
		"approx qty": "qty",
		"rate": "rate",
		"rate aed": "rate",
		"display": "display",
		"display mode": "display",
	}
	return mapping.get(label, label.replace(" ", "_"))


def _parse_excel_boq_rows(rows: list[dict], company: str | None = None) -> list[dict]:
	two_level = _is_two_level_import(company)
	lines: list[dict] = []
	current_section = ""
	parent_counter = 0
	sub_counters: dict[str, int] = {}
	letters = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"

	for row in rows:
		row_type = _excel_row_type(row)
		description = _excel_cell_text(row.get("description"))
		section_value = _excel_cell_text(row.get("section"))
		no_value = _excel_cell_text(row.get("no"))

		if row_type == "section":
			title = description or section_value or no_value
			if title.startswith("#"):
				title = title.lstrip("#").strip()
			current_section = title
			if title and not two_level:
				lines.append(
					{
						"line_type": "Section",
						"section_title": title,
						"description": title,
					}
				)
			continue

		if row_type == "parent":
			parent_no = no_value or str(_next_parent_no(parent_counter))
			parent_counter = int(parent_no) if parent_no.isdigit() else parent_counter + 1
			sub_counters[str(parent_counter)] = 0
			parent_row = {
				"line_type": "Parent",
				"parent_no": str(parent_counter),
				"description": description,
			}
			section_for_parent = section_value or current_section
			if two_level and section_for_parent:
				parent_row["section_title"] = section_for_parent
			lines.append(parent_row)
			continue

		# Sub row (explicit or inferred)
		if parent_counter == 0:
			parent_counter = 1
			sub_counters["1"] = 0
			parent_row = {
				"line_type": "Parent",
				"parent_no": "1",
				"description": "General",
			}
			if two_level and current_section:
				parent_row["section_title"] = current_section
			lines.append(parent_row)

		parent_key = str(parent_counter)
		sub_counters.setdefault(parent_key, 0)
		sub_no = no_value.upper() if no_value and len(no_value) == 1 and no_value.isalpha() else ""
		if not sub_no:
			sub_no = letters[sub_counters[parent_key] % len(letters)]
		sub_counters[parent_key] += 1

		uom = _excel_cell_text(row.get("unit")) or "Nos"
		qty = flt(row.get("qty"))
		rate = flt(row.get("rate"))
		display_mode = _excel_display_mode(row.get("display"), rate, qty)

		lines.append(
			{
				"line_type": "Sub",
				"parent_no": parent_key,
				"sub_no": sub_no,
				"description": description,
				"uom": uom,
				"qty": qty,
				"rate": rate,
				"amount": qty * rate if display_mode == "Normal" else 0,
				"is_fixed_rate": 0,
				"display_mode": display_mode,
			}
		)

	return lines


def _excel_row_type(row: dict) -> str:
	raw_type = _excel_cell_text(row.get("type")).lower()
	if raw_type in {"section", "parent", "sub"}:
		return raw_type

	no_value = _excel_cell_text(row.get("no"))
	description = _excel_cell_text(row.get("description"))
	if description.startswith("#") or raw_type in {"#", "section header"}:
		return "section"
	if no_value.isdigit() or raw_type in {"1", "parent"}:
		return "parent"
	if (len(no_value) == 1 and no_value.isalpha()) or raw_type in {"a", "sub"}:
		return "sub"

	# Infer from qty/rate columns when type is blank
	if row.get("unit") or row.get("qty") or row.get("rate"):
		return "sub"
	if description and not no_value:
		return "parent"
	return "sub"


def _excel_display_mode(display_value, rate: float, qty: float) -> str:
	display = _excel_cell_text(display_value)
	if display.lower() in {"n/a", "na"}:
		return "N/A"
	if display.lower() in {"rate only", "rateonly"}:
		return "Rate Only"
	if display.lower() == "normal":
		return "Normal"
	if rate == 0 and qty > 0:
		return "N/A"
	return "Normal"


def _excel_cell_text(value) -> str:
	if value is None:
		return ""
	return str(value).strip()


def _next_parent_no(current: int) -> int:
	return current + 1 if current else 1

@frappe.whitelist()
def build_quick_boq_lines(
	company: str | None,
	parent_description: str,
	sub_rows: str | list | None = None,
	section_title: str | None = None,
	parent_no: str | None = None,
) -> list[dict]:
	"""Build BOQ lines for one parent and optional section with multiple sub rows."""
	import json

	if isinstance(sub_rows, str):
		sub_rows = json.loads(sub_rows) if sub_rows else []

	two_level = _is_two_level_import(company)
	lines: list[dict] = []
	section_title = (section_title or "").strip()
	parent_description = (parent_description or "").strip()

	if section_title and not two_level:
		lines.append(
			{
				"line_type": "Section",
				"section_title": section_title,
				"description": section_title,
			}
		)

	parent_row = {
		"line_type": "Parent",
		"parent_no": (parent_no or "1").strip(),
		"description": parent_description,
	}
	if two_level and section_title:
		parent_row["section_title"] = section_title
	lines.append(parent_row)

	letters = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
	for idx, row in enumerate(sub_rows or []):
		if isinstance(row, dict):
			row_data = row
		else:
			continue

		description = (row_data.get("description") or "").strip()
		if not description:
			continue

		sub_no = (row_data.get("sub_no") or "").strip() or letters[idx % len(letters)]
		lines.append(
			{
				"line_type": "Sub",
				"parent_no": parent_row["parent_no"],
				"sub_no": sub_no,
				"description": description,
				"uom": row_data.get("uom") or "Nos",
				"qty": flt(row_data.get("qty")),
				"rate": flt(row_data.get("rate")),
				"amount": flt(row_data.get("qty")) * flt(row_data.get("rate")),
				"is_fixed_rate": 0,
				"display_mode": row_data.get("display_mode") or "Normal",
			}
		)

	return lines


def _parse_pasted_boq_text(text: str, company: str | None = None) -> list[dict]:
	"""Parse clipboard rows into Section / Parent / Sub BOQ lines."""
	if not text or not str(text).strip():
		return []

	two_level = _is_two_level_import(company)
	lines: list[dict] = []
	current_section = ""
	parent_counter = 0
	sub_counters: dict[str, int] = {}
	letters = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"

	for raw_line in str(text).splitlines():
		line = raw_line.strip()
		if not line:
			continue

		cells = [_clean_paste_cell(c) for c in _split_paste_line(line)]
		cells = [c for c in cells if c != ""]
		if not cells:
			continue

		if _is_section_paste_row(cells):
			current_section = _section_title_from_cells(cells)
			if not two_level:
				lines.append(
					{
						"line_type": "Section",
						"section_title": current_section,
						"description": current_section,
					}
				)
			continue

		parent_no = _parse_parent_no(cells[0])
		if parent_no and len(cells) <= 3 and not _looks_like_sub_row(cells):
			parent_counter = int(parent_no)
			sub_counters[parent_no] = 0
			description = cells[1] if len(cells) > 1 else cells[0]
			parent_row = {
				"line_type": "Parent",
				"parent_no": str(parent_counter),
				"description": description,
			}
			if two_level and current_section:
				parent_row["section_title"] = current_section
			lines.append(parent_row)
			continue

		sub_no = _parse_sub_no(cells[0])
		desc_idx = 1 if sub_no and len(cells) > 1 else 0
		description = cells[desc_idx] if len(cells) > desc_idx else cells[0]
		rest = cells[desc_idx + 1 :] if len(cells) > desc_idx + 1 else []

		if not sub_no and parent_counter == 0:
			parent_counter = 1
			sub_counters["1"] = 0
			parent_row = {
				"line_type": "Parent",
				"parent_no": "1",
				"description": "General",
			}
			if two_level and current_section:
				parent_row["section_title"] = current_section
			lines.append(parent_row)

		parent_key = str(parent_counter or 1)
		sub_counters.setdefault(parent_key, 0)
		if not sub_no:
			sub_no = letters[sub_counters[parent_key] % len(letters)]
		sub_counters[parent_key] += 1

		uom, qty, rate = _parse_qty_rate_cells(rest)
		display_mode = "Normal"
		if rate == 0 and qty > 0:
			display_mode = "Rate Only"
		if "n/a" in " ".join(rest).lower():
			display_mode = "N/A"

		lines.append(
			{
				"line_type": "Sub",
				"parent_no": parent_key,
				"sub_no": sub_no,
				"description": description,
				"uom": uom or "Nos",
				"qty": qty,
				"rate": rate,
				"amount": qty * rate if display_mode == "Normal" else 0,
				"is_fixed_rate": 0,
				"display_mode": display_mode,
			}
		)

	return lines


def _split_paste_line(line: str) -> list[str]:
	if "\t" in line:
		return line.split("\t")
	if "|" in line:
		return [part.strip() for part in line.split("|")]
	return [part.strip() for part in line.split(",")]


def _clean_paste_cell(value: str) -> str:
	return (value or "").strip().strip(".").strip("-")


def _is_section_paste_row(cells: list[str]) -> bool:
	first = (cells[0] or "").upper()
	if first.startswith("#") or first.startswith("SECTION"):
		return True
	if len(cells) == 1 and first.isupper() and len(first) > 8:
		return True
	return False


def _section_title_from_cells(cells: list[str]) -> str:
	title = cells[0]
	if title.upper().startswith("SECTION"):
		title = title.split(":", 1)[-1].strip()
	return title.lstrip("#").strip()


def _parse_parent_no(value: str) -> str | None:
	value = _clean_paste_cell(value)
	if value.isdigit():
		return value
	if value.endswith("-") and value[:-1].isdigit():
		return value[:-1]
	return None


def _parse_sub_no(value: str) -> str | None:
	value = _clean_paste_cell(value)
	if len(value) == 1 and value.isalpha():
		return value.upper()
	return None


def _looks_like_sub_row(cells: list[str]) -> bool:
	if len(cells) < 3:
		return False
	uom, qty, rate = _parse_qty_rate_cells(cells[1:])
	return bool(qty or rate or uom)


def _parse_qty_rate_cells(cells: list[str]) -> tuple[str, float, float]:
	if not cells:
		return "", 0.0, 0.0

	if len(cells) >= 3:
		uom = cells[0]
		qty = _parse_number(cells[1])
		rate = _parse_number(cells[2])
		return uom, qty, rate

	if len(cells) == 2:
		if _parse_number(cells[0]) and not _parse_number(cells[1]):
			return cells[1], _parse_number(cells[0]), 0.0
		return cells[0], _parse_number(cells[1]), 0.0

	return "", _parse_number(cells[0]), 0.0


def _parse_number(value: str) -> float:
	value = (value or "").strip().replace(",", "")
	if not value or value.lower() in {"-", "n/a", "na"}:
		return 0.0
	try:
		return flt(value)
	except Exception:
		return 0.0


def lines_to_boq_html(
	lines: list,
	include_totals: bool = True,
	company: str | None = None,
	include_vat: bool = True,
) -> str:
	"""Convert Quotation BOQ Line child rows to print HTML."""
	from construction_management.quotation_boq_hierarchy import resolve_hierarchy_mode
	from construction_management.quotation_boq_print_context import _group_boq_lines

	line_dicts = [
		row if isinstance(row, dict) else row.as_dict() if hasattr(row, "as_dict") else row.__dict__
		for row in lines
	]
	hierarchy_mode = resolve_hierarchy_mode(company, line_dicts)
	sections, total_excl = _group_boq_lines(line_dicts, hierarchy_mode=hierarchy_mode)

	rows: list[dict] = []
	for section in sections:
		if section.get("title"):
			rows.append({"kind": "section", "title": section["title"]})
		for parent in section.get("parents") or []:
			rows.append(
				{
					"kind": "parent",
					"parent_no": parent.get("no") or "",
					"description": parent.get("description") or "",
				}
			)
			for sub in parent.get("subs") or []:
				rows.append(
					{
						"kind": "sub",
						"sub_no": sub.get("sub_no") or "",
						"description": sub.get("description") or "",
						"uom": sub.get("uom") or "",
						"qty": sub.get("qty"),
						"rate": sub.get("rate_display"),
						"amount": sub.get("amount_display"),
						"is_fixed_rate": sub.get("include_in_total", True),
						"display_mode": sub.get("display_mode") or "Normal",
					}
				)

	return _render_boq_html(
		rows,
		include_totals=include_totals,
		total_excl=total_excl,
		include_vat=include_vat,
	)


def _render_boq_html(
	rows: list[dict],
	include_totals: bool = True,
	template_only: bool = False,
	total_excl: float = 0.0,
	include_vat: bool = True,
) -> str:
	if template_only:
		rows = [
			{"kind": "section", "title": "THERMAL AND MOISTURE PROTECTION TO SUBSTRUCTURE"},
			{
				"kind": "parent",
				"parent_no": "1",
				"description": "One layer of 1000 gauge polythene sheets; all in accordance with the drawings and specification",
			},
			{
				"kind": "sub",
				"sub_no": "A",
				"description": "To Raft Slab",
				"uom": "m2",
				"qty": "2,259.00",
				"rate": "-",
				"amount": "N/A",
			},
			{
				"kind": "sub",
				"sub_no": "B",
				"description": "Below Grade Slab",
				"uom": "m2",
				"qty": "1,800.00",
				"rate": "-",
				"amount": "N/A",
			},
			{
				"kind": "parent",
				"parent_no": "2",
				"description": "2 layers of 4mm thick SBS waterproofing membrane 180gm/m2 from Awazel or Equivalent",
			},
			{
				"kind": "sub",
				"sub_no": "A",
				"description": "Horizontally to Raft Slab",
				"uom": "m2",
				"qty": "2,184.00",
				"rate": "62.00",
				"amount": "135,408.00",
			},
		]

	empty = f'<td style="{BOQ_CELL_STYLE}"></td>'
	body_rows = []
	for row in rows:
		kind = row.get("kind")
		if kind == "section":
			# Six bordered cells (no colspan) so print engines keep vertical lines.
			body_rows.append(
				f'<tr><td style="{BOQ_CELL_STYLE}"></td>'
				f'<td style="{BOQ_SECTION_STYLE}">{html.escape(row.get("title") or "")}</td>'
				f"{empty}{empty}{empty}{empty}</tr>"
			)
		elif kind == "parent":
			parent_no = html.escape(str(row.get("parent_no") or ""))
			desc = html.escape(row.get("description") or "")
			sl = f"{parent_no}-" if parent_no else ""
			body_rows.append(
				f'<tr><td style="{BOQ_CELL_STYLE};text-align:center;">{sl}</td>'
				f'<td style="{BOQ_CELL_STYLE};font-weight:bold;">{desc}</td>'
				f"{empty}{empty}{empty}{empty}</tr>"
			)
		elif kind == "sub":
			sub_no = html.escape(str(row.get("sub_no") or ""))
			desc = html.escape(row.get("description") or "")
			uom = html.escape(str(row.get("uom") or ""))
			qty = _format_cell(row.get("qty"))
			rate = _format_cell(row.get("rate"))
			amount = _format_amount_cell(row)
			sl = f"{sub_no}." if sub_no else ""
			body_rows.append(
				f'<tr><td style="{BOQ_CELL_STYLE};text-align:center;">{sl}</td>'
				f'<td style="{BOQ_CELL_STYLE};padding-left:18px;">{desc}</td>'
				f'<td style="{BOQ_CELL_STYLE};text-align:center;">{uom}</td>'
				f'<td style="{BOQ_NUM_STYLE}">{qty}</td>'
				f'<td style="{BOQ_NUM_STYLE}">{rate}</td>'
				f'<td style="{BOQ_NUM_STYLE}">{amount}</td></tr>'
			)

	table_html = f"""
<table style="{BOQ_TABLE_STYLE}">
	<thead>
		<tr>
			<th style="{BOQ_HEAD_STYLE}">SL No.</th>
			<th style="{BOQ_HEAD_STYLE}">Description of Works</th>
			<th style="{BOQ_HEAD_STYLE}">Unit</th>
			<th style="{BOQ_HEAD_STYLE}">Approx. QTY.</th>
			<th style="{BOQ_HEAD_STYLE}">Rate AED</th>
			<th style="{BOQ_HEAD_STYLE}">Amount AED</th>
		</tr>
	</thead>
	<tbody>
		{''.join(body_rows)}
	</tbody>
</table>
""".strip()

	if not include_totals:
		return table_html

	if not include_vat:
		totals_html = f"""
<table style="width:55%;margin-left:auto;border-collapse:collapse;{BOQ_TABLE_STYLE}">
	<tr>
		<td style="{BOQ_CELL_STYLE};font-weight:bold;text-align:right;">Total Amount</td>
		<td style="{BOQ_NUM_STYLE}">{html.escape(f"{flt(total_excl):,.2f}")}</td>
	</tr>
</table>
<p><br></p>
""".strip()
		return f"{table_html}\n{totals_html}"

	vat = flt(total_excl) * 0.05
	incl = flt(total_excl) + vat
	totals_html = f"""
<table style="width:55%;margin-left:auto;border-collapse:collapse;{BOQ_TABLE_STYLE}">
	<tr>
		<td style="{BOQ_CELL_STYLE};font-weight:bold;text-align:right;">Total Amount Excl. VAT</td>
		<td style="{BOQ_NUM_STYLE}">{html.escape(f"{flt(total_excl):,.2f}")}</td>
	</tr>
	<tr>
		<td style="{BOQ_CELL_STYLE};font-weight:bold;text-align:right;">5% VAT</td>
		<td style="{BOQ_NUM_STYLE}">{html.escape(f"{vat:,.2f}")}</td>
	</tr>
	<tr>
		<td style="{BOQ_CELL_STYLE};font-weight:bold;text-align:right;">Total Amount Incl. VAT</td>
		<td style="{BOQ_NUM_STYLE}">{html.escape(f"{incl:,.2f}")}</td>
	</tr>
</table>
<p><br></p>
""".strip()

	return f"{table_html}\n{totals_html}"


def _format_cell(value) -> str:
	if value in (None, ""):
		return ""
	if isinstance(value, (int, float)):
		return html.escape(f"{flt(value):,.2f}")
	return html.escape(str(value))


def _format_amount_cell(row: dict) -> str:
	display_mode = row.get("display_mode") or "Normal"
	if display_mode == "N/A":
		return "N/A"
	if display_mode == "Rate Only":
		return "Rate only"
	return _format_cell(row.get("amount"))


def _get_estimation_source_items(estimation) -> list[dict]:
	boq_name = frappe.db.get_value(
		"Bill of Quantity",
		{"project_estimation": estimation.name},
		"name",
	)
	if boq_name:
		boq = frappe.get_doc("Bill of Quantity", boq_name)
		return [row.as_dict() for row in boq.get("items") or []]

	return [row.as_dict() for row in estimation.get("items") or []]


def _estimation_item_description(item: dict) -> str:
	parts = []
	for key in ("item", "activity_type", "task", "description"):
		value = item.get(key)
		if value:
			parts.append(str(value))
	return " - ".join(dict.fromkeys(parts)) if parts else "Item"
