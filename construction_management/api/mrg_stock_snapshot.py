"""Controlled import for the 07 October 2026 MRG stock snapshot."""

from __future__ import annotations

import csv
import hashlib
from collections import defaultdict
from io import StringIO
from pathlib import Path

import frappe
from frappe.utils import flt, getdate, now_datetime
from frappe.utils.file_manager import save_file
from openpyxl import load_workbook

from construction_management.api import controlled_procurement as catalog
from construction_management.api import item_tree_part_2_expanded_import as expanded_import


COMPANY = "M R G INSULATION WORKS L.L.C"
POSTING_DATE = "2026-10-07"
REPORT_WAREHOUSES_TO_REACTIVATE = ("1101-FIBREX WADI AL SAFA-2 - MRG",)
SOURCE_DIRECTORY = Path(__file__).resolve().parents[1] / "data"
ITEM_TREE_PATH = SOURCE_DIRECTORY / "mrg_item_tree_part_2_3.xlsx"
STOCK_REPORT_PATH = SOURCE_DIRECTORY / "mrg_stock_report_2026_10_07.xlsx"
CORRECTION_RATES = {"filler board": 39.0, "utility knife": 4.0}
# These rows were agreed corrections to the final Item Tree.  They are kept
# here because the supplied Sheet2 does not contain every corrected row.
CORRECTION_ITEM_METADATA = {
	"filler board": {"item_type": "Stockable", "stock_uom": "Nos"},
	"utility knife": {"item_type": "Consumable", "stock_uom": "Nos"},
	"fire blanket": {"item_type": "Consumable", "stock_uom": "Nos"},
	"safety goggles clear": {"item_type": "Consumable", "stock_uom": "PKT"},
	"safety jacket": {"item_type": "Consumable", "stock_uom": "Nos"},
}
CATALOG_CORRECTIONS = (
	{"item_code": "CP-AWAZEL_GROUT_EP_300_14_LT/SET", "item_name": "AWAZEL GROUT EP 300 14 LT/SET", "item_group": "GROUT AND TILES", "item_type": "Stockable", "stock_uom": "Nos", "supplier": "AWAZEL INTERNATIONAL CO LLC", "rate": 125},
	{"item_code": "FILLER BOARD 10MM THICK WITH 7CM CUTTING PCS", "item_name": "FILLER BOARD", "item_group": "PROTECTION BOARD", "item_type": "Stockable", "stock_uom": "Nos", "supplier": "SHERIDAN SPECIALIZED BUILDING PRODUCTS  L.L.C", "rate": 39},
	{"item_code": "FIRE BLANKET", "item_name": "FIRE BLANKET", "item_group": "HARDWARE ITEMS", "item_type": "Consumable", "stock_uom": "Nos", "supplier": "ROYAL TOOLS TRADING  L.L.C", "rate": 17},
	{"item_code": "CP-GAS_PIPE", "item_name": "GAS PIPE", "item_group": "WATERPROOFING GASS", "item_type": "Consumable", "stock_uom": "Nos", "supplier": "MMT TRADING LLC", "rate": 6.5},
	{"item_code": "HOSE CLIP", "item_name": "HOSE CLIP", "item_group": "WATERPROOFING GASS", "item_type": "Consumable", "stock_uom": "Nos", "supplier": "JANNAT AL FERDOS BUILDING MATERIALS TRADING LLC", "rate": 1.2},
	{"item_code": "MEASURING TAPE 10 MTR", "item_name": "MEASURING TAPE 10 MTR", "item_group": "HARDWARE ITEMS", "item_type": "Consumable", "stock_uom": "Nos", "supplier": "MMT TRADING LLC", "rate": 11.5},
	{"item_code": "CP-RHEOGROUT_EP_102_(30KG_KIT)", "item_name": "RHEOGROUT EP 102 (30KG KIT)", "item_group": "GROUT AND TILES", "item_type": "Stockable", "stock_uom": "Nos", "supplier": "Petrozo Energy L.L.C", "rate": 118},
	{"item_code": "CP-ROCKET_THINNER_2.5_LITR", "item_name": "ROCKET THINNER 2.5 LITR", "item_group": "HARDWARE ITEMS", "item_type": "Consumable", "stock_uom": "GALLON", "supplier": "ROYAL TOOLS TRADING  L.L.C", "rate": 52},
	{"item_code": "SAFETY GOGGLES CLEAR", "item_name": "SAFETY GOGGLES CLEAR", "item_group": "HARDWARE ITEMS", "item_type": "Consumable", "stock_uom": "PKT", "supplier": "MMT TRADING LLC", "rate": 15},
	{"item_code": "SAFETY JACKET", "item_name": "SAFETY JACKET", "item_group": "SAFETY ITEMS", "item_type": "Consumable", "stock_uom": "Nos", "supplier": "CITY RANGERS TRADING  LLC", "rate": 18},
	{"item_code": "SCRAPER 2''", "item_name": "SCRAPER 2''", "item_group": "HARDWARE ITEMS", "item_type": "Consumable", "stock_uom": "Nos", "supplier": "ROYAL TOOLS TRADING  L.L.C", "rate": 2},
	{"item_code": "SCRAPPER 3", "item_name": "SCRAPPER 3", "item_group": "HARDWARE ITEMS", "item_type": "Consumable", "stock_uom": "Nos", "supplier": "ROYAL TOOLS TRADING  L.L.C", "rate": 2.75},
	{"item_code": "SCREW DRIVER 10 FLAT", "item_name": "SCREW DRIVER 10 FLAT", "item_group": "HARDWARE ITEMS", "item_type": "Consumable", "stock_uom": "Nos", "supplier": "ROYAL TOOLS TRADING  L.L.C", "rate": 9},
	{"item_code": "SCREW DRIVER 10'' STAR", "item_name": "SCREW DRIVER 10'' STAR", "item_group": "HARDWARE ITEMS", "item_type": "Consumable", "stock_uom": "Nos", "supplier": "ROYAL TOOLS TRADING  L.L.C", "rate": 9},
	{"item_code": "UTILITY KNIFE", "item_name": "UTILITY KNIFE", "item_group": "HARDWARE ITEMS", "item_type": "Consumable", "stock_uom": "Nos", "supplier": "JANNAT AL FERDOS BUILDING MATERIALS TRADING LLC", "rate": 4},
)


def preflight(
	item_tree_path: str = str(ITEM_TREE_PATH),
	stock_report_path: str = str(STOCK_REPORT_PATH),
) -> dict:
	"""Validate and normalize the positive MRG stock-snapshot balances."""
	tree_rates, tree_types, tree_uoms = _tree_metadata(Path(item_tree_path))
	rows = _report_rows(Path(stock_report_path))
	items, non_stock_items = _catalog_items()
	warehouses = _warehouses()
	targets, held, errors = _targets(
		rows, items, non_stock_items, warehouses, tree_rates, tree_types, tree_uoms,
	)
	_apply_valuation_rates(targets, tree_rates)
	missing_rates = [target for target in targets if target["valuation_rate"] is None]
	if missing_rates:
		errors.extend({
			"reason": "No existing valuation rate or finalized Item Tree buying rate",
			"item": target["item_name"],
			"warehouse": target["warehouse"],
		} for target in missing_rates)
	return {
		"company": COMPANY,
		"posting_date": POSTING_DATE,
		"source_rows": len(rows),
		"target_lines": len(targets),
		"warehouses": len({target["warehouse"] for target in targets}),
		"held_rows": held,
		"errors": errors,
		"targets": targets,
	}


def activate_report_warehouses() -> list[str]:
	"""Re-enable the exact MRG warehouse included in the approved stock report."""
	activated = []
	for name in REPORT_WAREHOUSES_TO_REACTIVATE:
		warehouse = frappe.db.get_value("Warehouse", name, ["company", "disabled"], as_dict=True)
		if not warehouse or warehouse.company != COMPANY or not warehouse.disabled:
			continue
		document = frappe.get_doc("Warehouse", name)
		document.disabled = 0
		document.save(ignore_permissions=True)
		activated.append(name)
	return activated


def run(
	item_tree_path: str = str(ITEM_TREE_PATH),
	stock_report_path: str = str(STOCK_REPORT_PATH),
) -> dict:
	"""Create missing unambiguous masters and reconcile only changed MRG balances."""
	frappe.set_user("Administrator")
	item_tree = Path(item_tree_path)
	stock_report = Path(stock_report_path)
	result = preflight(str(item_tree), str(stock_report))
	catalog_requests = _stage_catalog_corrections()
	catalog_requests.extend(_create_missing_catalog_requests(item_tree, result["errors"]))
	if catalog_requests:
		result = preflight(str(item_tree), str(stock_report))
	if result["errors"]:
		audit_file = _save_audit(result, [], catalog_requests, None, preflight_failed=True)
		return _run_result(
			result, catalog_requests, [], [], audit_file.file_url, status="preflight_failed",
		)

	mismatches = _mismatched_targets(result["targets"])
	if not mismatches:
		audit_file = _save_audit(result, mismatches, catalog_requests, None)
		return _run_result(result, catalog_requests, [], mismatches, audit_file.file_url)

	documents = _create_reconciliations(mismatches)
	audit_file = _save_audit(result, mismatches, catalog_requests, documents)
	for source in (item_tree, stock_report):
		save_file(source.name, source.read_bytes(), documents[0].doctype, documents[0].name, is_private=1)
	for document in documents:
		document.submit()
	return _run_result(result, catalog_requests, documents, mismatches, audit_file.file_url)


def submit(
	item_tree_path: str = str(ITEM_TREE_PATH),
	stock_report_path: str = str(STOCK_REPORT_PATH),
) -> dict:
	"""Backward-compatible alias for the idempotent MRG snapshot job."""
	return run(item_tree_path, stock_report_path)


def verify(
	item_tree_path: str = str(ITEM_TREE_PATH),
	stock_report_path: str = str(STOCK_REPORT_PATH),
) -> dict:
	"""Compare every eligible report target with the resulting MRG Bin balance."""
	result = preflight(item_tree_path, stock_report_path)
	mismatches = _mismatched_targets(result["targets"])
	return {
		"target_lines": result["target_lines"],
		"warehouses": result["warehouses"],
		"held_rows": result["held_rows"],
		"preflight_errors": result["errors"],
		"mismatches": mismatches,
	}


def _tree_metadata(path: Path) -> tuple[dict[str, float], dict[str, str], dict[str, str]]:
	book = load_workbook(path, read_only=True, data_only=True)
	if "Sheet2" not in book:
		frappe.throw("The finalized Item Tree must contain Sheet2.")
	rates = dict(CORRECTION_RATES)
	types = {name: metadata["item_type"] for name, metadata in CORRECTION_ITEM_METADATA.items()}
	uoms = {name: metadata["stock_uom"] for name, metadata in CORRECTION_ITEM_METADATA.items()}
	for row in book["Sheet2"].iter_rows(min_row=3, values_only=True):
		if not row[1]:
			continue
		name = _key(row[1])
		types.setdefault(name, _allocation_type(row[5]))
		uoms.setdefault(name, catalog._catalog_uom(row[3]))
		if flt(row[4]) > 0:
			rates.setdefault(name, flt(row[4]))
	return rates, types, uoms


def _report_rows(path: Path) -> list[dict]:
	book = load_workbook(path, read_only=True, data_only=True)
	sheet = book.active
	return [
		{
			"source_row": number,
			"item_name": str(row[0]).strip(),
			"warehouse": str(row[2]).strip(),
			"report_uom": str(row[3] or "").strip(),
			"qty": flt(row[4]),
		}
		for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2)
		if row[0]
	]


def _catalog_items() -> tuple[dict[str, list[dict]], dict[str, list[dict]]]:
	items, non_stock_items = defaultdict(list), defaultdict(list)
	for row in frappe.get_all(
		"Item",
		filters={"disabled": 0, "controlled_procurement_catalog": 1},
		fields=["name", "item_name", "stock_uom", "controlled_item_type", "is_stock_item"],
		limit_page_length=0,
	):
		if row.controlled_item_type in {"Stockable", "Consumable"}:
			(items if row.is_stock_item else non_stock_items)[_key(row.item_name)].append(row)
	return items, non_stock_items


def _warehouses() -> set[str]:
	return set(frappe.get_all(
		"Warehouse",
		filters={"company": COMPANY, "is_group": 0, "disabled": 0},
		pluck="name",
		limit_page_length=0,
	))


def _targets(
	rows: list[dict],
	items: dict,
	non_stock_items: dict,
	warehouses: set[str],
	tree_rates: dict,
	tree_types: dict,
	tree_uoms: dict,
) -> tuple[list[dict], list[dict], list[dict]]:
	targets, held, errors = {}, [], []
	for row in rows:
		if row["qty"] < 0:
			held.append({**row, "reason": "Negative report balance held for separate adjustment"})
			continue
		if tree_types.get(_key(row["item_name"])) == "Asset":
			held.append({**row, "reason": "Finalized Item Tree classifies this row as an Asset"})
			continue
		if row["warehouse"] not in warehouses:
			errors.append({**row, "reason": "Active MRG warehouse not found"})
			continue
		item_key = _key(row["item_name"])
		matches = _canonical_catalog_matches(items.get(item_key, []), row["item_name"], tree_uoms.get(item_key))
		if not matches and non_stock_items.get(_key(row["item_name"])):
			held.append({**row, "reason": "Controlled master cannot enable stock after submitted transactions"})
			continue
		if len(matches) != 1:
			errors.append({**row, "reason": "Canonical controlled stock item was not uniquely resolved"})
			continue
		item = matches[0]
		if item_key not in tree_rates:
			errors.append({**row, "reason": "Finalized Item Tree buying rate not found"})
			continue
		key = (item.name, row["warehouse"])
		target = targets.setdefault(key, {
			"item_code": item.name,
			"item_name": item.item_name,
			"stock_uom": item.stock_uom,
			"warehouse": row["warehouse"],
			"qty": 0.0,
			"source_rows": [],
			"valuation_rate": None,
		})
		target["qty"] += row["qty"]
		target["source_rows"].append(row["source_row"])
	return sorted(targets.values(), key=lambda target: (target["warehouse"], target["item_code"])), held, errors


def _canonical_catalog_matches(matches: list[dict], item_name: str, finalized_uom: str | None) -> list[dict]:
	"""Use the finalized UOM and first generated CP code to ignore duplicate imports."""
	if not matches:
		return []
	if finalized_uom:
		matches = [item for item in matches if _key(item.stock_uom) == _key(finalized_uom)]
	if len(matches) <= 1:
		return matches
	canonical_code = _canonical_item_code(item_name)
	canonical = [item for item in matches if item.name == canonical_code]
	if len(canonical) == 1:
		return canonical
	non_collision_codes = [item for item in matches if not item.name.rsplit("-", 1)[-1].isdigit()]
	return non_collision_codes if len(non_collision_codes) == 1 else matches


def _canonical_item_code(item_name: str) -> str:
	"""Return the stable first code, without the collision suffix used by Frappe."""
	return "CP-" + frappe.scrub(item_name).upper()[:100]


def _apply_valuation_rates(targets: list[dict], tree_rates: dict) -> None:
	bins = {
		(row.item_code, row.warehouse): row.valuation_rate
		for row in frappe.get_all(
			"Bin",
			fields=["item_code", "warehouse", "valuation_rate"],
			limit_page_length=0,
		)
	}
	for target in targets:
		bin_key = (target["item_code"], target["warehouse"])
		existing_rate = flt(bins[bin_key]) if bin_key in bins else 0
		target["valuation_rate"] = existing_rate or tree_rates.get(_key(target["item_name"]))


def _mismatched_targets(targets: list[dict]) -> list[dict]:
	"""Return only balances that need a Stock Reconciliation on this rerun."""
	bins = {
		(row.item_code, row.warehouse): flt(row.actual_qty)
		for row in frappe.get_all("Bin", fields=["item_code", "warehouse", "actual_qty"], limit_page_length=0)
	}
	return [
		{**target, "actual_qty": bins.get((target["item_code"], target["warehouse"]), 0.0)}
		for target in targets
		if abs(bins.get((target["item_code"], target["warehouse"]), 0.0) - target["qty"]) > 0.000001
	]


def _stage_catalog_corrections() -> list[str]:
	"""Stage finalized corrections that were held by source-row conflicts."""
	payloads = [
		payload for payload in CATALOG_CORRECTIONS
		if not _is_current_controlled_item(payload)
	]
	if not payloads:
		return []
	for group in {payload["item_group"] for payload in payloads}:
		catalog._ensure_group(group)
	prepared = [catalog._catalog_request_row(payload, COMPANY) for payload in payloads]
	checksum = hashlib.sha256(
		"|".join(sorted(
			f"{row['item_code']}:{row['stock_uom']}:{row['supplier']}:{row['rate']}"
			for row in prepared
		)).encode()
	).hexdigest()
	existing = frappe.db.get_value(
		"Controlled Catalog Request", {"company": COMPANY, "checksum": checksum}, "name",
	)
	if existing:
		return [existing]
	request = frappe.get_doc({
		"doctype": "Controlled Catalog Request",
		"company": COMPANY,
		"status": "Pending Accounts",
		"source": "Workbook Import",
		"requested_by": frappe.session.user,
		"checksum": checksum,
		"items": prepared,
		"approval_log": [{
			"action": "Created",
			"actor": frappe.session.user,
			"actioned_on": now_datetime(),
			"comment": "MRG finalized Item Tree corrections staged for approval before opening stock.",
		}],
	})
	request.flags.controlled_catalog_request_api = True
	request.insert(ignore_permissions=True)
	return [request.name]


def _is_current_controlled_item(payload: dict) -> bool:
	item = frappe.db.get_value(
		"Item", payload["item_code"],
		["controlled_procurement_catalog", "controlled_item_type", "stock_uom", "is_stock_item"],
		as_dict=True,
	)
	return bool(
		item
		and item.controlled_procurement_catalog
		and item.controlled_item_type == payload["item_type"]
		and item.stock_uom == payload["stock_uom"]
		and item.is_stock_item
	)


def _create_missing_catalog_requests(item_tree: Path, errors: list[dict]) -> list[str]:
	"""Create one approval-controlled request for unambiguous report-item gaps."""
	missing_names = {
		_key(error["item_name"])
		for error in errors
		if error["reason"] == "Canonical controlled stock item was not uniquely resolved"
	}
	if not missing_names:
		return []

	selected, _, _, _ = expanded_import._prepare(item_tree)
	stock_items, non_stock_items = _catalog_items()
	candidates = defaultdict(list)
	for row in selected:
		if _key(row["item_name"]) in missing_names and row["action"] in {"Create", "Adopt"}:
			candidates[_key(row["item_name"])].append(row)
	payloads = []
	for name, rows in candidates.items():
		if len(rows) != 1 or stock_items.get(name) or non_stock_items.get(name):
			continue
		payload = rows[0]
		if not _catalog_candidate_has_compatible_uom(payload):
			_record_catalog_staging_error(errors, name, payload)
			continue
		payloads.append(payload)
	if not payloads:
		return []

	for group in {row["item_group"] for row in payloads}:
		catalog._ensure_group(group)
	# Validate every row before creating the request so one incompatible legacy Item
	# cannot prevent the rest of the unambiguous request from being staged.
	prepared = []
	for payload in payloads:
		try:
			prepared.append(catalog._catalog_request_row(payload, COMPANY))
		except frappe.ValidationError as error:
			_record_catalog_staging_error(errors, _key(payload["item_name"]), payload, str(error))
	if not prepared:
		return []
	checksum = hashlib.sha256(
		"|".join(sorted(f"{row['item_name']}:{row['stock_uom']}:{row['supplier']}:{row['rate']}" for row in prepared)).encode()
	).hexdigest()
	existing = frappe.db.get_value(
		"Controlled Catalog Request", {"company": COMPANY, "checksum": checksum}, "name"
	)
	if existing:
		return [existing]
	request = frappe.get_doc({
		"doctype": "Controlled Catalog Request",
		"company": COMPANY,
		"status": "Pending Accounts",
		"source": "Workbook Import",
		"requested_by": frappe.session.user,
		"checksum": checksum,
		"items": prepared,
		"approval_log": [{
			"action": "Created",
			"actor": frappe.session.user,
			"actioned_on": now_datetime(),
			"comment": "MRG stock snapshot preflight staged this request for Accounts approval.",
		}],
	})
	request.flags.controlled_catalog_request_api = True
	request.insert(ignore_permissions=True)
	return [request.name]


def _catalog_candidate_has_compatible_uom(payload: dict) -> bool:
	item_code = (payload.get("item_code") or "").strip()
	existing = frappe.db.get_value("Item", item_code, ["name", "stock_uom"], as_dict=True) if item_code else None
	if not existing:
		existing = frappe.db.get_value(
			"Item", {"item_name": payload["item_name"], "disabled": 0}, ["name", "stock_uom"], as_dict=True,
		)
	return not existing or existing.stock_uom == payload["stock_uom"]


def _record_catalog_staging_error(errors: list[dict], item_name: str, payload: dict, reason: str = "") -> None:
	if not reason:
		item_code = (payload.get("item_code") or "").strip()
		existing = frappe.db.get_value("Item", item_code, "name") if item_code else frappe.db.get_value(
			"Item", {"item_name": payload["item_name"], "disabled": 0}, "name",
		)
		reason = f"Existing item {existing} has another stock UOM; catalog request was not staged"
	for error in errors:
		if _key(error.get("item_name")) == item_name and error["reason"] == "Canonical controlled stock item was not uniquely resolved":
			error["reason"] = reason


def _create_reconciliations(targets: list[dict]) -> list:
	by_warehouse = defaultdict(list)
	for target in targets:
		by_warehouse[target["warehouse"]].append(target)
	documents = []
	for warehouse, warehouse_targets in sorted(by_warehouse.items()):
		document = frappe.new_doc("Stock Reconciliation")
		document.company = COMPANY
		document.purpose = "Stock Reconciliation"
		document.set_posting_time = 1
		document.posting_date = getdate(POSTING_DATE)
		document.posting_time = "00:00:00"
		for target in warehouse_targets:
			document.append("items", {
				"item_code": target["item_code"],
				"warehouse": warehouse,
				"qty": target["qty"],
				"valuation_rate": target["valuation_rate"],
			})
		document.flags.ignore_permissions = True
		document.insert()
		documents.append(document)
	return documents


def _save_audit(
	result: dict,
	mismatches: list[dict],
	catalog_requests: list[str],
	documents: list | None = None,
	preflight_failed: bool = False,
):
	"""Persist a run-level audit, including no-op checks and business-held rows."""
	document_by_warehouse = {
		doc.items[0].warehouse: doc.name for doc in (documents or [])
	}
	mismatch_keys = {(target["item_code"], target["warehouse"]) for target in mismatches}
	audit = StringIO()
	writer = csv.writer(audit)
	writer.writerow(["MRG Catalog and Stock Check", "Value"])
	writer.writerow(["Posting Date", POSTING_DATE])
	writer.writerow(["Eligible Report Balances", result["target_lines"]])
	writer.writerow(["Matched / No-op Balances", 0 if preflight_failed else result["target_lines"] - len(mismatches)])
	writer.writerow(["Balances Requiring Reconciliation", result["target_lines"] if preflight_failed else len(mismatches)])
	writer.writerow(["Catalog Requests", ", ".join(catalog_requests)])
	writer.writerow(["Held Report Rows", len(result["held_rows"])])
	writer.writerow([])
	writer.writerow(["Status", "Stock Reconciliation", "Warehouse", "Item Code", "Item Name", "Final UOM", "Report Qty", "Current Qty", "Valuation Rate", "Source Rows"])
	for target in result["targets"]:
		writer.writerow([
			"Blocked" if preflight_failed else ("Reconcile" if (target["item_code"], target["warehouse"]) in mismatch_keys else "Matched"),
			document_by_warehouse.get(target["warehouse"], ""), target["warehouse"], target["item_code"],
			target["item_name"], target["stock_uom"], target["qty"], target.get("actual_qty", target["qty"]), target["valuation_rate"],
			", ".join(map(str, target["source_rows"])),
		])
	writer.writerow([])
	writer.writerow(["Held Source Row", "Item Name", "Warehouse", "Report Qty", "Reason"])
	for held in result["held_rows"]:
		writer.writerow([held["source_row"], held["item_name"], held["warehouse"], held["qty"], held["reason"]])
	content = audit.getvalue().encode()
	filename = f"MRG_Stock_Catalog_Check_{POSTING_DATE}_{hashlib.sha256(content).hexdigest()[:8]}.csv"
	if documents:
		doctype, name = documents[0].doctype, documents[0].name
	else:
		doctype, name = "Company", COMPANY
	existing = frappe.get_all(
		"File",
		filters={"file_name": filename, "attached_to_doctype": doctype, "attached_to_name": name},
		pluck="name",
		limit_page_length=1,
	)
	if existing:
		return frappe.get_doc("File", existing[0])
	return save_file(filename, content, doctype, name, is_private=1)


def _run_result(
	result: dict, catalog_requests: list[str], documents: list, mismatches: list[dict], audit_file: str,
	status: str = "completed",
) -> dict:
	preflight_failed = status == "preflight_failed"
	return {
		"status": status,
		"posting_date": POSTING_DATE,
		"catalog_requests": catalog_requests,
		"stock_reconciliations": [document.name for document in documents],
		"submitted_documents": len(documents),
		"target_lines": result["target_lines"],
		"matched_balances": 0 if preflight_failed else result["target_lines"] - len(mismatches),
		"reconciliation_balances": result["target_lines"] if preflight_failed else len(mismatches),
		"held_rows": result["held_rows"],
		"preflight_error_count": len(result["errors"]),
		"audit_file": audit_file,
		"no_op": not catalog_requests and not documents,
	}


def _key(value) -> str:
	return " ".join(str(value or "").split()).casefold()


def _allocation_type(value) -> str:
	return {
		"material in stock": "Stockable",
		"consumable": "Consumable",
		"asset": "Asset",
		"service": "Service",
	}.get(_key(value), "")
